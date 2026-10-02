"""Tests for docket.infra.parsing.xlsx_wrapper.XlsxParser.

No Ollama/GPU dependency -- openpyxl parsing is fully local and fast, so
nothing here needs the ``integration`` marker (mirrors
``test_docling_wrapper.py``'s own framing).

openpyxl cannot itself write a formula cell *with* a cached result through
its normal API: setting ``cell.value`` to a formula string only ever writes
``<f>formula</f>`` with an empty ``<v></v>`` (openpyxl never evaluates
formulas -- see the module docstring in ``xlsx_wrapper.py``). To exercise the
"formula with a cached result" and "formula with a cached error" cases
realistically, ``_inject_cached_formula_result`` below post-processes the
saved ``.xlsx``'s raw worksheet XML to inject the ``<v>`` Excel itself would
write after recalculating and saving -- the only way to produce a fixture
that looks like a real "stale cache" workbook rather than one that was never
calculated at all (which is already exercised separately, as the "missing
cache" case, with no XML surgery needed).
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import openpyxl
import pytest

from docket.infra.parsing.xlsx_wrapper import (
    MISSING_FORMULA_CACHE_ERROR,
    ParseError,
    ParsedWorkbook,
    UnsupportedSpreadsheetFormatError,
    XlsxParser,
)


@pytest.fixture(scope="module")
def parser() -> XlsxParser:
    return XlsxParser()


def _inject_cached_formula_result(
    xlsx_path: Path,
    cell_ref: str,
    cached_value: str,
    *,
    is_error: bool = False,
    sheet_file: str = "xl/worksheets/sheet1.xml",
) -> None:
    with zipfile.ZipFile(xlsx_path, "r") as zf:
        items = {name: zf.read(name) for name in zf.namelist()}

    xml = items[sheet_file].decode("utf-8")
    pattern = re.compile(
        rf'<c r="{re.escape(cell_ref)}"([^>]*)><f>([^<]*)</f>(?:<v>[^<]*</v>|<v\s*/>)</c>'
    )
    match = pattern.search(xml)
    assert match is not None, (
        f"expected an empty-cache formula cell {cell_ref!r} in {sheet_file}; "
        "openpyxl's own write shape for a freshly-set formula must have changed"
    )
    formula = match.group(2)
    type_attr = ' t="e"' if is_error else ""
    replacement = f'<c r="{cell_ref}"{type_attr}><f>{formula}</f><v>{cached_value}</v></c>'
    xml = xml[: match.start()] + replacement + xml[match.end() :]
    items[sheet_file] = xml.encode("utf-8")

    with zipfile.ZipFile(xlsx_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in items.items():
            zf.writestr(name, data)


# ---------------------------------------------------------------------------
# Literal values with headers
# ---------------------------------------------------------------------------


def test_literal_value_reads_with_header_and_sheet_identity(
    parser: XlsxParser, tmp_path: Path
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Revenue"
    ws["A1"] = "Month"
    ws["B1"] = "Amount"
    ws["A2"] = "January"
    ws["B2"] = 120.5
    path = tmp_path / "literal.xlsx"
    wb.save(str(path))

    workbook = parser.parse("src1", path)

    assert isinstance(workbook, ParsedWorkbook)
    assert [s.name for s in workbook.sheets] == ["Revenue"]
    sheet = workbook.sheets[0]
    assert sheet.headers == {1: "Month", 2: "Amount"}

    [row] = sheet.rows
    assert row.row == 2
    by_coord = {c.coordinate: c for c in row.cells}
    assert by_coord["A2"].value == "January"
    assert by_coord["B2"].value == 120.5
    assert by_coord["B2"].is_formula is False
    assert by_coord["B2"].is_error is False


def test_missing_cell_is_not_reported_as_zero(parser: XlsxParser, tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Month"
    ws["B1"] = "Amount"
    ws["A2"] = "January"
    # B2 deliberately left empty.
    path = tmp_path / "sparse.xlsx"
    wb.save(str(path))

    workbook = parser.parse("src1", path)
    [row] = workbook.sheets[0].rows
    coords = {c.coordinate for c in row.cells}
    assert "B2" not in coords  # never fabricated as a 0/empty-string entry


# ---------------------------------------------------------------------------
# Formulas: literal / cached result / cached error / missing cache
# ---------------------------------------------------------------------------


def test_formula_with_cached_result_retains_formula_text_and_cached_value(
    parser: XlsxParser, tmp_path: Path
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Base"
    ws["A2"] = 100
    ws["B2"] = "=A2*1.2"
    path = tmp_path / "formula.xlsx"
    wb.save(str(path))
    _inject_cached_formula_result(path, "B2", "120")

    workbook = parser.parse("src1", path)
    [row] = workbook.sheets[0].rows
    cell = next(c for c in row.cells if c.coordinate == "B2")

    assert cell.is_formula is True
    assert cell.formula_text == "=A2*1.2"
    assert cell.value == 120
    assert cell.is_error is False


def test_formula_with_cached_error_is_flagged_not_treated_as_value(
    parser: XlsxParser, tmp_path: Path
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Numerator"
    ws["A2"] = 10
    ws["B2"] = "=A2/0"
    path = tmp_path / "formula_error.xlsx"
    wb.save(str(path))
    _inject_cached_formula_result(path, "B2", "#DIV/0!", is_error=True)

    workbook = parser.parse("src1", path)
    [row] = workbook.sheets[0].rows
    cell = next(c for c in row.cells if c.coordinate == "B2")

    assert cell.is_formula is True
    assert cell.formula_text == "=A2/0"
    assert cell.is_error is True
    assert cell.error_text == "#DIV/0!"


def test_formula_with_no_cached_result_is_flagged_as_missing_cache(
    parser: XlsxParser, tmp_path: Path
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Base"
    ws["A2"] = 5
    ws["B2"] = "=A2*2"  # never recalculated/saved by Excel -- no cache at all
    path = tmp_path / "formula_no_cache.xlsx"
    wb.save(str(path))

    workbook = parser.parse("src1", path)
    [row] = workbook.sheets[0].rows
    cell = next(c for c in row.cells if c.coordinate == "B2")

    assert cell.is_formula is True
    assert cell.value is None
    assert cell.is_error is True
    assert cell.error_text == MISSING_FORMULA_CACHE_ERROR


# ---------------------------------------------------------------------------
# Merged cells, hidden rows/columns, filters
# ---------------------------------------------------------------------------


def test_merged_cells_recorded_as_metadata_not_duplicated_values(
    parser: XlsxParser, tmp_path: Path
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Section"
    ws.merge_cells("A1:C1")
    ws["A1"] = "Q1 Totals"
    ws["A2"] = "x"
    path = tmp_path / "merged.xlsx"
    wb.save(str(path))

    workbook = parser.parse("src1", path)
    sheet = workbook.sheets[0]
    assert sheet.merged_ranges == ["A1:C1"]

    # Header row (row 1) is consumed as the header row, not emitted as a
    # data row -- merged-header propagation is covered by
    # test_merged_header_propagates_label_across_its_column_span below.
    assert sheet.header_row == 1


def test_merged_header_propagates_label_across_its_column_span(
    parser: XlsxParser, tmp_path: Path
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Q1"
    ws.merge_cells("A1:B1")
    ws["C1"] = "Notes"
    ws["A2"] = 1
    ws["B2"] = 2
    ws["C2"] = "ok"
    path = tmp_path / "merged_header.xlsx"
    wb.save(str(path))

    workbook = parser.parse("src1", path)
    sheet = workbook.sheets[0]
    assert sheet.headers[1] == "Q1"
    assert sheet.headers[2] == "Q1"  # propagated across the merge's span
    assert sheet.headers[3] == "Notes"


def test_merged_non_anchor_cell_has_no_independent_value(
    parser: XlsxParser, tmp_path: Path
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "H1"
    ws["B1"] = "H2"
    ws["A2"] = "Title spanning two cells"
    ws.merge_cells("A2:B2")
    path = tmp_path / "merged_row.xlsx"
    wb.save(str(path))

    workbook = parser.parse("src1", path)
    [row] = workbook.sheets[0].rows
    by_coord = {c.coordinate: c for c in row.cells}
    assert by_coord["A2"].merged_range == "A2:B2"
    assert by_coord["A2"].merged_non_anchor is False
    assert by_coord["B2"].merged_non_anchor is True
    assert by_coord["B2"].value is None


def test_hidden_row_and_hidden_column_are_flagged(parser: XlsxParser, tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Label"
    ws["B1"] = "Secret"
    ws["A2"] = "visible"
    ws["B2"] = "shown"
    ws["A3"] = "hidden-row"
    ws["B3"] = "also-hidden"
    ws.row_dimensions[3].hidden = True
    ws.column_dimensions["B"].hidden = True
    path = tmp_path / "hidden.xlsx"
    wb.save(str(path))

    workbook = parser.parse("src1", path)
    sheet = workbook.sheets[0]
    assert sheet.hidden_columns == ["B"]
    hidden_flags = {row.row: row.hidden for row in sheet.rows}
    assert hidden_flags == {2: False, 3: True}


def test_autofilter_marks_sheet_as_filtered(parser: XlsxParser, tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Label"
    ws["A2"] = "x"
    ws.auto_filter.ref = "A1:A2"
    path = tmp_path / "filtered.xlsx"
    wb.save(str(path))

    workbook = parser.parse("src1", path)
    assert workbook.sheets[0].filtered is True


# ---------------------------------------------------------------------------
# Multiple sheets
# ---------------------------------------------------------------------------


def test_multiple_sheets_each_parsed_independently(parser: XlsxParser, tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Alpha"
    ws1["A1"] = "X"
    ws1["A2"] = 1
    ws2 = wb.create_sheet("Beta")
    ws2["A1"] = "Y"
    ws2["A2"] = "two"
    path = tmp_path / "multi.xlsx"
    wb.save(str(path))

    workbook = parser.parse("src1", path)
    assert [s.name for s in workbook.sheets] == ["Alpha", "Beta"]
    assert workbook.sheets[0].headers == {1: "X"}
    assert workbook.sheets[1].headers == {1: "Y"}


# ---------------------------------------------------------------------------
# Scaling-hint and percentage detection
# ---------------------------------------------------------------------------


def test_thousands_scaling_hint_detected_from_number_format(
    parser: XlsxParser, tmp_path: Path
) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Amount"
    ws["A2"] = 3.5
    ws["A2"].number_format = "#,##0.00,"
    path = tmp_path / "scaled.xlsx"
    wb.save(str(path))

    workbook = parser.parse("src1", path)
    [row] = workbook.sheets[0].rows
    cell = next(c for c in row.cells if c.coordinate == "A2")
    assert cell.scaling_hint == "thousands"


def test_percentage_format_flagged(parser: XlsxParser, tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Rate"
    ws["A2"] = 0.15
    ws["A2"].number_format = "0%"
    path = tmp_path / "percent.xlsx"
    wb.save(str(path))

    workbook = parser.parse("src1", path)
    [row] = workbook.sheets[0].rows
    cell = next(c for c in row.cells if c.coordinate == "A2")
    assert cell.is_percentage is True


# ---------------------------------------------------------------------------
# Errors: corrupt files, unsupported formats
# ---------------------------------------------------------------------------


def test_corrupt_file_raises_parse_error_not_raw_openpyxl_exception(
    parser: XlsxParser, tmp_path: Path
) -> None:
    fake_xlsx = tmp_path / "fake.xlsx"
    fake_xlsx.write_text("this is plain text, not a real .xlsx file")

    with pytest.raises(ParseError) as exc_info:
        parser.parse("src-corrupt", fake_xlsx)

    err = exc_info.value
    assert err.source_id == "src-corrupt"
    assert err.path == fake_xlsx
    assert isinstance(err.cause, Exception)
    assert not isinstance(err.cause, ParseError)


def test_password_protected_style_file_raises_parse_error(
    parser: XlsxParser, tmp_path: Path
) -> None:
    # A password-protected .xlsx is actually an OLE2/CFB compound file, not a
    # zip at all -- openpyxl can't open it (nor should it try to guess a
    # password). This writes the real OLE2 magic bytes so the failure mode
    # matches what a genuinely encrypted workbook produces: a non-zip read
    # failure, not a parsed-but-empty result.
    encrypted_like = tmp_path / "protected.xlsx"
    encrypted_like.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)

    with pytest.raises(ParseError) as exc_info:
        parser.parse("src-protected", encrypted_like)

    assert exc_info.value.path == encrypted_like


def test_xlsm_extension_raises_unsupported_format_error(
    parser: XlsxParser, tmp_path: Path
) -> None:
    wb = openpyxl.Workbook()
    wb.active["A1"] = "irrelevant"
    path = tmp_path / "macro.xlsm"
    wb.save(str(path))

    with pytest.raises(UnsupportedSpreadsheetFormatError):
        parser.parse("src1", path)


def test_xls_extension_raises_unsupported_format_error(
    parser: XlsxParser, tmp_path: Path
) -> None:
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"not a real xls file either")

    with pytest.raises(UnsupportedSpreadsheetFormatError):
        parser.parse("src1", path)
