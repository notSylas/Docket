"""Tests for docket.infra.parsing.pptx_wrapper.PptxParser.

No Ollama/GPU dependency -- python-pptx parsing is fully local and fast, so
nothing here needs the ``integration`` marker (mirrors
``test_xlsx_wrapper.py``'s own framing).
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from docket.infra.parsing.pptx_wrapper import (
    ParseError,
    ParsedPresentation,
    PptxParser,
    UnsupportedPresentationFormatError,
)


@pytest.fixture(scope="module")
def parser() -> PptxParser:
    return PptxParser()


def _blank_slide(presentation: Presentation):
    return presentation.slides.add_slide(presentation.slide_layouts[6])


# ---------------------------------------------------------------------------
# Plain slide text
# ---------------------------------------------------------------------------


def test_slide_text_extracted_with_shape_identity(parser: PptxParser, tmp_path: Path) -> None:
    presentation = Presentation()
    slide = _blank_slide(presentation)
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    textbox.text_frame.text = "Quarterly results overview"
    path = tmp_path / "text.pptx"
    presentation.save(str(path))

    parsed = parser.parse("src1", path)

    assert isinstance(parsed, ParsedPresentation)
    assert len(parsed.slides) == 1
    slide_data = parsed.slides[0]
    assert slide_data.slide_number == 1
    [shape_data] = slide_data.text_shapes
    assert shape_data.text == "Quarterly results overview"
    assert shape_data.shape_id == textbox.shape_id
    assert shape_data.shape_name == textbox.name
    assert slide_data.tables == []
    assert slide_data.charts == []
    assert slide_data.notes_text is None


def test_blank_text_frame_is_not_reported_as_a_shape(parser: PptxParser, tmp_path: Path) -> None:
    presentation = Presentation()
    slide = _blank_slide(presentation)
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    textbox.text_frame.text = "   "  # whitespace-only
    path = tmp_path / "blank_text.pptx"
    presentation.save(str(path))

    parsed = parser.parse("src1", path)
    assert parsed.slides[0].text_shapes == []


# ---------------------------------------------------------------------------
# Tables
# ---------------------------------------------------------------------------


def test_table_on_slide_preserves_row_column_grid(parser: PptxParser, tmp_path: Path) -> None:
    presentation = Presentation()
    slide = _blank_slide(presentation)
    table_shape = slide.shapes.add_table(3, 2, Inches(1), Inches(1), Inches(4), Inches(2))
    table = table_shape.table
    table.cell(0, 0).text = "Region"
    table.cell(0, 1).text = "Revenue"
    table.cell(1, 0).text = "East"
    table.cell(1, 1).text = "100"
    table.cell(2, 0).text = "West"
    table.cell(2, 1).text = "200"
    path = tmp_path / "table.pptx"
    presentation.save(str(path))

    parsed = parser.parse("src1", path)
    [table_data] = parsed.slides[0].tables
    assert table_data.shape_id == table_shape.shape_id
    assert table_data.n_rows == 3
    assert table_data.n_cols == 2
    assert table_data.header_row == ["Region", "Revenue"]
    assert table_data.rows == [
        ["Region", "Revenue"],
        ["East", "100"],
        ["West", "200"],
    ]
    assert parsed.slides[0].text_shapes == []


def test_single_row_table_has_no_header_row(parser: PptxParser, tmp_path: Path) -> None:
    presentation = Presentation()
    slide = _blank_slide(presentation)
    table_shape = slide.shapes.add_table(1, 2, Inches(1), Inches(1), Inches(4), Inches(1))
    table = table_shape.table
    table.cell(0, 0).text = "Only"
    table.cell(0, 1).text = "Row"
    path = tmp_path / "single_row_table.pptx"
    presentation.save(str(path))

    parsed = parser.parse("src1", path)
    [table_data] = parsed.slides[0].tables
    assert table_data.n_rows == 1
    assert table_data.header_row is None
    assert table_data.rows == [["Only", "Row"]]


# ---------------------------------------------------------------------------
# Chart data
# ---------------------------------------------------------------------------


def test_chart_data_extracted_via_native_chart_api(parser: PptxParser, tmp_path: Path) -> None:
    presentation = Presentation()
    slide = _blank_slide(presentation)
    chart_data = CategoryChartData()
    chart_data.categories = ["East", "West"]
    chart_data.add_series("Revenue", (1.1, 2.2))
    chart_shape = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(4), Inches(3), chart_data
    )
    path = tmp_path / "chart.pptx"
    presentation.save(str(path))

    parsed = parser.parse("src1", path)
    [chart] = parsed.slides[0].charts
    assert chart.shape_id == chart_shape.shape_id
    assert chart.chart_type == "COLUMN_CLUSTERED"
    assert chart.categories == ["East", "West"]
    [series] = chart.series
    assert series.name == "Revenue"
    assert series.values == [1.1, 2.2]
    assert parsed.slides[0].tables == []
    assert parsed.slides[0].text_shapes == []


# ---------------------------------------------------------------------------
# Speaker notes -- kept distinct from slide text
# ---------------------------------------------------------------------------


def test_speaker_notes_kept_separate_from_slide_text(parser: PptxParser, tmp_path: Path) -> None:
    presentation = Presentation()
    slide = _blank_slide(presentation)
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    textbox.text_frame.text = "Visible on the slide"
    slide.notes_slide.notes_text_frame.text = "Only the presenter should see this"
    path = tmp_path / "notes.pptx"
    presentation.save(str(path))

    parsed = parser.parse("src1", path)
    slide_data = parsed.slides[0]
    assert slide_data.notes_text == "Only the presenter should see this"
    [shape_data] = slide_data.text_shapes
    assert shape_data.text == "Visible on the slide"
    # Notes text must never leak into slide_text or vice versa.
    assert "presenter" not in shape_data.text
    assert "Visible" not in slide_data.notes_text


def test_slide_with_no_notes_slide_has_none_notes_text(
    parser: PptxParser, tmp_path: Path
) -> None:
    presentation = Presentation()
    slide = _blank_slide(presentation)
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text_frame.text = "Hi"
    path = tmp_path / "no_notes.pptx"
    presentation.save(str(path))

    parsed = parser.parse("src1", path)
    assert parsed.slides[0].notes_text is None


def test_empty_notes_text_frame_is_none_not_empty_string(
    parser: PptxParser, tmp_path: Path
) -> None:
    presentation = Presentation()
    slide = _blank_slide(presentation)
    slide.notes_slide.notes_text_frame.text = "   "  # whitespace-only
    path = tmp_path / "whitespace_notes.pptx"
    presentation.save(str(path))

    parsed = parser.parse("src1", path)
    assert parsed.slides[0].notes_text is None


# ---------------------------------------------------------------------------
# No extractable content (image-only slide) -- must not crash
# ---------------------------------------------------------------------------


def test_picture_only_slide_parses_without_crashing_and_flags_unextracted(
    parser: PptxParser, tmp_path: Path
) -> None:
    pytest.importorskip("PIL")
    from PIL import Image

    presentation = Presentation()
    slide = _blank_slide(presentation)
    buf = io.BytesIO()
    Image.new("RGB", (10, 10), color="red").save(buf, format="PNG")
    buf.seek(0)
    picture_shape = slide.shapes.add_picture(buf, Inches(1), Inches(1), Inches(1), Inches(1))
    path = tmp_path / "picture_only.pptx"
    presentation.save(str(path))

    parsed = parser.parse("src1", path)
    slide_data = parsed.slides[0]
    assert slide_data.text_shapes == []
    assert slide_data.tables == []
    assert slide_data.charts == []
    assert slide_data.notes_text is None
    [unextracted] = slide_data.unextracted_shapes
    assert unextracted.shape_id == picture_shape.shape_id
    assert unextracted.shape_type == "PICTURE"


def test_fully_blank_slide_parses_without_crashing(parser: PptxParser, tmp_path: Path) -> None:
    presentation = Presentation()
    _blank_slide(presentation)
    path = tmp_path / "blank_slide.pptx"
    presentation.save(str(path))

    parsed = parser.parse("src1", path)
    slide_data = parsed.slides[0]
    assert slide_data.text_shapes == []
    assert slide_data.tables == []
    assert slide_data.charts == []
    assert slide_data.notes_text is None
    assert slide_data.unextracted_shapes == []


# ---------------------------------------------------------------------------
# Grouped shapes -- content nested inside a group must still be discovered
# ---------------------------------------------------------------------------


def test_text_inside_a_group_shape_is_still_discovered(
    parser: PptxParser, tmp_path: Path
) -> None:
    presentation = Presentation()
    slide = _blank_slide(presentation)
    textbox1 = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1))
    textbox1.text_frame.text = "First"
    textbox2 = slide.shapes.add_textbox(Inches(3), Inches(1), Inches(2), Inches(1))
    textbox2.text_frame.text = "Second"
    slide.shapes.add_group_shape([textbox1, textbox2])
    path = tmp_path / "grouped.pptx"
    presentation.save(str(path))

    parsed = parser.parse("src1", path)
    texts = {s.text for s in parsed.slides[0].text_shapes}
    assert texts == {"First", "Second"}


# ---------------------------------------------------------------------------
# Multiple slides
# ---------------------------------------------------------------------------


def test_multiple_slides_numbered_in_order(parser: PptxParser, tmp_path: Path) -> None:
    presentation = Presentation()
    slide1 = _blank_slide(presentation)
    slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "One"
    slide2 = _blank_slide(presentation)
    slide2.shapes.add_textbox(Inches(1), Inches(1), Inches(2), Inches(1)).text_frame.text = "Two"
    path = tmp_path / "multi_slide.pptx"
    presentation.save(str(path))

    parsed = parser.parse("src1", path)
    assert [s.slide_number for s in parsed.slides] == [1, 2]
    assert parsed.slides[0].text_shapes[0].text == "One"
    assert parsed.slides[1].text_shapes[0].text == "Two"


# ---------------------------------------------------------------------------
# Errors: corrupt files, unsupported formats
# ---------------------------------------------------------------------------


def test_corrupt_file_raises_parse_error_not_raw_pptx_exception(
    parser: PptxParser, tmp_path: Path
) -> None:
    fake_pptx = tmp_path / "fake.pptx"
    fake_pptx.write_text("this is plain text, not a real .pptx file")

    with pytest.raises(ParseError) as exc_info:
        parser.parse("src-corrupt", fake_pptx)

    err = exc_info.value
    assert err.source_id == "src-corrupt"
    assert err.path == fake_pptx
    assert isinstance(err.cause, Exception)
    assert not isinstance(err.cause, ParseError)


def test_ppt_extension_raises_unsupported_format_error(
    parser: PptxParser, tmp_path: Path
) -> None:
    path = tmp_path / "legacy.ppt"
    path.write_bytes(b"not a real legacy ppt file")

    with pytest.raises(UnsupportedPresentationFormatError):
        parser.parse("src1", path)
