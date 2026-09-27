"""Tests for `docket.ingestion.pipeline.IngestionPipeline` -- the end-to-end
incremental ingestion wiring: real Docling parsing (via small generated
`.docx` fixture files, same technique as `test_docling_wrapper.py`'s
real-file smoke test), `FakeInferenceGateway` for embeddings (no Ollama
dependency, so this stays a fast-ish unit test despite exercising real
Docling), and a real migrated SQLite DB + real LanceDB dir (both under
`tmp_path`).

This is the regression suite proving the whole pipeline -- not just
`IndexManager` in isolation (see `test_index_manager.py`) -- correctly skips
re-embedding unchanged files end to end, correctly reconciles stale chunks
when a file's content changes, and correctly isolates one bad file's failure
from the rest of the batch.
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import docx
import pytest
from sqlalchemy import Engine, select

from docket.db.engine import get_session_factory
from docket.db.models import Chunk, EvidenceVersion, IngestionJob, IngestionJobStatus
from docket.evidence.manager import EvidenceManager
from docket.evidence.store import ContentAddressedStore
from docket.index.fts_index import FtsIndexWriter
from docket.index.manager import IndexManager
from docket.index.vector_index import LanceIndexWriter
from docket.inference.gateway import FakeInferenceGateway
from docket.ingestion.pipeline import IngestionPipeline
from docket.parsing.docling_wrapper import DoclingParser
from docket.parsing.recipes import DEFAULT_SPLITTER, ChunkRecipe
from docket.sources.manager import SourceManager


def _write_docx(path: Path, heading: str, body: str) -> None:
    document = docx.Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(body)
    document.save(str(path))


@pytest.fixture(scope="module")
def parser() -> DoclingParser:
    # Constructed once per test module: Docling loads its models on
    # DocumentConverter() construction (expensive the first time).
    return DoclingParser()


@pytest.fixture()
def env(
    migrated_sqlite_engine: Engine, tmp_path: Path, parser: DoclingParser
) -> SimpleNamespace:
    session_factory = get_session_factory(migrated_sqlite_engine)

    store_root = tmp_path / "evidence_store"
    for sub in ("objects", "manifests", "quarantine", "trash"):
        (store_root / sub).mkdir(parents=True, exist_ok=True)
    store = ContentAddressedStore(store_root)
    evidence_manager = EvidenceManager(store, session_factory)

    gateway = FakeInferenceGateway()
    fts = FtsIndexWriter(migrated_sqlite_engine)
    vector = LanceIndexWriter(tmp_path / "lancedb")
    index_manager = IndexManager(fts, vector, gateway)

    chunk_recipe = ChunkRecipe(
        chunk_size=200,
        overlap=40,
        splitter=DEFAULT_SPLITTER,
        parser_name=parser.parser_name,
        parser_version=parser.parser_version,
    )

    pipeline = IngestionPipeline(
        session_factory=session_factory,
        evidence_manager=evidence_manager,
        parser=parser,
        index_manager=index_manager,
        chunk_recipe=chunk_recipe,
    )

    source_manager = SourceManager(session_factory)
    folder = tmp_path / "docs"
    folder.mkdir()
    source = source_manager.register_source(folder)

    return SimpleNamespace(
        session_factory=session_factory,
        gateway=gateway,
        fts=fts,
        vector=vector,
        pipeline=pipeline,
        source_manager=source_manager,
        source=source,
        folder=folder,
    )


def _chunk_ids_for_file(env: SimpleNamespace, file_path: Path) -> set[str]:
    """Chunk ids belonging to `file_path`'s *current* EvidenceVersion only
    -- a file's earlier (superseded) version keeps its own old Chunk rows
    around in SQLite for audit purposes even after reconciliation removes
    them from the indexes, so this must not match those too."""
    with env.session_factory() as session:
        rows = session.execute(
            select(Chunk.id)
            .join(EvidenceVersion, Chunk.evidence_version_id == EvidenceVersion.id)
            .where(
                EvidenceVersion.file_path == str(file_path),
                EvidenceVersion.is_current.is_(True),
            )
        ).scalars().all()
        return set(rows)


def _job_status(env: SimpleNamespace, job_id: str) -> IngestionJobStatus:
    with env.session_factory() as session:
        job = session.get(IngestionJob, job_id)
        return job.status


# ---------------------------------------------------------------------------
# First run: everything is new.
# ---------------------------------------------------------------------------


def test_first_run_ingests_all_files_and_indexes(env: SimpleNamespace) -> None:
    doc1 = env.folder / "doc1.docx"
    doc2 = env.folder / "doc2.docx"
    _write_docx(doc1, "Section One", "Alpha content about penguins and glaciers.")
    _write_docx(doc2, "Section Two", "Beta content about volcanoes and coffee.")

    result = env.pipeline.run_ingestion_for_source(env.source.id)

    assert result.status == "succeeded"
    assert result.files_processed == 2
    assert result.files_failed == 0
    assert {r.status for r in result.file_results} == {"ingested"}
    assert all(r.chunks_written > 0 for r in result.file_results)

    assert _job_status(env, result.job_id) == IngestionJobStatus.SUCCEEDED

    with env.session_factory() as session:
        chunk_count = len(
            session.execute(
                select(Chunk.id).where(Chunk.source_id == env.source.id)
            ).scalars().all()
        )
    assert chunk_count == sum(r.chunks_written for r in result.file_results)
    assert chunk_count > 0

    # Chunks actually landed in both indexes (not just SQLite).
    assert len(env.gateway.embed_calls) > 0
    for file_path in (doc1, doc2):
        chunk_ids = _chunk_ids_for_file(env, file_path)
        assert chunk_ids
        assert env.vector.existing_chunk_ids(chunk_ids) == chunk_ids
        assert env.fts.existing_chunk_ids(chunk_ids) == chunk_ids


# ---------------------------------------------------------------------------
# Second run, unchanged: the critical no-re-embed regression assertion.
# ---------------------------------------------------------------------------


def test_second_run_unchanged_files_are_skipped_and_not_reembedded(
    env: SimpleNamespace,
) -> None:
    doc1 = env.folder / "doc1.docx"
    doc2 = env.folder / "doc2.docx"
    _write_docx(doc1, "Section One", "Alpha content about penguins and glaciers.")
    _write_docx(doc2, "Section Two", "Beta content about volcanoes and coffee.")

    first = env.pipeline.run_ingestion_for_source(env.source.id)
    assert {r.status for r in first.file_results} == {"ingested"}
    embed_calls_after_first_run = len(env.gateway.embed_calls)
    assert embed_calls_after_first_run > 0

    second = env.pipeline.run_ingestion_for_source(env.source.id)

    assert second.status == "succeeded"
    assert second.files_processed == 2
    assert second.files_failed == 0
    assert {r.status for r in second.file_results} == {"unchanged"}
    assert all(r.chunks_written == 0 for r in second.file_results)

    # The whole-pipeline regression assertion: re-running over unchanged
    # files must not trigger a single additional embed call anywhere in the
    # stack (evidence layer -> parser -> chunker -> index manager).
    assert len(env.gateway.embed_calls) == embed_calls_after_first_run


# ---------------------------------------------------------------------------
# Content change: re-ingest + reconcile stale chunks.
# ---------------------------------------------------------------------------


def test_modified_file_is_reingested_and_stale_chunks_are_reconciled(
    env: SimpleNamespace,
) -> None:
    doc1 = env.folder / "doc1.docx"
    doc2 = env.folder / "doc2.docx"
    _write_docx(doc1, "Section One", "Alpha content about penguins and glaciers.")
    _write_docx(doc2, "Section Two", "Beta content about volcanoes and coffee.")

    env.pipeline.run_ingestion_for_source(env.source.id)
    old_doc1_chunk_ids = _chunk_ids_for_file(env, doc1)
    old_doc2_chunk_ids = _chunk_ids_for_file(env, doc2)
    assert old_doc1_chunk_ids and old_doc2_chunk_ids

    _write_docx(doc1, "Section One Revised", "Completely different words: orchestras and tides.")
    result = env.pipeline.run_ingestion_for_source(env.source.id)

    results_by_path = {r.path: r for r in result.file_results}
    assert results_by_path[doc1].status == "ingested"
    assert results_by_path[doc2].status == "unchanged"

    new_doc1_chunk_ids = _chunk_ids_for_file(env, doc1)
    assert new_doc1_chunk_ids
    assert new_doc1_chunk_ids.isdisjoint(old_doc1_chunk_ids)

    # Old doc1 chunks are gone from both indexes...
    assert env.vector.existing_chunk_ids(old_doc1_chunk_ids) == set()
    assert env.fts.existing_chunk_ids(old_doc1_chunk_ids) == set()
    # ...new doc1 chunks are indexed...
    assert env.vector.existing_chunk_ids(new_doc1_chunk_ids) == new_doc1_chunk_ids
    assert env.fts.existing_chunk_ids(new_doc1_chunk_ids) == new_doc1_chunk_ids
    # ...and doc2's untouched chunks were NOT swept up in the reconcile.
    assert env.vector.existing_chunk_ids(old_doc2_chunk_ids) == old_doc2_chunk_ids
    assert env.fts.existing_chunk_ids(old_doc2_chunk_ids) == old_doc2_chunk_ids


# ---------------------------------------------------------------------------
# One bad file must not abort the batch.
# ---------------------------------------------------------------------------


def test_one_corrupt_file_fails_without_aborting_the_batch(env: SimpleNamespace) -> None:
    doc1 = env.folder / "doc1.docx"
    doc2 = env.folder / "doc2.docx"
    corrupt = env.folder / "corrupt.docx"
    _write_docx(doc1, "Section One", "Alpha content about penguins and glaciers.")
    _write_docx(doc2, "Section Two", "Beta content about volcanoes and coffee.")
    corrupt.write_text("this is plain text, not a real .docx file")

    result = env.pipeline.run_ingestion_for_source(env.source.id)

    assert result.files_processed == 3
    assert result.files_failed == 1
    assert result.status == "partial"
    assert _job_status(env, result.job_id) == IngestionJobStatus.PARTIAL

    results_by_path = {r.path: r for r in result.file_results}
    assert results_by_path[doc1].status == "ingested"
    assert results_by_path[doc2].status == "ingested"
    assert results_by_path[corrupt].status == "failed"
    assert results_by_path[corrupt].error is not None


# ---------------------------------------------------------------------------
# Progress callback + interruption.
# ---------------------------------------------------------------------------


def test_progress_events_order_and_counts(env: SimpleNamespace) -> None:
    _write_docx(env.folder / "a.docx", "A", "Alpha content about penguins.")
    _write_docx(env.folder / "b.docx", "B", "Beta content about volcanoes.")
    events = []

    result = env.pipeline.run_ingestion_for_source(env.source.id, progress=events.append)

    assert [(e.kind, e.index, e.total) for e in events] == [
        ("start", 1, 2),
        ("done", 1, 2),
        ("start", 2, 2),
        ("done", 2, 2),
    ]
    assert events[0].result is None
    assert events[1].result.status == "ingested"
    assert [e.path.name for e in events[::2]] == ["a.docx", "b.docx"]
    assert result.files_processed == 2


def test_progress_callback_exception_is_swallowed(env: SimpleNamespace) -> None:
    _write_docx(env.folder / "a.docx", "A", "Alpha content about penguins.")

    def boom(event):
        raise RuntimeError("callback bug")

    result = env.pipeline.run_ingestion_for_source(env.source.id, progress=boom)
    assert result.status == "succeeded"
    assert result.files_processed == 1


def test_interrupted_run_finalizes_job_as_failed_and_reraises(
    env: SimpleNamespace, monkeypatch: pytest.MonkeyPatch
) -> None:
    _write_docx(env.folder / "a.docx", "A", "Alpha content about penguins.")

    def interrupt(source_id, path):
        raise KeyboardInterrupt

    monkeypatch.setattr(env.pipeline, "_ingest_one_file", interrupt)
    with pytest.raises(KeyboardInterrupt):
        env.pipeline.run_ingestion_for_source(env.source.id)

    with env.session_factory() as session:
        job = session.execute(select(IngestionJob)).scalars().one()
        assert job.status == IngestionJobStatus.FAILED
        assert job.error == "interrupted"
        assert job.finished_at is not None


def test_formula_regions_persist_and_legacy_backfill_keeps_chunks(env, monkeypatch):
    import json
    from docket.parsing.docling_wrapper import ParsedDocument
    path = env.folder / "formula.docx"
    _write_docx(path, "Formula", "An unreadable equation follows.")
    regions = [{"page_no": 2, "bbox": {"l": 1, "t": 2, "r": 3, "b": 4}}]
    parsed = ParsedDocument(text="Formula <!-- formula-not-decoded -->", source_path=path,
                            parser_name="fixture", parser_version="1", formula_regions=regions)
    calls = []

    def parse(source_id, source_path):
        calls.append(source_path)
        return parsed

    monkeypatch.setattr(env.pipeline._parser, "parse", parse)
    env.pipeline.run_ingestion_for_source(env.source.id)
    with env.session_factory() as session:
        version = session.execute(select(EvidenceVersion)).scalar_one()
        assert json.loads(version.formula_regions_json) == regions
        version.formula_regions_json = None
        before = list(session.execute(select(Chunk.id)).scalars())
        session.commit()
    result = env.pipeline.run_ingestion_for_source(env.source.id)
    with env.session_factory() as session:
        assert list(session.execute(select(Chunk.id)).scalars()) == before
        assert json.loads(session.execute(select(EvidenceVersion.formula_regions_json)).scalar_one()) == regions
    assert result.file_results[0].status == "unchanged"
    assert len(calls) == 2
    env.pipeline.run_ingestion_for_source(env.source.id)
    assert len(calls) == 2


# ---------------------------------------------------------------------------
# Page provenance (Chunk.page_start/page_end) -- visual retrieval checkpoint 1.
# ---------------------------------------------------------------------------


def test_chunk_page_start_and_page_end_persist_and_survive_readback(env, monkeypatch):
    """A mocked multi-page parse (same fixture-injection technique as the
    formula-regions test above) proves page_start/page_end make it all the
    way from `ParsedDocument.text_with_page_markers` through
    `chunk_document` into real `Chunk` rows in the database, and that they
    read back correctly from a *fresh* session (not just an in-memory
    artifact of the write)."""
    from docket.parsing.docling_wrapper import ParsedDocument

    path = env.folder / "multipage.docx"
    _write_docx(path, "Multi", "placeholder body")

    # Markers sit right before each heading (mirroring how docling_wrapper
    # inserts a marker right before the first matched text item on a new
    # page), so each section here lands entirely on one page.
    annotated = (
        "<!--PAGE:1-->\n"
        "# Section One\n"
        "Alpha bravo charlie delta echo foxtrot golf hotel india juliet.\n\n"
        "<!--PAGE:2-->\n"
        "# Section Two\n"
        "Kilo lima mike november oscar papa quebec romeo sierra tango.\n"
    )
    plain = annotated.replace("<!--PAGE:1-->", "").replace("<!--PAGE:2-->", "")
    parsed = ParsedDocument(
        text=plain,
        source_path=path,
        parser_name="fixture",
        parser_version="1",
        text_with_page_markers=annotated,
    )

    def parse(source_id, source_path):
        return parsed

    monkeypatch.setattr(env.pipeline._parser, "parse", parse)
    result = env.pipeline.run_ingestion_for_source(env.source.id)
    assert result.status == "succeeded"

    with env.session_factory() as session:
        rows = session.execute(
            select(Chunk.heading, Chunk.page_start, Chunk.page_end).order_by(Chunk.ordinal)
        ).all()
    assert rows == [("Section One", 1, 1), ("Section Two", 2, 2)]
    for heading, page_start, page_end in rows:
        assert page_start is not None and page_end is not None

    # Fresh session: proves this round-tripped through the DB, not just an
    # artifact of the write session's identity map.
    with env.session_factory() as session:
        readback = session.execute(
            select(Chunk.page_start, Chunk.page_end).order_by(Chunk.ordinal)
        ).all()
    assert readback == [(1, 1), (2, 2)]


def test_chunk_page_span_is_none_when_parser_gives_no_marker_info(env, monkeypatch):
    """A `ParsedDocument` built without `text_with_page_markers` (its
    default) must chunk exactly as before this checkpoint: every chunk gets
    page_start=page_end=None rather than a guessed page number."""
    from docket.parsing.docling_wrapper import ParsedDocument

    path = env.folder / "nopage.docx"
    _write_docx(path, "NoPage", "placeholder body")
    parsed = ParsedDocument(
        text="# Heading\nSome ordinary content with no page info at all.",
        source_path=path,
        parser_name="fixture",
        parser_version="1",
    )

    def parse(source_id, source_path):
        return parsed

    monkeypatch.setattr(env.pipeline._parser, "parse", parse)
    result = env.pipeline.run_ingestion_for_source(env.source.id)
    assert result.status == "succeeded"

    with env.session_factory() as session:
        rows = session.execute(select(Chunk.page_start, Chunk.page_end)).all()
    assert rows
    assert all(page_start is None and page_end is None for page_start, page_end in rows)


# ---------------------------------------------------------------------------
# Page images (EvidenceVersion.page_images_json + visual index) --
# visual retrieval checkpoint 2.
# ---------------------------------------------------------------------------


def _multipage_parsed_document(path, page_images: dict[int, bytes]):
    from docket.parsing.docling_wrapper import ParsedDocument

    return ParsedDocument(
        text="# Section One\nAlpha content.\n# Section Two\nBeta content.\n",
        source_path=path,
        parser_name="fixture",
        parser_version="1",
        page_images=page_images,
    )


def test_page_images_stored_and_indexed_when_visual_index_enabled(
    migrated_sqlite_engine, tmp_path, parser, monkeypatch
):
    """With `visual_index_enabled=True` and a `FakeInferenceGateway`,
    ingesting a small multi-page fixture must: populate
    `EvidenceVersion.page_images_json`, make each page's bytes retrievable
    from the `ContentAddressedStore`, and write rows (with the fake
    gateway's canned description) to the pages LanceDB table."""
    import json as json_mod

    from docket.config import Settings
    from docket.db.engine import get_session_factory
    from docket.db.models import EvidenceVersion
    from docket.evidence.manager import EvidenceManager
    from docket.evidence.store import ContentAddressedStore
    from docket.index.fts_index import FtsIndexWriter
    from docket.index.manager import IndexManager
    from docket.index.vector_index import LanceIndexWriter
    from docket.index.visual_index import LancePageIndexWriter
    from docket.inference.gateway import FakeInferenceGateway
    from docket.ingestion.pipeline import IngestionPipeline
    from docket.parsing.recipes import DEFAULT_SPLITTER, ChunkRecipe
    from docket.sources.manager import SourceManager

    session_factory = get_session_factory(migrated_sqlite_engine)
    store_root = tmp_path / "evidence_store"
    for sub in ("objects", "manifests", "quarantine", "trash"):
        (store_root / sub).mkdir(parents=True, exist_ok=True)
    store = ContentAddressedStore(store_root)
    evidence_manager = EvidenceManager(store, session_factory)

    gateway = FakeInferenceGateway(canned_description="a page describing penguins")
    fts = FtsIndexWriter(migrated_sqlite_engine)
    vector = LanceIndexWriter(tmp_path / "lancedb")
    visual_writer = LancePageIndexWriter(tmp_path / "lancedb")
    index_manager = IndexManager(fts, vector, gateway)

    chunk_recipe = ChunkRecipe(
        chunk_size=200,
        overlap=40,
        splitter=DEFAULT_SPLITTER,
        parser_name=parser.parser_name,
        parser_version=parser.parser_version,
    )

    pipeline = IngestionPipeline(
        session_factory=session_factory,
        evidence_manager=evidence_manager,
        parser=parser,
        index_manager=index_manager,
        chunk_recipe=chunk_recipe,
        gateway=gateway,
        visual_index_writer=visual_writer,
        settings=Settings(visual_index_enabled=True),
    )

    source_manager = SourceManager(session_factory)
    folder = tmp_path / "docs"
    folder.mkdir()
    source = source_manager.register_source(folder)

    path = folder / "multipage.docx"
    _write_docx(path, "Multi", "placeholder body")
    page_images = {1: b"fake-png-bytes-page-1", 2: b"fake-png-bytes-page-2"}
    monkeypatch.setattr(
        pipeline._parser, "parse", lambda source_id, source_path: _multipage_parsed_document(
            source_path, page_images
        )
    )

    result = pipeline.run_ingestion_for_source(source.id)
    assert result.status == "succeeded"

    with session_factory() as session:
        version = session.execute(select(EvidenceVersion)).scalar_one()
        assert version.page_images_json is not None
        hashes = json_mod.loads(version.page_images_json)
        assert set(hashes.keys()) == {"1", "2"}

    for page_no, image_bytes in page_images.items():
        content_hash = hashes[str(page_no)]
        assert store.get(content_hash) == image_bytes

    assert len(gateway.describe_image_calls) == 2
    described_images = {call["image_bytes"] for call in gateway.describe_image_calls}
    assert described_images == set(page_images.values())
    for call in gateway.describe_image_calls:
        assert call["model"] == "qwen2.5vl:7b"

    table = visual_writer.table
    assert table is not None
    rows = {row["page_no"]: row["description"] for row in table.search().to_list()}
    assert rows == {1: "a page describing penguins", 2: "a page describing penguins"}
    for row in table.search().to_list():
        assert row["evidence_version_id"] == version.id
        assert row["source_id"] == source.id


def test_page_images_disabled_by_default_no_vlm_or_index_calls(env, monkeypatch):
    """With `visual_index_enabled=False` (the default -- `env`'s pipeline was
    built without passing `gateway`/`visual_index_writer`/`settings` at all,
    same as every other test in this file), none of the visual-index
    machinery runs: no `describe_image`/extra `embed` calls, no pages table
    writes. Page images are still captured/stored (item 3/4 of the
    checkpoint is unconditional) -- only the VLM+embedding+LanceDB-write step
    is gated."""
    path = env.folder / "multipage.docx"
    _write_docx(path, "Multi", "placeholder body")
    page_images = {1: b"fake-png-bytes-page-1", 2: b"fake-png-bytes-page-2"}
    monkeypatch.setattr(
        env.pipeline._parser,
        "parse",
        lambda source_id, source_path: _multipage_parsed_document(source_path, page_images),
    )

    embed_calls_before = len(env.gateway.embed_calls)
    result = env.pipeline.run_ingestion_for_source(env.source.id)
    assert result.status == "succeeded"

    assert env.gateway.describe_image_calls == []
    # Chunk-text embedding still happens (unrelated to the visual index);
    # what must NOT happen is any *extra* embed call for a page description.
    # With a fake gateway and no page-description embeds, the count should
    # equal exactly the number of chunk-text embed calls made this run.
    assert len(env.gateway.embed_calls) - embed_calls_before == sum(
        r.chunks_written for r in result.file_results
    )

    import json as json_mod
    from docket.db.models import EvidenceVersion

    with env.session_factory() as session:
        version = session.execute(
            select(EvidenceVersion).where(EvidenceVersion.file_path == str(path))
        ).scalar_one()
        assert version.page_images_json is not None
        assert json_mod.loads(version.page_images_json) != {}

    # No "pages" table created at all under the same lancedb dir the
    # pipeline's own vector writer uses.
    assert "pages" not in env.vector._db.list_tables().tables


# ---------------------------------------------------------------------------
# Formula transcriptions (EvidenceVersion.formula_transcriptions_json) --
# Phase B checkpoint 1 of "verified formula transcription".
# ---------------------------------------------------------------------------


def _synthetic_page_png(width_px: int = 600, height_px: int = 800) -> bytes:
    from io import BytesIO
    from PIL import Image

    image = Image.new("RGB", (width_px, height_px), (255, 255, 255))
    buffer = BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


# One region comfortably above MIN_FORMULA_REGION_AREA_PT2 (50x20=1000 pt^2)
# and one comfortably below it (9x10=90 pt^2) -- both on the same page, same
# BOTTOMLEFT-origin/page-size shape `_formula_regions` actually produces.
_ABOVE_THRESHOLD_REGION = {
    "item_ref": "#/texts/1",
    "page_no": 1,
    "coordinate_origin": "BOTTOMLEFT",
    "page_width": 300.0,
    "page_height": 400.0,
    "bbox": {"l": 100.0, "t": 300.0, "r": 150.0, "b": 280.0},
}
_BELOW_THRESHOLD_REGION = {
    "item_ref": "#/texts/2",
    "page_no": 1,
    "coordinate_origin": "BOTTOMLEFT",
    "page_width": 300.0,
    "page_height": 400.0,
    "bbox": {"l": 10.0, "t": 20.0, "r": 19.0, "b": 10.0},
}


def _formula_parsed_document(path, page_images, formula_regions):
    from docket.parsing.docling_wrapper import ParsedDocument

    return ParsedDocument(
        text="# Section One\nAlpha content.\n# Section Two\nBeta content.\n",
        source_path=path,
        parser_name="fixture",
        parser_version="1",
        page_images=page_images,
        formula_regions=formula_regions,
    )


def _build_formula_pipeline(migrated_sqlite_engine, tmp_path, parser, *, formula_transcription_enabled):
    from docket.config import Settings
    from docket.db.engine import get_session_factory
    from docket.evidence.manager import EvidenceManager
    from docket.evidence.store import ContentAddressedStore
    from docket.index.fts_index import FtsIndexWriter
    from docket.index.manager import IndexManager
    from docket.index.vector_index import LanceIndexWriter
    from docket.inference.gateway import FakeInferenceGateway
    from docket.ingestion.pipeline import IngestionPipeline
    from docket.parsing.recipes import DEFAULT_SPLITTER, ChunkRecipe
    from docket.sources.manager import SourceManager

    session_factory = get_session_factory(migrated_sqlite_engine)
    store_root = tmp_path / "evidence_store"
    for sub in ("objects", "manifests", "quarantine", "trash"):
        (store_root / sub).mkdir(parents=True, exist_ok=True)
    store = ContentAddressedStore(store_root)
    evidence_manager = EvidenceManager(store, session_factory)

    gateway = FakeInferenceGateway(canned_description="V = IR (Ohm's law)")
    fts = FtsIndexWriter(migrated_sqlite_engine)
    vector = LanceIndexWriter(tmp_path / "lancedb")
    index_manager = IndexManager(fts, vector, gateway)

    chunk_recipe = ChunkRecipe(
        chunk_size=200,
        overlap=40,
        splitter=DEFAULT_SPLITTER,
        parser_name=parser.parser_name,
        parser_version=parser.parser_version,
    )

    pipeline = IngestionPipeline(
        session_factory=session_factory,
        evidence_manager=evidence_manager,
        parser=parser,
        index_manager=index_manager,
        chunk_recipe=chunk_recipe,
        gateway=gateway,
        settings=Settings(formula_transcription_enabled=formula_transcription_enabled),
    )

    source_manager = SourceManager(session_factory)
    folder = tmp_path / "docs"
    folder.mkdir()
    source = source_manager.register_source(folder)

    from types import SimpleNamespace

    return SimpleNamespace(
        session_factory=session_factory,
        gateway=gateway,
        pipeline=pipeline,
        source_manager=source_manager,
        source=source,
        folder=folder,
    )


def test_formula_transcriptions_populated_above_threshold_when_enabled(
    migrated_sqlite_engine, tmp_path, parser, monkeypatch
):
    """With `formula_transcription_enabled=True`, ingesting a fixture whose
    formula regions include one above and one below
    `MIN_FORMULA_REGION_AREA_PT2` must: call the VLM exactly once (only for
    the above-threshold region), and store both a populated
    `formula_transcriptions_json` (one entry, tagged with the canned
    transcription and model) built from that one call."""
    import json as json_mod

    from docket.db.models import EvidenceVersion
    from sqlalchemy import select as sa_select

    env = _build_formula_pipeline(
        migrated_sqlite_engine, tmp_path, parser, formula_transcription_enabled=True
    )

    path = env.folder / "formulas.docx"
    _write_docx(path, "Formulas", "placeholder body")
    page_images = {1: _synthetic_page_png()}
    formula_regions = [_ABOVE_THRESHOLD_REGION, _BELOW_THRESHOLD_REGION]
    monkeypatch.setattr(
        env.pipeline._parser,
        "parse",
        lambda source_id, source_path: _formula_parsed_document(
            source_path, page_images, formula_regions
        ),
    )

    result = env.pipeline.run_ingestion_for_source(env.source.id)
    assert result.status == "succeeded"

    # Only the above-threshold region should ever reach the gateway -- the
    # below-threshold ("trivial") region must never trigger a describe_image
    # call at all.
    assert len(env.gateway.describe_image_calls) == 1

    with env.session_factory() as session:
        version = session.execute(
            sa_select(EvidenceVersion).where(EvidenceVersion.file_path == str(path))
        ).scalar_one()
        assert version.formula_transcriptions_json is not None
        transcriptions = json_mod.loads(version.formula_transcriptions_json)

    assert len(transcriptions) == 1
    (entry,) = transcriptions
    assert entry["item_ref"] == "#/texts/1"
    assert entry["page_no"] == 1
    assert entry["transcription"] == "V = IR (Ohm's law)"
    assert entry["model"] == env.pipeline._settings.vision_model


def test_formula_transcriptions_disabled_by_default_no_gateway_calls_column_none(
    migrated_sqlite_engine, tmp_path, parser, monkeypatch
):
    """With `formula_transcription_enabled=False` (the default), ingesting
    the same fixture must make zero `describe_image` calls and leave
    `formula_transcriptions_json` as `None` -- same rigor as checkpoint 2's
    off-by-default proof (checking the actual call-count/DB-column
    evidence, not just that ingestion succeeded)."""
    import json as json_mod

    from docket.db.models import EvidenceVersion
    from sqlalchemy import select as sa_select

    env = _build_formula_pipeline(
        migrated_sqlite_engine, tmp_path, parser, formula_transcription_enabled=False
    )
    assert env.pipeline._settings.formula_transcription_enabled is False

    path = env.folder / "formulas.docx"
    _write_docx(path, "Formulas", "placeholder body")
    page_images = {1: _synthetic_page_png()}
    formula_regions = [_ABOVE_THRESHOLD_REGION, _BELOW_THRESHOLD_REGION]
    monkeypatch.setattr(
        env.pipeline._parser,
        "parse",
        lambda source_id, source_path: _formula_parsed_document(
            source_path, page_images, formula_regions
        ),
    )

    describe_calls_before = len(env.gateway.describe_image_calls)
    result = env.pipeline.run_ingestion_for_source(env.source.id)
    assert result.status == "succeeded"

    assert len(env.gateway.describe_image_calls) - describe_calls_before == 0
    assert env.gateway.describe_image_calls == []

    with env.session_factory() as session:
        version = session.execute(
            sa_select(EvidenceVersion).where(EvidenceVersion.file_path == str(path))
        ).scalar_one()
        assert version.formula_transcriptions_json is None
        # formula_regions_json is unrelated and unconditional -- still
        # populated regardless of the transcription flag.
        assert json_mod.loads(version.formula_regions_json) == formula_regions


def test_formula_transcription_backfill_reuses_stored_page_images_no_reparse(
    migrated_sqlite_engine, tmp_path, parser, monkeypatch
):
    """Simulates turning `formula_transcription_enabled` on for a source
    that was already ingested with it off: formula_regions_json and
    page_images_json are already populated, only formula_transcriptions_json
    is missing. The backfill branch must transcribe using the already-
    stored page images/regions WITHOUT calling `parser.parse` again (no
    re-render)."""
    import json as json_mod

    from docket.db.models import EvidenceVersion
    from sqlalchemy import select as sa_select

    env = _build_formula_pipeline(
        migrated_sqlite_engine, tmp_path, parser, formula_transcription_enabled=False
    )

    path = env.folder / "formulas.docx"
    _write_docx(path, "Formulas", "placeholder body")
    page_images = {1: _synthetic_page_png()}
    formula_regions = [_ABOVE_THRESHOLD_REGION, _BELOW_THRESHOLD_REGION]
    parse_calls = {"count": 0}

    def _fake_parse(source_id, source_path):
        parse_calls["count"] += 1
        return _formula_parsed_document(source_path, page_images, formula_regions)

    monkeypatch.setattr(env.pipeline._parser, "parse", _fake_parse)

    # First run: flag off. formula_regions_json/page_images_json get
    # populated; formula_transcriptions_json stays None.
    result = env.pipeline.run_ingestion_for_source(env.source.id)
    assert result.status == "succeeded"
    assert parse_calls["count"] == 1
    assert env.gateway.describe_image_calls == []

    # Flip the flag on and re-run against the *same unchanged file* --
    # EvidenceManager sees identical bytes, so this exercises the
    # unchanged-file backfill branch, not a fresh ingest.
    env.pipeline._settings.formula_transcription_enabled = True
    result2 = env.pipeline.run_ingestion_for_source(env.source.id)
    assert result2.status == "succeeded"
    assert result2.file_results[0].status == "unchanged"

    # The backfill must reuse the already-stored regions/page images --
    # parser.parse must NOT have been called again.
    assert parse_calls["count"] == 1

    assert len(env.gateway.describe_image_calls) == 1

    with env.session_factory() as session:
        version = session.execute(
            sa_select(EvidenceVersion).where(EvidenceVersion.file_path == str(path))
        ).scalar_one()
        assert version.formula_transcriptions_json is not None
        transcriptions = json_mod.loads(version.formula_transcriptions_json)
    assert len(transcriptions) == 1
    assert transcriptions[0]["item_ref"] == "#/texts/1"
