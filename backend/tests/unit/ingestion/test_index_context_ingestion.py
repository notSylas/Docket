"""End to end: ingestion embeds/FTS-indexes the context prefix, everything
stored or quoted stays verbatim, and `docket reindex` rebuilds identical rows."""

from __future__ import annotations

from types import SimpleNamespace

import docx
from sqlalchemy import select

from conftest import _write_xlsx
from docket.core.db.models import Chunk
from docket.infra.index.manifest import read_manifest
from docket.infra.index.reindex import reindex
from docket.infra.inference.gateway import FakeInferenceGateway
from docket.infra.retrieval.resolver import EvidenceResolver


def _write_nested_docx(path) -> None:
    document = docx.Document()
    document.add_heading("Zanzibar Expedition", level=1)
    document.add_heading("Provisions", level=2)
    document.add_paragraph("Bring rice and lentils for the trek.")
    document.save(str(path))


def _fts_rows(env) -> dict[str, str]:
    with env.session_factory() as session:
        rows = session.connection().exec_driver_sql("SELECT chunk_id, text FROM fts_chunks")
        return {r[0]: r[1] for r in rows}


def _fts_match(env, query: str) -> set[str]:
    with env.session_factory() as session:
        rows = session.connection().exec_driver_sql(
            "SELECT chunk_id FROM fts_chunks WHERE fts_chunks MATCH ?", (query,)
        )
        return {r[0] for r in rows}


def _ingest(env):
    _write_nested_docx(env.folder / "trek_notes.docx")
    _write_xlsx(
        env.folder / "payroll_ledger.xlsx", "Salaries", ["name", "amount"], [["ann", 5], ["bo", 7]]
    )
    result = env.pipeline.run_ingestion_for_source(env.source.id)
    assert result.status == "succeeded"
    with env.session_factory() as session:
        return {c.id: c for c in session.execute(select(Chunk)).scalars()}


def test_prefix_is_indexed_and_stored_text_is_verbatim(env: SimpleNamespace) -> None:
    chunks = _ingest(env)
    fts = _fts_rows(env)
    vector_rows = {r["chunk_id"]: r for r in env.vector._open_table().search().to_list()}
    docx_chunk = next(c for c in chunks.values() if "lentils" in c.text)
    xlsx_chunk = next(c for c in chunks.values() if "ann" in c.text and "payroll" not in c.text)

    expected_docx = f"trek_notes.docx > Zanzibar Expedition > Provisions\n{docx_chunk.text}"
    assert fts[docx_chunk.id] == expected_docx
    assert fts[xlsx_chunk.id].startswith("payroll_ledger.xlsx > Salaries\n")
    assert expected_docx in env.gateway.embed_calls

    for chunk in chunks.values():
        # No directory or user name ever reaches the indexed text.
        assert str(env.folder) not in fts[chunk.id]
        # Stored and LanceDB text are the verbatim chunk text.
        assert vector_rows[chunk.id]["text"] == chunk.text
        assert not chunk.text.startswith(("trek_notes", "payroll_ledger"))

    # Words only in the file name / an ancestor heading now retrieve the chunk
    # (the lentils chunk's own text has neither "zanzibar" nor "trek_notes").
    assert "zanzibar" not in docx_chunk.text.lower()
    assert docx_chunk.id in _fts_match(env, "zanzibar")
    assert docx_chunk.id in _fts_match(env, "trek_notes")
    assert xlsx_chunk.id in _fts_match(env, "payroll_ledger")


def test_resolver_quotes_verbatim_chunk_text(env: SimpleNamespace) -> None:
    chunks = _ingest(env)
    resolver = EvidenceResolver(env.session_factory)
    for chunk_id, chunk in chunks.items():
        assert resolver.resolve(chunk_id).text == chunk.text


def test_reindex_rebuilds_identical_fts_text_and_vectors(
    env: SimpleNamespace, migrated_sqlite_engine, tmp_path
) -> None:
    chunks = _ingest(env)
    fts_before = _fts_rows(env)
    vec_before = {
        r["chunk_id"]: (r["text"], list(r["vector"]))
        for r in env.vector._open_table().search().to_list()
    }

    manifest_path = tmp_path / "index_manifest.json"
    reindex(
        engine=migrated_sqlite_engine,
        db_path=tmp_path / "lancedb",
        gateway=FakeInferenceGateway(),
        manifest_path=manifest_path,
        embed_model="m1",
        sleep=lambda s: None,
    )

    assert _fts_rows(env) == fts_before
    import lancedb

    table = lancedb.connect(str(tmp_path / "lancedb")).open_table("chunks")
    vec_after = {r["chunk_id"]: (r["text"], list(r["vector"])) for r in table.search().to_list()}
    assert set(vec_after) == set(vec_before) == set(chunks)
    for chunk_id, (text, vector) in vec_after.items():
        assert text == chunks[chunk_id].text
        assert vector == vec_before[chunk_id][1]
    assert read_manifest(manifest_path).index_text_version == 1
