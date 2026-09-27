"""Shared fixtures/helpers for the ingestion-pipeline tests (mirrors
`tests/unit/eval/conftest.py`'s precedent for this package's tests): real
Docling parsing (via small generated `.docx` fixture files, same technique
as `test_docling_wrapper.py`'s real-file smoke test), `FakeInferenceGateway`
for embeddings (no Ollama dependency, so this stays a fast-ish unit test
despite exercising real Docling), and a real migrated SQLite DB + real
LanceDB dir (both under `tmp_path`).
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import docx
import pytest
from sqlalchemy import Engine, select

from docket.core.db.engine import get_session_factory
from docket.core.db.models import Chunk, EvidenceVersion, IngestionJob, IngestionJobStatus
from docket.infra.evidence.manager import EvidenceManager
from docket.infra.evidence.store import ContentAddressedStore
from docket.infra.index.fts_index import FtsIndexWriter
from docket.infra.index.manager import IndexManager
from docket.infra.index.vector_index import LanceIndexWriter
from docket.inference.gateway import FakeInferenceGateway
from docket.ingestion.pipeline import IngestionPipeline
from docket.parsing.docling_wrapper import DoclingParser
from docket.parsing.recipes import DEFAULT_SPLITTER, ChunkRecipe
from docket.sources.manager import SourceManager


def _write_docx(path: Path, heading: str, body: str) -> None:
    document = docx.Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(body)
    document.save(str(path))


@pytest.fixture(scope="module")
def parser() -> DoclingParser:
    # Constructed once per test module: Docling loads its models on
    # DocumentConverter() construction (expensive the first time).
    return DoclingParser()


@pytest.fixture()
def env(
    migrated_sqlite_engine: Engine, tmp_path: Path, parser: DoclingParser
) -> SimpleNamespace:
    session_factory = get_session_factory(migrated_sqlite_engine)

    store_root = tmp_path / "evidence_store"
    for sub in ("objects", "manifests", "quarantine", "trash"):
        (store_root / sub).mkdir(parents=True, exist_ok=True)
    store = ContentAddressedStore(store_root)
    evidence_manager = EvidenceManager(store, session_factory)

    gateway = FakeInferenceGateway()
    fts = FtsIndexWriter(migrated_sqlite_engine)
    vector = LanceIndexWriter(tmp_path / "lancedb")
    index_manager = IndexManager(fts, vector, gateway)

    chunk_recipe = ChunkRecipe(
        chunk_size=200,
        overlap=40,
        splitter=DEFAULT_SPLITTER,
        parser_name=parser.parser_name,
        parser_version=parser.parser_version,
    )

    pipeline = IngestionPipeline(
        session_factory=session_factory,
        evidence_manager=evidence_manager,
        parser=parser,
        index_manager=index_manager,
        chunk_recipe=chunk_recipe,
    )

    source_manager = SourceManager(session_factory)
    folder = tmp_path / "docs"
    folder.mkdir()
    source = source_manager.register_source(folder)

    return SimpleNamespace(
        session_factory=session_factory,
        gateway=gateway,
        fts=fts,
        vector=vector,
        pipeline=pipeline,
        source_manager=source_manager,
        source=source,
        folder=folder,
    )


def _chunk_ids_for_file(env: SimpleNamespace, file_path: Path) -> set[str]:
    """Chunk ids belonging to `file_path`'s *current* EvidenceVersion only
    -- a file's earlier (superseded) version keeps its own old Chunk rows
    around in SQLite for audit purposes even after reconciliation removes
    them from the indexes, so this must not match those too."""
    with env.session_factory() as session:
        rows = session.execute(
            select(Chunk.id)
            .join(EvidenceVersion, Chunk.evidence_version_id == EvidenceVersion.id)
            .where(
                EvidenceVersion.file_path == str(file_path),
                EvidenceVersion.is_current.is_(True),
            )
        ).scalars().all()
        return set(rows)


def _job_status(env: SimpleNamespace, job_id: str) -> IngestionJobStatus:
    with env.session_factory() as session:
        job = session.get(IngestionJob, job_id)
        return job.status
