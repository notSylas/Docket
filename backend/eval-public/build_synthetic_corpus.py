"""Deterministically generate the synthetic evaluation corpus.

Everything here is fictional: "Tarnwick Provisions Pvt Ltd" and its people do
not exist. Values are literal numbers (no formulas) so ground truth is known by
construction. Run from anywhere:

    python eval-public/build_synthetic_corpus.py

and commit the regenerated files under eval-public/corpus-synthetic/.
Traps deliberately built in (see eval-public/README.md):
  * three near-identical revenue workbooks (FY2024-25 actual in INR, FY2025-26
    actual in INR thousands, FY2025-26 budget in INR thousands);
  * quarter-subtotal and TOTAL rows (naive summing double counts), units note
    (header text in one workbook, separate Notes sheet in the others),
    percentage column, blank vs 0 cells, one hidden row;
  * Expenses-FY2025-26.xlsx keeps a merged title row above the header row on
    purpose: the xlsx parser takes the FIRST non-empty row as the header, so its
    chunks are labelled with the title instead of column names (a real-world
    parser trap; questions on it are tagged in gold-extended.yaml);
  * a deck and a memo that overlap with the workbooks (some numbers agree,
    some differ);
  * people/ is the folder intended to be revoked.
"""

from __future__ import annotations

import datetime as dt
import re
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.util import Inches

OUT = Path(__file__).resolve().parent / "corpus-synthetic"
FIXED = dt.datetime(2025, 10, 1, 9, 0, 0)
PRODUCTS = ["Marigold Chai", "Saffron Rusk", "Juniper Masala"]

# (label, chai, rusk, masala, mom growth fraction or None). None cells stay blank.
Q = "Q1 subtotal", "Q2 subtotal"


def _revenue_rows(months: dict[str, tuple], units_note: str) -> list[list]:
    """months: name -> (chai, rusk, masala) with None for blank. Builds rows with
    quarter subtotals and a grand TOTAL, plus MoM growth fraction."""
    rows: list[list] = []
    prev_total = None
    sums = {"Q1": [0, 0, 0], "Q2": [0, 0, 0]}
    # a product line with no figure at all in the period keeps blank (not 0) subtotals
    absent = [all(months[m][j] is None for m in months) for j in range(3)]
    order = ["Apr", "May", "Jun", "Jul", "Aug", "Sep"]
    for i, m in enumerate(order):
        vals = months[m]
        total = sum(v or 0 for v in vals)
        growth = None if prev_total is None else round((total - prev_total) / prev_total, 3)
        prev_total = total
        rows.append([m, *vals, total, growth])
        q = "Q1" if i < 3 else "Q2"
        for j, v in enumerate(vals):
            sums[q][j] += v or 0
        if m in ("Jun", "Sep"):
            s = sums[q]
            rows.append([f"{q} subtotal", *[None if absent[j] else s[j] for j in range(3)], sum(s), None])
    g = [sums["Q1"][j] + sums["Q2"][j] for j in range(3)]
    rows.append(["TOTAL (Apr-Sep)", *[None if absent[j] else g[j] for j in range(3)], sum(g), None])
    return rows


def _save_wb(wb: Workbook, path: Path) -> None:
    wb.properties.creator = "Tarnwick Finance (synthetic)"
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def revenue_workbook(path: Path, units_note: str, months: dict, hidden_memo: tuple | None,
                     units_in_header: bool, label: str = "", sheet: str = "Monthly Revenue") -> None:
    """Header sits in row 1 (so the parser sees real column names). Units are
    either in the header text or only on a separate "Notes" sheet."""
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    suffix = f" ({units_note})" if units_in_header else ""
    ws.append(["Month", *[p + suffix for p in PRODUCTS], "Total" + suffix, "MoM Growth %"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in _revenue_rows(months, units_note):
        ws.append(r)
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=6).number_format = "0.0%"
        for col in range(2, 6):
            ws.cell(row=row, column=col).number_format = "#,##0"
    if hidden_memo:
        ws.append(list(hidden_memo))
        ws.row_dimensions[ws.max_row].hidden = True
    if not units_in_header:
        notes = wb.create_sheet("Notes")
        notes.append(["Note", "Detail"])
        notes.append(["Units", f"{units_note} ({label})"])
        notes.append(["Fiscal year", "April to March"])
    _save_wb(wb, path)


def build_finance() -> None:
    d = OUT / "finance"
    # FY2024-25 actual, full INR. Juniper Masala did not exist yet: column left blank.
    fy24 = {
        "Apr": (3650000, 2010000, None), "May": (3820000, 2090000, None),
        "Jun": (3975000, 2150000, None), "Jul": (4040000, 2205000, None),
        "Aug": (4215000, 2330000, None), "Sep": (4090000, 2260000, None),
    }
    revenue_workbook(d / "Revenue-FY2024-25.xlsx", "INR", fy24, None, units_in_header=True)
    # FY2025-26 actual, INR thousands. Masala: April blank (not launched), May 0 (stock-out).
    fy25 = {
        "Apr": (4120, 2310, None), "May": (4380, 2455, 0), "Jun": (4610, 2590, 780),
        "Jul": (4725, 2640, 915), "Aug": (4905, 2712, 1043), "Sep": (4388, 2540, 1198),
    }
    revenue_workbook(d / "Revenue-FY2025-26.xlsx", "INR thousands", fy25,
                     ("Memo: one-off export order (not in totals)", None, None, None, 1250, None),
                     units_in_header=False, label="actuals")
    # FY2025-26 budget, INR thousands. Masala budgeted at 0 in Apr-May.
    bud = {
        "Apr": (4000, 2250, 0), "May": (4200, 2350, 0), "Jun": (4500, 2500, 700),
        "Jul": (4650, 2600, 850), "Aug": (4800, 2650, 1000), "Sep": (4900, 2700, 1100),
    }
    revenue_workbook(d / "Revenue-FY2025-26-Budget.xlsx", "INR thousands", bud, None, units_in_header=False, label="budget")

    wb = Workbook()
    ws = wb.active
    ws.title = "Operating Costs"
    ws["A1"] = "Tarnwick Provisions - Expenses FY2025-26, Q2"
    ws.merge_cells("A1:F1")
    ws["A2"] = "All figures in INR thousands"
    ws.append(["Month", "Staff", "Logistics", "Marketing", "Total", "% of Revenue"])
    for r in [["Jul", 2950, 1480, 620, 5050, 0.610], ["Aug", 3010, 1540, 700, 5250, 0.606],
              ["Sep", 3040, 1610, 585, 5235, 0.644], ["Q2 TOTAL", 9000, 4630, 1905, 15535, 0.620]]:
        ws.append(r)
    for row in range(4, ws.max_row + 1):
        ws.cell(row=row, column=6).number_format = "0.0%"
    cap = wb.create_sheet("Capital Spend")
    cap["A1"] = "Tarnwick Provisions - Capital spend FY2025-26, Q2"
    cap.merge_cells("A1:D1")
    cap["A2"] = "All figures in INR thousands"
    cap.append(["Month", "Plant & Machinery", "Vehicles", "Total"])
    for r in [["Jul", 1200, 0, 1200], ["Aug", 0, 850, 850], ["Sep", 640, None, 640],
              ["Q2 TOTAL", 1840, 850, 2690]]:
        cap.append(r)
    _save_wb(wb, d / "Expenses-FY2025-26.xlsx")


def build_operations() -> None:
    d = OUT / "operations"
    wb = Workbook()
    ws = wb.active
    ws.title = "Warehouse Throughput"
    ws.append(["Month", "Orders Shipped", "On-Time Delivery %", "Return Rate %", "Avg Pick Time (min)"])
    for r in [["Jul", 18420, 0.931, 0.018, 7.4], ["Aug", 19860, 0.946, 0.016, 7.1],
              ["Sep", 17935, 0.924, 0.022, 8.0]]:
        ws.append(r)
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=3).number_format = "0.0%"
        ws.cell(row=row, column=4).number_format = "0.0%"
        ws.cell(row=row, column=2).number_format = "#,##0"
    _save_wb(wb, d / "Warehouse-Throughput-Q2.xlsx")

    prs = Presentation()
    prs.core_properties.created = FIXED
    prs.core_properties.modified = FIXED
    prs.core_properties.author = "Tarnwick Operations (synthetic)"

    def slide(title: str, bullets: list[str]) -> None:
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = title
        tf = s.placeholders[1].text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b

    t = prs.slides.add_slide(prs.slide_layouts[0])
    t.shapes.title.text = "Tarnwick Provisions Q2 FY2025-26 Operations Review"
    t.placeholders[1].text = "Operations leadership, October 2025"
    slide("Revenue highlights", [
        "Q2 revenue closed at INR 25.07 million across all three product lines",
        "Juniper Masala crossed INR 1 million in a single month for the first time in August",
        "Marigold Chai remains the largest line; Saffron Rusk growth is flat",
    ])
    slide("Warehouse performance", [
        "On-time delivery in September was 92.9%, below the 95% target",
        "Return rate peaked in September after the monsoon packaging issue",
        "Peak daily dispatch reached 912 orders on 14 August",
    ])
    slide("Targets for Q3 FY2025-26", [
        "Q3 revenue target: INR 28.5 million",
        "Open two new distribution hubs, in Pune and Nagpur",
        "Reduce average pick time to 6.5 minutes",
    ])
    prs.save(d / "Ops-Review-Q2-FY2025-26.pptx")


def build_strategy() -> None:
    d = OUT / "strategy"
    d.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.core_properties.created = FIXED
    doc.core_properties.modified = FIXED
    doc.core_properties.author = "Tarnwick Finance (synthetic)"
    doc.add_heading("Board Memo: Q2 FY2025-26 Results", level=1)
    doc.add_paragraph("To: Board of Directors, Tarnwick Provisions Pvt Ltd")
    doc.add_paragraph("From: Office of the CFO")
    doc.add_paragraph("Date: 14 October 2025")
    doc.add_heading("Summary", level=2)
    doc.add_paragraph(
        "Revenue for Q2 FY2025-26 was INR 25.1 million, up 31.0% on Q2 of FY2024-25. "
        "August was the strongest month of the half-year."
    )
    doc.add_paragraph(
        "Operating costs for the quarter were INR 15.5 million, or 62.0% of revenue. "
        "Marigold Chai contributed 55.9% of quarterly revenue."
    )
    doc.add_heading("Outlook", level=2)
    doc.add_paragraph(
        "The Q3 FY2025-26 revenue target is INR 28.0 million. Management expects on-time "
        "delivery, which averaged 93.4% in Q2, to recover to the 95% target by December."
    )
    doc.add_paragraph("Capital spend in Q2 was concentrated on packaging machinery and two delivery vans.")
    doc.save(d / "Board-Memo-Q2-FY2025-26.docx")


def build_people() -> None:
    d = OUT / "people"
    wb = Workbook()
    ws = wb.active
    ws.title = "Team Directory"
    ws.append(["Employee", "Role", "Department", "Location", "Monthly Salary (INR)"])
    for r in [
        ["Tavish Orlund", "Head of Logistics", "Logistics", "Nagpur", 184500],
        ["Pemma Dasgupta", "Finance Controller", "Finance", "Pune", 172250],
        ["Ishaan Valdez", "Warehouse Supervisor", "Logistics", "Nagpur", 61750],
        ["Rhea Kantor", "Brand Manager", "Marketing", "Pune", 96300],
        ["Oduya Menon", "Procurement Lead", "Operations", "Pune", 88900],
    ]:
        ws.append(r)
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=5).number_format = "#,##0"
    _save_wb(wb, d / "Team-Directory-Sep-2025.xlsx")

    doc = Document()
    doc.core_properties.created = FIXED
    doc.core_properties.modified = FIXED
    doc.core_properties.author = "Tarnwick People Team (synthetic)"
    doc.add_heading("Compensation Policy Note", level=1)
    doc.add_paragraph("Tarnwick Provisions reviews salaries once a year, in April.")
    doc.add_paragraph("The annual performance bonus is 8% of annual base salary for grade B and above.")
    doc.add_paragraph("Employees in Nagpur receive a relocation allowance of INR 15000 in their first year.")
    doc.save(d / "Compensation-Policy-Note.docx")


def _freeze_zip_timestamps(root: Path) -> None:
    """Office files are zips; rewrite them with a fixed entry timestamp so the
    generated bytes are identical on every run."""
    for path in sorted(root.rglob("*")):
        if path.suffix not in {".xlsx", ".docx", ".pptx"}:
            continue
        with zipfile.ZipFile(path) as zin:
            items = [(info.filename, zin.read(info.filename)) for info in zin.infolist()]
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
            for name, data in items:
                if name == "docProps/core.xml":  # openpyxl stamps "modified" with the current time
                    data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2025-10-01T09:00:00Z", data)
                info = zipfile.ZipInfo(name, date_time=(2025, 10, 1, 9, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o644 << 16
                zout.writestr(info, data)


def main() -> None:
    build_finance()
    build_operations()
    build_strategy()
    build_people()
    _freeze_zip_timestamps(OUT)
    print(f"wrote synthetic corpus to {OUT}")


if __name__ == "__main__":
    main()
