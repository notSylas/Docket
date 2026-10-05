"""Deterministically generate the G1 "hard numeric" corpus and gold set.

Everything here is fictional: "Brindle Loom Textiles Pvt Ltd" does not exist
(and is unrelated to the Tarnwick Provisions corpus). Values are literal
numbers (no formulas) held in the data tables below; every derived gold answer
(sums, differences, percentage change, averages, unit conversions) is computed
from those tables in this file and never typed by hand. Run from anywhere:

    python eval-public/build_synthetic_corpus_g1.py

It rewrites eval-public/corpus-synthetic-g1/ and eval-public/gold-g1.yaml.
See Upgrade/08-evaluation-and-benchmarks.md section 5.2 (G1) and 5.4.

Structure is deliberately different from the Tarnwick corpus:
  * calendar-year (Jan-Dec) workbooks, quarter subtotal rows and a full-year
    TOTAL row (naive summing double counts);
  * units differ per workbook: CY2024 sales in full INR (units in header
    text), CY2025 sales in INR thousands (units only on a Notes sheet),
    costs in INR lakh, capex in INR crore (1 lakh = 100 thousand;
    1 crore = 100 lakh = 10,000 thousand), production in metres;
  * CY2025 actual and CY2025 budget workbooks have the same shape (actual vs
    budget in the same year);
  * blank vs literal 0: Silk Blend is blank before launch (CY2024 Jan-Jun,
    CY2025 Jan) and a literal 0 in CY2025 Feb; production has a blank March
    (no reading) and a literal 0 in loom hours for August (shutdown);
  * hidden rows: a bulk export memo (sales actual) and a trial-run memo
    (production), both "not in totals";
  * Costs-CY2025.xlsx carries a merged title row on its SECOND sheet only
    (Plant Capex): the parser labels every cell of that sheet with the title;
  * a value that exists only on a Notes sheet (inventory write-off);
  * the fiscal year is April-March but every workbook is calendar-year, so
    "first half of FY2025-26" is April to September, not Q1+Q2 subtotals.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import re
import zipfile
from pathlib import Path

import yaml
from docx import Document
from openpyxl import Workbook
from openpyxl.styles import Font

HERE = Path(__file__).resolve().parent
OUT = HERE / "corpus-synthetic-g1"
GOLD = HERE / "gold-g1.yaml"
FIXED = dt.datetime(2025, 12, 1, 9, 0, 0)
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
QUARTERS = ["Q1", "Q2", "Q3", "Q4"]
PRODUCTS = ["Cotton Yarn", "Linen Cloth", "Silk Blend"]

# --------------------------------------------------------------------------
# Data tables (the only hand-typed numbers). None = blank cell.
# --------------------------------------------------------------------------

SALES24 = {  # full INR
    "Cotton Yarn": [3120000, 3185000, 3290000, 3340000, 3415000, 3380000,
                    3520000, 3610000, 3575000, 3690000, 3815000, 4040000],
    "Linen Cloth": [1760000, 1795000, 1840000, 1880000, 1925000, 1910000,
                    1985000, 2040000, 2015000, 2090000, 2170000, 2310000],
    "Silk Blend": [None] * 6 + [410000, 520000, 610000, 655000, 720000, 805000],
}
SALES25 = {  # INR thousands, actual
    "Cotton Yarn": [3310, 3395, 3480, 3560, 3655, 3610, 3745, 3860, 3790, 3925, 4070, 4310],
    "Linen Cloth": [1855, 1890, 1950, 2005, 2070, 2040, 2120, 2190, 2145, 2230, 2315, 2470],
    "Silk Blend": [None, 0, 640, 712, 790, 845, 918, 1005, 972, 1088, 1160, 1290],
}
BUDGET25 = {  # INR thousands, budget
    "Cotton Yarn": [3300, 3400, 3500, 3600, 3700, 3700, 3800, 3900, 3900, 4000, 4100, 4300],
    "Linen Cloth": [1850, 1900, 1950, 2000, 2050, 2050, 2100, 2150, 2150, 2200, 2300, 2400],
    "Silk Blend": [0, 0, 600, 700, 800, 850, 900, 1000, 1000, 1050, 1150, 1250],
}
EXPORT_MEMO_25 = 1830  # hidden row in the actual workbook, not in totals

COST_LAKH = {  # per quarter, INR lakh
    "Raw Material": [74.2, 76.9, 79.5, 83.7],
    "Labour": [31.5, 32.0, 32.8, 33.6],
    "Energy": [12.8, 13.4, 14.1, 14.9],
}
CAPEX_CRORE = {  # per quarter, INR crore; None = blank
    "Plant & Machinery": [1.25, 0.8, 0, 2.1],
    "Vehicles": [None, 0, 0.45, 0.3],
}
WRITE_OFF_LAKH = 7.5  # appears only on the Notes sheet of Costs-CY2025.xlsx

OUTPUT_M = [41200, 42350, None, 43900, 44750, 44100, 45600, 0, 46800, 48250, 47900, 49300]
DEFECT_PCT = [2.4, 2.3, 2.5, 2.2, 2.1, 2.3, 2.0, 1.9, 2.1, 1.8, 1.9, 1.7]  # in percent
LOOM_HOURS = [5120, 5090, 5240, 5310, 5380, 5290, 5410, 0, 5520, 5640, 5580, 5700]
TRIAL_RUN_M = 3600  # hidden row in the production workbook, not in totals

MEMO_TARGET_CRORE = 17.5  # CY2026 revenue target stated only in the board memo


# --------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------


def s(values):  # sum ignoring blanks
    return sum(v or 0 for v in values)


def present(values):  # values that are recorded (0 counts, blank does not)
    return [v for v in values if v is not None]


def rnd(x: float, places: int = 4) -> float:
    return round(x + 0.0, places)


def fmt(x) -> str:
    """Compact decimal text: 7.5, 12, 0.35 (never 7.500000000001)."""
    x = round(x, 4)
    if x == int(x):
        return str(int(x))
    return f"{x:.4f}".rstrip("0")


class Table:
    """One sheet as indexed by the parser: header labels per column and rows
    keyed by their first cell. `label_override` models the merged-title sheet,
    where the parser labels every cell with the title text."""

    def __init__(self, source, file, sheet, cols, rows, label_override=None):
        self.source, self.file, self.sheet = source, file, sheet
        self.cols = cols  # header texts, col 0 is the row label
        self.rows = rows  # list of lists, same width as cols
        self.label_override = label_override

    def row(self, label):
        for r in self.rows:
            if r[0] == label:
                return r
        raise KeyError(label)

    def label(self, i):
        return self.label_override or self.cols[i]

    def col(self, name):
        for i, c in enumerate(self.cols):
            if c == name or c.startswith(name + " ("):
                return i
        raise KeyError(name)

    def val(self, row_label, col_name):
        return self.row(row_label)[self.col(col_name)]

    def span(self, row_label, *col_names, lo_limit=20, hi_limit=60):
        """A verbatim run of `Label: value` fragments (cells present in that
        row, in order) covering the named columns, 20-60 chars. Prefers a run
        that starts at the row label column."""
        r = self.row(row_label)
        frags = [(i, f"{self.label(i)}: {r[i]}") for i in range(len(r)) if r[i] is not None]
        want = {self.col(c) if isinstance(c, str) else c for c in col_names}
        best = None
        for a in range(len(frags)):
            for b in range(a, len(frags)):
                idx = {i for i, _ in frags[a:b + 1]}
                if not want <= idx:
                    continue
                text = " ".join(f for _, f in frags[a:b + 1])
                if not lo_limit <= len(text) <= hi_limit:
                    continue
                score = (0 if frags[a][0] == 0 else 1, len(text))
                if best is None or score < best[0]:
                    best = (score, text)
        if best is None:
            raise ValueError(f"no span fits for {self.file}/{self.sheet}/{row_label}/{col_names}")
        return best[1]


def _sales_table(source, file, sheet, data, hidden_memo, suffix):
    cols = ["Month"] + [p + suffix for p in PRODUCTS] + ["Total" + suffix, "Silk Share %"]
    rows = []
    sums = [[0, 0, 0] for _ in QUARTERS]
    seen = [[False, False, False] for _ in QUARTERS]
    for m in range(12):
        vals = [data[p][m] for p in PRODUCTS]
        total = s(vals)
        share = None if vals[2] is None else round(vals[2] / total, 3)
        rows.append([MONTHS[m], *vals, total, share])
        q = m // 3
        for j, v in enumerate(vals):
            if v is not None:
                sums[q][j] += v
                seen[q][j] = True
        if m % 3 == 2:
            rows.append([f"{QUARTERS[q]} subtotal",
                         *[sums[q][j] if seen[q][j] else None for j in range(3)], sum(sums[q]), None])
    grand = [sum(sums[q][j] for q in range(4)) for j in range(3)]
    any_seen = [any(seen[q][j] for q in range(4)) for j in range(3)]
    rows.append(["TOTAL (Jan-Dec)", *[grand[j] if any_seen[j] else None for j in range(3)], sum(grand), None])
    if hidden_memo is not None:
        rows.append(["Memo: bulk export order (not in totals)", None, None, None, hidden_memo, None])
    return Table(source, file, sheet, cols, rows)


def _save_wb(wb: Workbook, path: Path) -> None:
    wb.properties.creator = "Brindle Loom Finance (synthetic)"
    wb.properties.created = FIXED
    wb.properties.modified = FIXED
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _write_sheet(ws, table: Table, hidden_labels=(), pct_cols=()):
    ws.append(table.cols)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in table.rows:
        ws.append(r)
        if r[0] in hidden_labels:
            ws.row_dimensions[ws.max_row].hidden = True
    for row in range(2, ws.max_row + 1):
        for col in pct_cols:
            ws.cell(row=row, column=col).number_format = "0.0%"


def _notes_sheet(wb, rows):
    n = wb.create_sheet("Notes")
    n.append(["Note", "Detail"])
    for r in rows:
        n.append(list(r))


# --------------------------------------------------------------------------
# Tables
# --------------------------------------------------------------------------

T_S24 = _sales_table("finance", "Sales-CY2024.xlsx", "Sales by Month", SALES24, None, " (INR)")
T_S25 = _sales_table("finance", "Sales-CY2025-Actual.xlsx", "Sales Actual", SALES25, EXPORT_MEMO_25, "")
T_B25 = _sales_table("finance", "Sales-CY2025-Budget.xlsx", "Sales Budget", BUDGET25, None, "")


def _cost_table():
    cols = ["Quarter", "Raw Material (INR lakh)", "Labour (INR lakh)", "Energy (INR lakh)", "Total (INR lakh)"]
    rows = []
    for q in range(4):
        parts = [COST_LAKH[k][q] for k in COST_LAKH]
        rows.append([f"{QUARTERS[q]} CY2025", *parts, rnd(sum(parts), 1)])
    yr = [rnd(sum(COST_LAKH[k]), 1) for k in COST_LAKH]
    rows.append(["CY2025 TOTAL", *yr, rnd(sum(yr), 1)])
    return Table("finance", "Costs-CY2025.xlsx", "Cost Summary", cols, rows)


CAPEX_TITLE = "Plant Capex CY2025"


def _capex_table():
    cols = ["Quarter", "Plant & Machinery", "Vehicles", "Total"]
    rows = []
    for q in range(4):
        pm, ve = CAPEX_CRORE["Plant & Machinery"][q], CAPEX_CRORE["Vehicles"][q]
        rows.append([f"{QUARTERS[q]} CY2025", pm, ve, rnd(s([pm, ve]), 2)])
    pm_t = rnd(s(CAPEX_CRORE["Plant & Machinery"]), 2)
    ve_t = rnd(s(CAPEX_CRORE["Vehicles"]), 2)
    rows.append(["CY2025 TOTAL", pm_t, ve_t, rnd(pm_t + ve_t, 2)])
    return Table("finance", "Costs-CY2025.xlsx", "Plant Capex", cols, rows, label_override=CAPEX_TITLE)


T_COST = _cost_table()
T_CAPEX = _capex_table()


def _production_table():
    cols = ["Month", "Output (metres)", "Defect Rate (%)", "Loom Hours"]
    rows = [[MONTHS[m], OUTPUT_M[m], DEFECT_PCT[m], LOOM_HOURS[m]] for m in range(12)]
    rows.append(["TOTAL (Jan-Dec)", s(OUTPUT_M), None, s(LOOM_HOURS)])
    rows.append(["Memo: trial run (not in totals)", TRIAL_RUN_M, None, None])
    return Table("operations", "Production-CY2025.xlsx", "Loom Output", cols, rows)


T_PROD = _production_table()


def build_finance() -> None:
    d = OUT / "finance"
    wb = Workbook()
    _write_sheet(wb.active, T_S24, pct_cols=(6,))
    wb.active.title = T_S24.sheet
    for r in range(2, wb.active.max_row + 1):
        for c in range(2, 6):
            wb.active.cell(row=r, column=c).number_format = "#,##0"
    _save_wb(wb, d / T_S24.file)

    for t, label in ((T_S25, "actuals"), (T_B25, "budget")):
        wb = Workbook()
        wb.active.title = t.sheet
        _write_sheet(wb.active, t, hidden_labels=("Memo: bulk export order (not in totals)",), pct_cols=(6,))
        _notes_sheet(wb, [
            ("Units", f"INR thousands ({label})"),
            ("Period", "Calendar year, January to December 2025"),
            ("Fiscal year", "April to March"),
        ])
        _save_wb(wb, d / t.file)

    wb = Workbook()
    ws = wb.active
    ws.title = T_COST.sheet
    _write_sheet(ws, T_COST)
    cap = wb.create_sheet(T_CAPEX.sheet)
    cap["A1"] = CAPEX_TITLE
    cap.merge_cells("A1:D1")
    cap["A2"] = "All figures in INR crore"
    cap.append(T_CAPEX.cols)
    for r in T_CAPEX.rows:
        cap.append(r)
    _notes_sheet(wb, [
        ("Units", "Cost Summary is in INR lakh; Plant Capex is in INR crore (1 crore = 100 lakh)"),
        ("Inventory write-off", f"INR {fmt(WRITE_OFF_LAKH)} lakh booked in Q3 CY2025, not in Cost Summary"),
        ("Fiscal year", "April to March"),
    ])
    _save_wb(wb, d / T_COST.file)


def build_operations() -> None:
    d = OUT / "operations"
    wb = Workbook()
    ws = wb.active
    ws.title = T_PROD.sheet
    _write_sheet(ws, T_PROD, hidden_labels=("Memo: trial run (not in totals)",))
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=2).number_format = "#,##0"
    _notes_sheet(wb, [
        ("Units", "Output in metres of finished cloth; defect rate in percent (2.4 means 2.4%)"),
        ("March", "No output reading was recorded for March (meter fault); the cell is left blank"),
    ])
    _save_wb(wb, d / T_PROD.file)


MEMO_PARAS = {}


def _total(t):
    return t.val("TOTAL (Jan-Dec)", "Total")


def build_strategy() -> None:
    d = OUT / "strategy"
    d.mkdir(parents=True, exist_ok=True)
    rev25_k = _total(T_S25)
    rev24_k = _total(T_S24) / 1000
    growth = (rev25_k - rev24_k) / rev24_k * 100
    cost_lakh = T_COST.val("CY2025 TOTAL", "Total")
    cost_ratio = cost_lakh * 100 / rev25_k * 100  # lakh -> thousands, then percent
    MEMO_PARAS["revenue"] = (
        f"Revenue for CY2025 was INR {rev25_k / 10000:.2f} crore, up {growth:.1f}% on CY2024. "
        "Silk Blend was the fastest growing line."
    )
    MEMO_PARAS["cost"] = (
        f"Operating costs for CY2025 were INR {cost_lakh / 100:.2f} crore, or {cost_ratio:.1f}% of revenue."
    )
    MEMO_PARAS["target"] = f"The CY2026 revenue target is INR {fmt(MEMO_TARGET_CRORE)} crore."
    doc = Document()
    doc.core_properties.created = FIXED
    doc.core_properties.modified = FIXED
    doc.core_properties.author = "Brindle Loom Finance (synthetic)"
    doc.add_heading("Board Memo: CY2025 Annual Review", level=1)
    doc.add_paragraph("To: Board of Directors, Brindle Loom Textiles Pvt Ltd")
    doc.add_paragraph("From: Office of the CFO")
    doc.add_paragraph("Date: 20 January 2026")
    doc.add_heading("Summary", level=2)
    doc.add_paragraph(MEMO_PARAS["revenue"])
    doc.add_paragraph(MEMO_PARAS["cost"])
    doc.add_heading("Outlook", level=2)
    doc.add_paragraph(MEMO_PARAS["target"])
    doc.add_paragraph("Management plans to add a second silk-blend line in the second half of the year.")
    doc.save(d / "Annual-Review-Memo-CY2025.docx")


def _freeze_zip_timestamps(root: Path) -> None:
    for path in sorted(root.rglob("*")):
        if path.suffix not in {".xlsx", ".docx", ".pptx"}:
            continue
        with zipfile.ZipFile(path) as zin:
            items = [(info.filename, zin.read(info.filename)) for info in zin.infolist()]
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
            for name, data in items:
                if name == "docProps/core.xml":
                    data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2025-12-01T09:00:00Z", data)
                info = zipfile.ZipInfo(name, date_time=(2025, 12, 1, 9, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.external_attr = 0o644 << 16
                zout.writestr(info, data)


# --------------------------------------------------------------------------
# Gold set (all derived answers computed here from the tables above)
# --------------------------------------------------------------------------

NOT_FOUND = (r"re:blank|empty|no (silk blend |silk )?(revenue|figure|value|entry|data|amount|sales)|"
             r"not (recorded|reported|listed|available|launched|yet)|nothing")
ZERO = r"re:(\b0\b|zero|nil)"


def num_re(value, places=None) -> str:
    """Regex for a number as it may appear in an (already normalized) answer:
    thousands separators are stripped by normalization, trailing zeros are
    optional, and the match cannot sit inside a longer number."""
    text = fmt(value) if places is None else f"{value:.{places}f}".rstrip("0").rstrip(".")
    esc = re.escape(text)
    tail = r"(?:\.0+)?" if "." not in text else r"0*"
    return rf"(?<![\d.]){esc}{tail}(?!\d|\.\d)"


def pat(*values, places=None) -> str:
    return "re:" + "|".join(num_re(v, places) for v in values)


def ppat(p: float) -> str:
    """Percentage, matched at 1 or 2 decimals; sign ignored (the question names the direction)."""
    p = abs(p)
    return "re:" + num_re(rnd(p, 1), 1) + "|" + num_re(rnd(p, 2), 2)


FULL_MONTH = {"Jan": "january", "Feb": "february", "Mar": "march", "Apr": "april", "May": "may", "Jun": "june",
              "Jul": "july", "Aug": "august", "Sep": "september", "Oct": "october", "Nov": "november", "Dec": "december"}

QUESTIONS: list[dict] = []


def q(qid, qtype, question, *, must=(), not_=(), spans=(), docs=(), origin, answerable=True):
    QUESTIONS.append({
        "id": qid, "type": qtype, "question": question, "answerable": answerable,
        "must_contain": list(must), "must_not_contain": list(not_),
        "gold_spans": list(spans), "reviewed": True,
        "origin": f"synthetic corpus (eval-public/build_synthetic_corpus_g1.py); {origin}",
        "source_documents": list(docs), "formula_dependent": False,
    })


def gold_questions() -> None:
    S24, S25, B25, C, CX, P = T_S24, T_S25, T_B25, T_COST, T_CAPEX, T_PROD
    f24, f25, fb, fc, fp = S24.file, S25.file, B25.file, C.file, P.file
    memo = "Annual-Review-Memo-CY2025.docx"
    months = {m: i for i, m in enumerate(MONTHS)}

    def mv(table, product, month):
        return table.val(month, product)

    def month_sum(table, col, first, last):
        a, b = months[first], months[last]
        return s(table.val(m, col) for m in MONTHS[a:b + 1])

    def mspans(table, col, first, last):
        a, b = months[first], months[last]
        return [table.span(m, col) for m in MONTHS[a:b + 1]]

    # ---- literal lookups (trap coverage: units, actual vs budget, blank vs zero, hidden)
    v = mv(S25, "Cotton Yarn", "Jul")
    q("g1-cotton-jul25-actual", "numeric", "What was the actual Cotton Yarn revenue in July 2025?",
      must=[pat(v, v * 1000)], not_=[pat(mv(B25, "Cotton Yarn", "Jul"))],
      spans=[S25.span("Jul", "Cotton Yarn")], docs=[f25],
      origin=f"literal cell, INR thousands (Notes sheet); budget workbook has {mv(B25, 'Cotton Yarn', 'Jul')} for the same cell")

    v = mv(B25, "Linen Cloth", "Sep")
    q("g1-linen-sep25-budget", "numeric", "What was the budgeted Linen Cloth revenue for September 2025?",
      must=[pat(v, v * 1000)], not_=[pat(mv(S25, "Linen Cloth", "Sep"))],
      spans=[B25.span("Sep", "Linen Cloth")], docs=[fb],
      origin=f"literal cell in the budget workbook; actual is {mv(S25, 'Linen Cloth', 'Sep')}")

    assert mv(S25, "Silk Blend", "Feb") == 0
    q("g1-silk-feb25-zero", "numeric", "How much Silk Blend revenue was recorded in February 2025 (actual)?",
      must=[ZERO], spans=[S25.span("Feb", "Silk Blend")], docs=[f25],
      origin="blank-vs-zero: February cell is a literal 0 (failed trial lot), January cell is blank")

    assert mv(S25, "Silk Blend", "Jan") is None
    q("g1-silk-jan25-blank", "numeric", "What was the actual Silk Blend revenue in January 2025?",
      must=[NOT_FOUND], not_=[r"re:(?<![\d.])0(?![\d.])\s*(inr|thousand)"],
      spans=[S25.span("Jan", "Cotton Yarn", "Linen Cloth", "Total")], docs=[f25],
      origin="blank-vs-zero: the January Silk Blend cell is blank (not launched), not 0")

    v = S24.val("Mar", "Total")
    q("g1-total-mar24", "numeric", "What was total revenue in March 2024?",
      must=[pat(v, v / 1000)], not_=[pat(S25.val("Mar", "Total"))],
      spans=[S24.span("Mar", "Total")], docs=[f24],
      origin="CY2024 workbook is in full INR, not thousands (units in header text)")

    v = S25.val("Q3 subtotal", "Total")
    q("g1-q3-2025-total", "numeric", "What was total actual revenue for the third quarter of 2025 (July to September)?",
      must=[pat(v, v * 1000)], not_=[pat(B25.val("Q3 subtotal", "Total"))],
      spans=[S25.span("Q3 subtotal", "Total")], docs=[f25],
      origin=f"literal subtotal row, INR thousands; budget Q3 subtotal is {B25.val('Q3 subtotal', 'Total')}")

    q("g1-hidden-export-memo", "numeric", "What was the one-off bulk export order noted in the CY2025 actual sales workbook?",
      must=[pat(EXPORT_MEMO_25, EXPORT_MEMO_25 * 1000)],
      spans=[S25.span("Memo: bulk export order (not in totals)", "Total", lo_limit=20)],
      docs=[f25], origin="hidden row, INR thousands, marked as not included in totals")

    # ---- sums across rows (subtotal traps)
    v = month_sum(S25, "Cotton Yarn", "Mar", "May")
    q("g1-cotton-mar-may25-sum", "numeric", "What was the combined actual Cotton Yarn revenue for March, April and May 2025?",
      must=[pat(v, v * 1000)], spans=mspans(S25, "Cotton Yarn", "Mar", "May"), docs=[f25],
      origin="sum of three monthly cells that straddle the Q1/Q2 subtotal row (a subtotal row sits between Mar and Apr)")

    v = month_sum(S25, "Linen Cloth", "Jan", "Jun")
    wrong = v * 2
    assert v == S25.val("Q1 subtotal", "Linen Cloth") + S25.val("Q2 subtotal", "Linen Cloth")
    q("g1-linen-h1-2025-sum", "numeric", "What was total actual Linen Cloth revenue from January to June 2025?",
      must=[pat(v, v * 1000)], not_=[pat(wrong)],
      spans=[S25.span("Q1 subtotal", "Linen Cloth"), S25.span("Q2 subtotal", "Linen Cloth")], docs=[f25],
      origin=f"= Q1 + Q2 subtotals; adding months and subtotals double counts to {wrong}")

    v = _total(S25)
    both = v + EXPORT_MEMO_25
    q("g1-total-cy2025-excl-memo", "numeric",
      "What was total actual revenue for calendar year 2025, excluding the one-off export order?",
      must=[pat(v, v * 1000)], not_=[pat(both)],
      spans=[S25.span("TOTAL (Jan-Dec)", "Total")], docs=[f25],
      origin=f"literal TOTAL row; the hidden export memo ({EXPORT_MEMO_25}) is not in totals; including it would give {both}")
    q("g1-total-cy2025-incl-memo", "numeric",
      "What was total actual revenue for calendar year 2025 including the one-off export order?",
      must=[pat(both, both * 1000)], not_=[pat(v)] if both != v else [],
      spans=[S25.span("TOTAL (Jan-Dec)", "Total"), S25.span("Memo: bulk export order (not in totals)", "Total")],
      docs=[f25], origin="derived: TOTAL row plus the hidden export memo row")

    # ---- actual vs budget in the same year
    d = S25.val("Jul", "Total") - B25.val("Jul", "Total")
    q("g1-jul25-variance", "numeric", "By how much did actual total revenue differ from budget in July 2025 (give the size of the gap), in INR thousands?",
      must=[pat(abs(d), abs(d) * 1000)], spans=[S25.span("Jul", "Total"), B25.span("Jul", "Total")], docs=[f25, fb],
      origin=f"derived: {S25.val('Jul','Total')} actual minus {B25.val('Jul','Total')} budget")

    d = S25.val("Q2 subtotal", "Total") - B25.val("Q2 subtotal", "Total")
    q("g1-q2-2025-variance", "numeric", "By how much did actual Q2 2025 total revenue differ from the Q2 budget (give the size of the gap), in INR thousands?",
      must=[pat(abs(d), abs(d) * 1000)], spans=[S25.span("Q2 subtotal", "Total"), B25.span("Q2 subtotal", "Total")],
      docs=[f25, fb], origin=f"derived: {S25.val('Q2 subtotal','Total')} minus {B25.val('Q2 subtotal','Total')}")

    d = S25.val("TOTAL (Jan-Dec)", "Cotton Yarn") - B25.val("TOTAL (Jan-Dec)", "Cotton Yarn")
    q("g1-cotton-fy-variance", "numeric",
      "What was the size of the full-year 2025 gap between actual and budgeted Cotton Yarn revenue, in INR thousands?",
      must=[pat(abs(d), abs(d) * 1000)], spans=[S25.span("TOTAL (Jan-Dec)", "Cotton Yarn"), B25.span("TOTAL (Jan-Dec)", "Cotton Yarn")],
      docs=[f25, fb], origin=f"derived: {S25.val('TOTAL (Jan-Dec)','Cotton Yarn')} actual vs {B25.val('TOTAL (Jan-Dec)','Cotton Yarn')} budget")

    a, b = _total(S25), _total(B25)
    p = (a - b) / b * 100
    q("g1-fy-vs-budget-pct", "numeric", "By what percentage did full-year 2025 actual revenue differ from the full-year budget (give the size of the gap)?",
      must=[ppat(p)],
      spans=[S25.span("TOTAL (Jan-Dec)", "Total"), B25.span("TOTAL (Jan-Dec)", "Total")], docs=[f25, fb],
      origin=f"derived: ({a} - {b}) / {b} = {p:.3f}%")

    # ---- percentage change
    a, b = mv(S25, "Cotton Yarn", "Jan"), mv(S25, "Cotton Yarn", "Dec")
    p = (b - a) / a * 100
    q("g1-cotton-jan-dec-pct", "numeric", "By what percentage did actual Cotton Yarn revenue grow from January to December 2025?",
      must=[ppat(p)], spans=[S25.span("Jan", "Cotton Yarn"), S25.span("Dec", "Cotton Yarn")], docs=[f25],
      origin=f"derived: ({b} - {a}) / {a} = {p:.3f}%")

    a, b = S25.val("Jun", "Total"), S25.val("Jul", "Total")
    p = (b - a) / a * 100
    q("g1-total-jun-jul-pct", "numeric", "What was the month-on-month percentage growth in actual total revenue from June to July 2025?",
      must=[ppat(p)], spans=[S25.span("Jun", "Total"), S25.span("Jul", "Total")], docs=[f25],
      origin=f"derived: ({b} - {a}) / {a} = {p:.3f}%")

    a, b = S25.val("Dec", "Silk Blend"), S25.val("Aug", "Silk Blend")
    p = (a - b) / b * 100
    q("g1-silk-aug-dec-pct", "numeric", "By what percentage did actual Silk Blend revenue rise from August to December 2025?",
      must=[ppat(p)], spans=[S25.span("Aug", "Silk Blend"), S25.span("Dec", "Silk Blend")], docs=[f25],
      origin=f"derived: ({a} - {b}) / {b} = {p:.3f}%")

    # ---- averages with blanks vs zeros
    vals = SALES25["Silk Blend"]
    rec = present(vals)
    avg = s(rec) / len(rec)
    not_ = {rnd(s(rec) / 12, 1), rnd(s(rec) / len([x for x in rec if x]), 1)}
    not_.discard(rnd(avg, 1))
    q("g1-silk-avg-2025", "numeric",
      "What was the average monthly Silk Blend revenue in 2025 (actual), counting a blank cell as not recorded and a 0 as a real month of zero revenue?",
      must=[pat(rnd(avg, 1), places=1) + "|" + num_re(rnd(avg, 2), 2)],
      not_=[pat(x, places=1) for x in sorted(not_)],
      spans=[S25.span("TOTAL (Jan-Dec)", "Silk Blend"), S25.span("Feb", "Silk Blend"), S25.span("Jan", "Cotton Yarn", "Total")],
      docs=[f25],
      origin=f"derived: sum {s(rec)} over {len(rec)} recorded months (Jan blank excluded, Feb 0 included) = {avg:.3f}; /12 or excluding the zero month gives different values")

    vals = SALES24["Silk Blend"]
    rec = present(vals)
    avg = s(rec) / len(rec)
    wrong = rnd(s(rec) / 12, 1)
    q("g1-silk-avg-2024", "numeric",
      "What was the average monthly Silk Blend revenue in 2024, over the months in which the product had a recorded figure, in INR?",
      must=[pat(rnd(avg, 1), places=1) + "|" + num_re(rnd(avg, 2), 2)], not_=[pat(wrong, places=1)],
      spans=[S24.span("TOTAL (Jan-Dec)", "Silk Blend"), S24.span("Jul", "Silk Blend")], docs=[f24],
      origin=f"derived: {s(rec)} over {len(rec)} months (Jan-Jun blank) = {avg:.1f}; dividing by 12 gives {wrong}")

    avg = _total(S24) / 12
    q("g1-total-avg-2024", "numeric", "What was the average monthly total revenue in 2024, in INR?",
      must=[pat(rnd(avg, 1), places=1) + "|" + num_re(rnd(avg, 2), 2) + "|" + num_re(round(avg))],
      spans=[S24.span("TOTAL (Jan-Dec)", "Total")], docs=[f24],
      origin=f"derived: {_total(S24)} / 12 = {avg:.2f}; full INR units")

    # ---- calendar vs fiscal year
    v = month_sum(S25, "Total", "Apr", "Sep")
    h1 = S25.val("Q1 subtotal", "Total") + S25.val("Q2 subtotal", "Total")
    q("g1-fy2526-h1-actual", "numeric",
      "The fiscal year runs April to March. What was total actual revenue for the first half of FY2025-26 (April to September 2025), in INR thousands?",
      must=[pat(v, v * 1000)], not_=[pat(h1)], spans=mspans(S25, "Total", "Apr", "Sep"), docs=[f25],
      origin=f"calendar vs fiscal trap: Apr-Sep sum, not Q1+Q2 subtotals ({h1}); fiscal year note on the Notes sheet")

    v = S25.val("Q1 subtotal", "Total")
    q("g1-fy2425-q4-revenue", "numeric",
      "The fiscal year runs April to March. What was total actual revenue for the last quarter of FY2024-25 (January to March 2025), in INR thousands?",
      must=[pat(v, v * 1000)], not_=[pat(S25.val("Q4 subtotal", "Total"))],
      spans=[S25.span("Q1 subtotal", "Total")], docs=[f25],
      origin="calendar vs fiscal trap: the calendar Q1 subtotal row is the fiscal Q4 of FY2024-25")

    v = month_sum(S25, "Total", "Apr", "Dec")
    q("g1-fy2526-ytd-dec", "numeric",
      "The fiscal year runs April to March. What was total actual revenue from the start of FY2025-26 through December 2025, in INR thousands?",
      must=[pat(v, v * 1000)], not_=[pat(_total(S25))], spans=mspans(S25, "Total", "Apr", "Dec"), docs=[f25],
      origin=f"derived: Apr-Dec monthly sum {v}; the calendar-year TOTAL is {_total(S25)}")

    # ---- mixed units across workbooks
    a_k, b_inr = _total(S25), _total(S24)
    inc_inr = a_k * 1000 - b_inr
    q("g1-yoy-increase-inr", "multi_doc",
      "By how many rupees did total revenue in 2025 (actual) exceed total revenue in 2024?",
      must=[pat(inc_inr, inc_inr / 1000, inc_inr / 100000)],
      spans=[S25.span("TOTAL (Jan-Dec)", "Total"), S24.span("TOTAL (Jan-Dec)", "Total")], docs=[f25, f24],
      origin=f"derived: {a_k} thousand = {a_k*1000} INR minus {b_inr} INR (units differ between workbooks)")

    p = (a_k * 1000 - b_inr) / b_inr * 100
    q("g1-yoy-pct", "multi_doc", "By what percentage did total revenue grow from 2024 to 2025 (actual)?",
      must=[ppat(p)],
      spans=[S25.span("TOTAL (Jan-Dec)", "Total"), S24.span("TOTAL (Jan-Dec)", "Total")], docs=[f25, f24],
      origin=f"derived: CY2025 {a_k} thousand vs CY2024 {b_inr} INR = {p:.3f}%; ignoring units gives a nonsense result")

    d = S25.val("Dec", "Silk Blend") * 1000 - S24.val("Dec", "Silk Blend")
    q("g1-silk-dec-yoy-inr", "multi_doc", "By how many rupees did Silk Blend revenue in December 2025 (actual) exceed December 2024?",
      must=[pat(d, d / 1000)], spans=[S25.span("Dec", "Silk Blend"), S24.span("Dec", "Silk Blend")], docs=[f25, f24],
      origin=f"derived: {S25.val('Dec','Silk Blend')} thousand minus {S24.val('Dec','Silk Blend')} INR, in INR")

    # ---- costs (INR lakh), Notes-only value, capex (INR crore, merged title)
    v = C.val("Q3 CY2025", "Labour")
    q("g1-labour-q3", "numeric", "What was Labour cost in the third quarter of CY2025?",
      must=[pat(v)], spans=[C.span("Q3 CY2025", "Labour")], docs=[fc],
      origin="literal cell, INR lakh")

    q("g1-writeoff-notes-only", "numeric", "What inventory write-off was booked in 2025, and in which quarter?",
      must=[pat(WRITE_OFF_LAKH), r"re:q3|third quarter|quarter 3"],
      spans=[f"Detail: INR {fmt(WRITE_OFF_LAKH)} lakh booked in Q3 CY2025"],
      docs=[fc], origin="value that exists only on the Notes sheet")

    v = rnd(C.val("Q3 CY2025", "Total") + WRITE_OFF_LAKH, 2)
    q("g1-q3-cost-with-writeoff", "multi_doc",
      "What were total costs for Q3 CY2025 including the inventory write-off, in INR lakh?",
      must=[pat(v)], not_=[pat(C.val("Q3 CY2025", "Total"))],
      spans=[C.span("Q3 CY2025", "Total"), f"Detail: INR {fmt(WRITE_OFF_LAKH)} lakh booked in Q3 CY2025"],
      docs=[fc], origin=f"derived: Cost Summary Q3 total {C.val('Q3 CY2025','Total')} plus Notes-only write-off {WRITE_OFF_LAKH}")

    v = rnd(C.val("Q1 CY2025", "Energy") + C.val("Q2 CY2025", "Energy"), 2)
    q("g1-energy-h1-sum", "numeric", "What was the combined Energy cost for Q1 and Q2 of CY2025, in INR lakh?",
      must=[pat(v)], spans=[C.span("Q1 CY2025", "Energy"), C.span("Q2 CY2025", "Energy")], docs=[fc],
      origin="derived: sum of two quarterly cells, INR lakh")

    assert CX.val("Q2 CY2025", "Vehicles") == 0 and CX.val("Q1 CY2025", "Vehicles") is None
    q("g1-capex-q2-vehicles-zero", "numeric", "How much was spent on Vehicles capex in Q2 CY2025?",
      must=[ZERO], spans=[CX.span("Q2 CY2025", "Vehicles")], docs=[fc],
      origin="blank-vs-zero on the merged-title Plant Capex sheet: Q2 Vehicles is a literal 0, Q1 Vehicles is blank; INR crore")

    v = CX.val("CY2025 TOTAL", "Total")
    q("g1-capex-total-crore", "numeric", "What was total capex for CY2025, in INR crore?",
      must=[pat(v)], spans=[CX.span("CY2025 TOTAL", "Total")], docs=[fc],
      origin="literal total on the merged-title Plant Capex sheet; unit is INR crore (units row), not lakh")

    v = rnd(CX.val("Q4 CY2025", "Total") * 100, 2)
    q("g1-capex-q4-lakh", "numeric", "What was total capex in Q4 CY2025, expressed in INR lakh?",
      must=[pat(v)], spans=[CX.span("Q4 CY2025", "Total")], docs=[fc],
      origin=f"derived unit conversion: {CX.val('Q4 CY2025','Total')} crore x 100 = {v} lakh")

    rev = S25.val("Q1 subtotal", "Total")
    cost = C.val("Q1 CY2025", "Total")
    d = rnd(rev - cost * 100, 2)
    q("g1-q1-margin-thousands", "multi_doc",
      "What was Q1 2025 actual revenue minus Q1 CY2025 total costs, in INR thousands?",
      must=[pat(d, d * 1000)], not_=[pat(rev - cost)],
      spans=[S25.span("Q1 subtotal", "Total"), C.span("Q1 CY2025", "Total")], docs=[f25, fc],
      origin=f"derived: {rev} thousand minus {cost} lakh x 100 = {d}; revenue is in thousands, costs in lakh")

    pct = cost * 100 / rev * 100
    q("g1-q1-cost-ratio", "multi_doc", "What percentage of Q1 2025 actual revenue was consumed by Q1 CY2025 total costs?",
      must=[ppat(pct)],
      spans=[S25.span("Q1 subtotal", "Total"), C.span("Q1 CY2025", "Total")], docs=[f25, fc],
      origin=f"derived: {cost} lakh = {cost*100} thousand over {rev} thousand = {pct:.3f}%")

    outlay = rnd(C.val("CY2025 TOTAL", "Total") + CX.val("CY2025 TOTAL", "Total") * 100, 2)
    q("g1-outlay-lakh", "multi_doc", "What was total CY2025 operating cost plus total CY2025 capex, in INR lakh?",
      must=[pat(outlay)],
      spans=[C.span("CY2025 TOTAL", "Total"), CX.span("CY2025 TOTAL", "Total")], docs=[fc],
      origin=f"derived: {C.val('CY2025 TOTAL','Total')} lakh plus {CX.val('CY2025 TOTAL','Total')} crore x 100 (both on separate sheets of one workbook)")

    # ---- production (blank March, zero August, hidden trial run)
    rec = present(OUTPUT_M[:6])
    avg = s(rec) / len(rec)
    wrong = s(rec) / 6
    q("g1-output-avg-h1", "numeric",
      "What was the average monthly output in metres from January to June 2025, counting a blank month as not recorded?",
      must=[pat(rnd(avg, 1), places=1) + "|" + num_re(rnd(avg, 2), 2)], not_=[pat(rnd(wrong, 1), places=1)],
      spans=[P.span(m, "Output") for m in MONTHS[:6] if P.val(m, "Output") is not None], docs=[fp],
      origin=f"derived: {s(rec)} over {len(rec)} recorded months (March blank) = {avg:.2f}; /6 gives {wrong:.2f}")

    v = s(OUTPUT_M)
    q("g1-output-total-excl-trial", "numeric", "What was total output in metres for 2025, excluding the trial run?",
      must=[pat(v)], not_=[pat(v + TRIAL_RUN_M)], spans=[P.span("TOTAL (Jan-Dec)", "Output")], docs=[fp],
      origin=f"literal TOTAL row; hidden trial-run memo ({TRIAL_RUN_M}) is not in totals")

    best = max(range(12), key=lambda i: OUTPUT_M[i] or -1)
    q("g1-output-peak-month", "numeric", "Which month of 2025 had the highest output in metres, and what was it?",
      must=[pat(OUTPUT_M[best]), rf"re:\b({MONTHS[best].casefold()}|{FULL_MONTH[MONTHS[best]]})\b"],
      spans=[P.span(MONTHS[best], "Output")], docs=[fp],
      origin=f"derived: argmax over recorded months = {MONTHS[best]} ({OUTPUT_M[best]})")

    avg = sum(DEFECT_PCT) / 12
    q("g1-defect-avg", "numeric", "What was the average monthly defect rate across 2025, in percent?",
      must=[ppat(avg)],
      spans=[P.span(m, "Defect Rate") for m in ("Jan", "Dec")], docs=[fp],
      origin=f"derived: mean of 12 monthly defect-rate cells = {avg:.4f}% (cells are in percent, not fractions)")

    a, b = OUTPUT_M[0], OUTPUT_M[11]
    p = (b - a) / a * 100
    q("g1-output-jan-dec-pct", "numeric", "By what percentage did monthly output change from January to December 2025?",
      must=[ppat(p)], spans=[P.span("Jan", "Output"), P.span("Dec", "Output")], docs=[fp],
      origin=f"derived: ({b} - {a}) / {a} = {p:.3f}%")

    assert LOOM_HOURS[7] == 0
    q("g1-loomhours-aug-zero", "numeric", "How many loom hours were recorded in August 2025?",
      must=[ZERO], spans=[P.span("Aug", "Loom Hours")], docs=[fp],
      origin="blank-vs-zero: August is a literal 0 (shutdown) in both Output and Loom Hours; March output is blank")

    # ---- memo (rounded, different units) and cross-document
    q("g1-memo-target", "numeric", "What revenue target does the board memo set for CY2026?",
      must=[pat(MEMO_TARGET_CRORE, MEMO_TARGET_CRORE * 10, MEMO_TARGET_CRORE * 10000)],
      spans=[MEMO_PARAS["target"]], docs=[memo], origin="literal sentence in the memo, INR crore")

    g = re.search(r"up ([\d.]+)%", MEMO_PARAS["revenue"]).group(1)
    q("g1-memo-growth", "numeric", "By what percentage does the board memo say CY2025 revenue grew over CY2024?",
      must=[pat(float(g), places=1)], spans=[MEMO_PARAS["revenue"][:60].rstrip()], docs=[memo],
      origin="literal in memo; computed by this generator from the two sales workbooks, must agree with g1-yoy-pct")

    tgt_k = MEMO_TARGET_CRORE * 10000
    p = (tgt_k - _total(S25)) / _total(S25) * 100
    q("g1-target-vs-2025-pct", "multi_doc",
      "By what percentage does the memo's CY2026 revenue target exceed actual CY2025 revenue from the sales workbook?",
      must=[ppat(p)], not_=[],
      spans=[MEMO_PARAS["target"], S25.span("TOTAL (Jan-Dec)", "Total")], docs=[memo, f25],
      origin=f"derived: {MEMO_TARGET_CRORE} crore = {tgt_k:g} thousand vs {_total(S25)} thousand = {p:.3f}%")

    # ---- out of corpus
    q("g1-ooc-silk-budget-2026", "out_of_corpus", "What is the budgeted Silk Blend revenue for 2026?",
      answerable=False, origin="no 2026 budget exists; the memo gives only a total revenue target")
    q("g1-ooc-revenue-2023", "out_of_corpus", "What was total revenue in 2023?",
      answerable=False, origin="the corpus starts at CY2024")
    q("g1-ooc-ebitda-2025", "out_of_corpus", "What was the EBITDA margin for CY2025?",
      answerable=False, origin="revenue and costs exist but no EBITDA definition, depreciation or tax data")


def _split_for(doc: str) -> str:
    bucket = int(hashlib.sha256(doc.encode("utf-8")).hexdigest(), 16) % 100
    return "dev" if bucket < 70 else "test"


def assign_splits(questions: list[dict]) -> None:
    """The loader rejects a source document that appears in both splits, so
    questions that share documents (multi-document questions) form groups
    that must share one split; each group takes the default hash split of its
    alphabetically first document."""
    parent: dict[str, str] = {}

    def find(x):
        parent.setdefault(x, x)
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    for qq in questions:
        docs = qq["source_documents"]
        for d in docs:
            find(d)
        for d in docs[1:]:
            parent[find(d)] = find(docs[0])
    groups: dict[str, list[str]] = {}
    for d in list(parent):
        groups.setdefault(find(d), []).append(d)
    # Sales-CY2025-Actual.xlsx links most questions into one dev group, so the
    # standalone Production workbook is pinned to test to keep a held-out slice.
    pinned = {"Production-CY2025.xlsx": "test"}
    split_of = {d: pinned.get(min(g), _split_for(min(g))) for g in groups.values() for d in g}
    for qq in questions:
        if qq["source_documents"]:
            qq["split"] = split_of[qq["source_documents"][0]]
        else:
            qq["split"] = _split_for(qq["id"])


def write_gold() -> None:
    gold_questions()
    assign_splits(QUESTIONS)
    for qq in QUESTIONS:
        if not qq["gold_spans"]:
            qq.pop("gold_spans")
        for k in ("must_contain", "must_not_contain"):
            if not qq[k]:
                qq.pop(k)
        if not qq["source_documents"]:
            qq.pop("source_documents")
            qq.pop("formula_dependent")
    header = (
        "# G1 hard numeric gold set (Upgrade/08 section 5.2). GENERATED by eval-public/build_synthetic_corpus_g1.py:\n"
        "# derived answers are computed from the generator's data tables, never typed by hand. SYNTHETIC:\n"
        "# not a milestone population. Do not edit by hand; rerun the generator.\n"
    )
    body = yaml.safe_dump({"version": 1, "population": "synthetic_public", "questions": QUESTIONS},
                          sort_keys=False, allow_unicode=True, width=110)
    GOLD.write_text(header + body, encoding="utf-8")


def main() -> None:
    build_finance()
    build_operations()
    build_strategy()
    _freeze_zip_timestamps(OUT)
    write_gold()
    print(f"wrote synthetic G1 corpus to {OUT} and {len(QUESTIONS)} questions to {GOLD}")


if __name__ == "__main__":
    main()
