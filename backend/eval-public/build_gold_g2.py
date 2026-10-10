"""Generate the G2 multi-step / cross-document gold sets (Upgrade/08 section 5.2).

Writes two gold files over the EXISTING corpora (nothing in either corpus is
modified, and no product code is touched):

  gold-g2-tarnwick.yaml   over eval-public/corpus-synthetic/     (Tarnwick Provisions)
  gold-g2-brindle.yaml    over eval-public/corpus-synthetic-g1/  (Brindle Loom Textiles)

Every expected value is READ from the generated corpus files (workbook cells
via openpyxl, memo/deck sentences via python-docx/python-pptx) and every
derived answer (differences, percentages, "list every month where ...",
argmax) is COMPUTED here; nothing is typed from memory. After building, each
gold span is verified verbatim (after `normalize_text`) against the text the
indexer stores: `XlsxParser` + `chunk_workbook` chunks for workbooks,
`PptxParser` + `chunk_presentation` chunks for the deck, paragraph text for
memos (Docling is not run).

Category tags are carried in the question id (g2t-/g2b- + category) and in
`origin`, because the Question schema forbids extra fields:

  cmp   comparison across workbooks / periods / actual vs budget
  tl    timeline or per-period trace
  enum  exhaustive enumeration ("list every month where ...")
  mis   misleading top hit (decoy row, wrong unit, wrong sheet)
  dis   workbook vs deck/memo (or deck vs memo) disagreement
  amb   ambiguity: a period or version is unstated (clarification wanted)
  miss  missing evidence (out_of_corpus) or revoked source

Run from anywhere:   python eval-public/build_gold_g2.py
"""

from __future__ import annotations

import hashlib
import os
import re
import sys
from pathlib import Path

import yaml
from docx import Document
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
os.environ.setdefault(
    "DOCKET_DATA_DIR",
    "/tmp/claude-1000/-home-pc-Desktop-Rajdeep-Dev-Local-Agent/b10fefe4-c2be-4f17-a32d-75b8491f9030/scratchpad/data",
)

from build_synthetic_corpus_g1 import FULL_MONTH, fmt, num_re, rnd, ZERO, NOT_FOUND  # noqa: E402

TARNWICK = HERE / "corpus-synthetic"
BRINDLE = HERE / "corpus-synthetic-g1"


# --------------------------------------------------------------------------
# Reading the corpora
# --------------------------------------------------------------------------


def fv(v) -> str:
    """Cell value as the indexer prints it (8.0 -> 8)."""
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    return str(v)


class Table:
    """One sheet: header labels per column, rows keyed by first cell. When the
    sheet carries a merged title row above the header the parser labels every
    cell with that title (`label_override`)."""

    def __init__(self, file, sheet, cols, rows, label_override=None):
        self.file, self.sheet, self.cols, self.rows = file, sheet, cols, rows
        self.label_override = label_override

    def row(self, label):
        for r in self.rows:
            if r[0] == label:
                return r
        raise KeyError(f"{self.file}/{self.sheet}: {label}")

    def label(self, i):
        return self.label_override or self.cols[i]

    def col(self, name):
        for i, c in enumerate(self.cols):
            if c == name or c.startswith(name + " ("):
                return i
        raise KeyError(f"{self.file}/{self.sheet}: {name}")

    def val(self, row_label, col_name):
        return self.row(row_label)[self.col(col_name)]

    def span(self, row_label, *col_names, lo=20, hi=60):
        """Longest verbatim run of `Label: value` fragments (present cells, in
        order) of the row, 20-60 chars, that covers the named columns."""
        r = self.row(row_label)
        frags = [(i, f"{self.label(i)}: {fv(r[i])}") for i in range(len(r)) if r[i] is not None]
        want = {self.col(c) for c in col_names}
        best = None
        for a in range(len(frags)):
            for b in range(a, len(frags)):
                idx = {i for i, _ in frags[a : b + 1]}
                if not want <= idx:
                    continue
                text = " ".join(f for _, f in frags[a : b + 1])
                if not lo <= len(text) <= hi:
                    continue
                score = (-len(text), 0 if frags[a][0] == 0 else 1)
                if best is None or score < best[0]:
                    best = (score, text)
        if best is None:
            raise ValueError(f"no span fits for {self.file}/{self.sheet}/{row_label}/{col_names}")
        return best[1]


def load_tables(path: Path) -> dict[str, Table]:
    wb = load_workbook(path, data_only=True)
    out: dict[str, Table] = {}
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c is not None for c in r)]
        h = next(i for i, r in enumerate(rows) if sum(c is not None for c in r) >= 2)
        title = str(rows[0][0]) if h > 0 else None
        cols = [str(c) for c in rows[h]]
        out[ws.title] = Table(path.name, ws.title, cols, [r[: len(cols)] for r in rows[h + 1 :]], title)
    return out


def doc_paragraphs(path: Path) -> list[str]:
    return [p.text for p in Document(path).paragraphs if p.text.strip()]


def deck_lines(path: Path) -> list[str]:
    from pptx import Presentation

    out = []
    for slide in Presentation(path).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
    return out


def find(lines: list[str], pattern: str) -> re.Match:
    for line in lines:
        m = re.search(pattern, line)
        if m:
            return m
    raise KeyError(pattern)


def text_span(m: re.Match) -> str:
    """A 20-60 char verbatim slice of the matched line that contains the match."""
    line = m.string
    start, end = m.start(), m.end()
    assert end - start <= 60, "match too long for a span"
    seg = line[start : start + 60]
    if start + 60 < len(line):
        cut = seg.rfind(" ")
        if cut >= end - start:
            seg = seg[:cut]
    if len(seg) < 20:  # short match near the line end: extend backwards
        start = max(0, end - 60)
        seg = line[start:end]
    seg = seg.strip()
    assert len(seg) >= 20 and line.find(seg) <= start + 1 and m.group(0) in seg, "bad span"
    return seg


# --------------------------------------------------------------------------
# Question plumbing
# --------------------------------------------------------------------------

MON_ABBR = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def month_item(m: str) -> str:
    if m == "May":  # the modal verb must not satisfy the item
        return r"re:\bmay(?!\s+(?:be|have|not|also|include|reflect|indicate|suggest|mean))\b"
    return rf"re:\b(?:{m.lower()}|{FULL_MONTH[m]})\b"


def quarter_item(qn: str) -> str:
    words = {"Q1": "first", "Q2": "second", "Q3": "third", "Q4": "fourth"}[qn]
    return rf"re:\b(?:{qn.lower()}|{words} quarter)\b"


def pat(*values) -> str:
    return "re:" + "|".join(num_re(v) for v in values)


def ppat(p: float) -> str:
    p = abs(p)
    return "re:" + num_re(rnd(p, 1), 1) + "|" + num_re(rnd(p, 2), 2)


def inr_pat(v_inr: float) -> str:
    """A rupee amount as it may be written: full, thousands, millions, lakh."""
    alts = [num_re(v_inr), num_re(v_inr / 1000)]
    alts += [rf"(?<![\d.]){re.escape(fmt(v_inr / 1e6))}0*\s*million", rf"(?<![\d.]){re.escape(fmt(v_inr / 1e5))}0*\s*lakh"]
    return "re:" + "|".join(alts)


class Gold:
    def __init__(self, prefix: str, corpus_name: str, generator: str = "eval-public/build_gold_g2.py",
                 tag: str = "g2 category"):
        self.prefix, self.corpus_name = prefix, corpus_name
        self.generator, self.tag = generator, tag
        self.questions: list[dict] = []

    def q(self, cat, slug, qtype, question, *, must=(), not_=(), spans=(), docs=(), origin,
          enumeration=(), answerable=True, setup=None):
        d = {
            "id": f"{self.prefix}-{cat}-{slug}",
            "type": qtype,
            "question": question,
            "answerable": answerable,
        }
        if must:
            d["must_contain"] = list(must)
        if not_:
            d["must_not_contain"] = list(not_)
        if spans:
            d["gold_spans"] = list(spans)
        if enumeration:
            d["enumeration"] = list(enumeration)
        if setup:
            d["setup"] = {"revoke": list(setup)}
        d["reviewed"] = True
        d["origin"] = f"{self.corpus_name} ({self.generator}); {self.tag} {cat}; {origin}"
        d["source_documents"] = list(docs)
        d["formula_dependent"] = False
        self.questions.append(d)


def _split_for(doc: str) -> str:
    return "dev" if int(hashlib.sha256(doc.encode()).hexdigest(), 16) % 100 < 70 else "test"


def assign_splits(questions: list[dict]) -> None:
    """Documents shared by a question must share a split (the loader rejects a
    document in both): union-find groups, each taking the hash split of its
    alphabetically first document."""
    parent: dict[str, str] = {}

    def find_(x):
        parent.setdefault(x, x)
        while parent[x] != x:
            parent[x] = parent[parent[x]]
            x = parent[x]
        return x

    for qq in questions:
        docs = qq["source_documents"]
        for d in docs:
            find_(d)
        for d in docs[1:]:
            parent[find_(d)] = find_(docs[0])
    groups: dict[str, list[str]] = {}
    for d in list(parent):
        groups.setdefault(find_(d), []).append(d)
    split_of = {d: _split_for(min(g)) for g in groups.values() for d in g}
    for qq in questions:
        qq["split"] = split_of[qq["source_documents"][0]] if qq["source_documents"] else _split_for(qq["id"])


# --------------------------------------------------------------------------
# Tarnwick Provisions
# --------------------------------------------------------------------------


def build_tarnwick() -> Gold:
    g = Gold("g2t", "synthetic corpus eval-public/corpus-synthetic (Tarnwick)")
    fin, ops, strat, ppl = (TARNWICK / d for d in ("finance", "operations", "strategy", "people"))
    R24 = load_tables(fin / "Revenue-FY2024-25.xlsx")["Monthly Revenue"]
    R25 = load_tables(fin / "Revenue-FY2025-26.xlsx")["Monthly Revenue"]
    RB = load_tables(fin / "Revenue-FY2025-26-Budget.xlsx")["Monthly Revenue"]
    ex = load_tables(fin / "Expenses-FY2025-26.xlsx")
    OPEX, CAPEX = ex["Operating Costs"], ex["Capital Spend"]
    WH = load_tables(ops / "Warehouse-Throughput-Q2.xlsx")["Warehouse Throughput"]
    DIR = load_tables(ppl / "Team-Directory-Sep-2025.xlsx")["Team Directory"]
    memo = doc_paragraphs(strat / "Board-Memo-Q2-FY2025-26.docx")
    deck = deck_lines(ops / "Ops-Review-Q2-FY2025-26.pptx")
    policy = doc_paragraphs(ppl / "Compensation-Policy-Note.docx")

    f24, f25, fb = "Revenue-FY2024-25.xlsx", "Revenue-FY2025-26.xlsx", "Revenue-FY2025-26-Budget.xlsx"
    fex, fwh = "Expenses-FY2025-26.xlsx", "Warehouse-Throughput-Q2.xlsx"
    fdeck, fmemo = "Ops-Review-Q2-FY2025-26.pptx", "Board-Memo-Q2-FY2025-26.docx"
    fdir, fpol = "Team-Directory-Sep-2025.xlsx", "Compensation-Policy-Note.docx"
    M6 = ["Apr", "May", "Jun", "Jul", "Aug", "Sep"]
    Q2M = ["Jul", "Aug", "Sep"]

    # ---- cmp: comparison across workbooks and periods
    a, b = R24.val("Q2 subtotal", "Total"), R25.val("Q2 subtotal", "Total") * 1000
    p = (b - a) / a * 100
    g.q("cmp", "q2-yoy-pct", "multi_doc",
        "By what percentage did total revenue for Q2 (July to September) grow from FY2024-25 to FY2025-26?",
        must=[ppat(p)], spans=[R24.span("Q2 subtotal", "Total"), R25.span("Q2 subtotal", "Total")], docs=[f24, f25],
        origin=f"derived: FY2025-26 {R25.val('Q2 subtotal','Total')} thousand vs FY2024-25 {a} INR = {p:.3f}% (units differ)")

    d = R25.val("Sep", "Marigold Chai") * 1000 - R24.val("Sep", "Marigold Chai")
    g.q("cmp", "chai-sep-gap", "multi_doc",
        "By how much, in INR thousands, did Marigold Chai revenue in September 2025 exceed September 2024?",
        must=[pat(d / 1000, d)], spans=[R24.span("Sep", "Marigold Chai"), R25.span("Sep", "Marigold Chai")], docs=[f24, f25],
        origin=f"derived: {R25.val('Sep','Marigold Chai')} thousand minus {R24.val('Sep','Marigold Chai')} INR = {d} INR")

    d = RB.val("Q2 subtotal", "Total") - R25.val("Q2 subtotal", "Total")
    assert d > 0
    g.q("cmp", "q2-budget-shortfall", "multi_doc",
        "By how much, in INR thousands, did actual Q2 FY2025-26 total revenue fall short of the Q2 budget?",
        must=[pat(d, d * 1000)], spans=[R25.span("Q2 subtotal", "Total"), RB.span("Q2 subtotal", "Total")], docs=[f25, fb],
        origin=f"derived: budget {RB.val('Q2 subtotal','Total')} minus actual {R25.val('Q2 subtotal','Total')} = {d}")

    s24 = R24.val("Q2 subtotal", "Marigold Chai") / R24.val("Q2 subtotal", "Total") * 100
    s25 = R25.val("Q2 subtotal", "Marigold Chai") / R25.val("Q2 subtotal", "Total") * 100
    g.q("cmp", "chai-share-q2", "multi_doc",
        "What share of Q2 total revenue came from Marigold Chai in FY2024-25 and in FY2025-26, and did the share rise or fall?",
        must=[ppat(s24), ppat(s25), r"re:fell|fall|declin|drop|decreas|lower|down|reduc"],
        spans=[R24.span("Q2 subtotal", "Marigold Chai"), R24.span("Q2 subtotal", "Total"),
               R25.span("Q2 subtotal", "Marigold Chai"), R25.span("Q2 subtotal", "Total")],
        docs=[f24, f25], origin=f"derived: {s24:.2f}% then {s25:.2f}% (a share, so units cancel)")

    # ---- tl: timeline / trace
    vals = [(m, R25.val(m, "Juniper Masala")) for m in M6]
    assert vals[0][1] is None and vals[1][1] == 0
    nz = [(m, v) for m, v in vals if v]
    g.q("tl", "masala-ramp", "enumeration",
        "Trace actual Juniper Masala revenue month by month from June to September 2025, in INR thousands.",
        enumeration=[pat(v) for _, v in nz], spans=[R25.span(m, "Juniper Masala") for m, _ in nz], docs=[f25],
        origin="per-month trace of " + ", ".join(f"{m}={v}" for m, v in nz) + "; April blank, May literal 0 precede it")

    tot24 = {m: R24.val(m, "Total") for m in M6}
    hi, lo_ = max(tot24, key=tot24.get), min(tot24, key=tot24.get)
    g.q("tl", "fy24-peak-trough", "numeric",
        "Across April to September FY2024-25, which month had the highest total revenue and which had the lowest, and what were the two amounts in INR?",
        must=[month_item(hi), month_item(lo_), inr_pat(tot24[hi]), inr_pat(tot24[lo_])],
        spans=[R24.span(hi, "Total"), R24.span(lo_, "Total")], docs=[f24],
        origin=f"derived: argmax {hi}={tot24[hi]}, argmin {lo_}={tot24[lo_]} over six monthly rows (subtotal and TOTAL rows excluded)")

    tc = {m: OPEX.val(m, "Total") for m in Q2M}
    top = max(tc, key=tc.get)
    g.q("tl", "opex-trend", "numeric",
        "How did total operating costs move month by month from July to September 2025, and which month was highest?",
        must=[pat(tc["Jul"]), pat(tc["Aug"]), pat(tc["Sep"]), month_item(top)],
        spans=[OPEX.span(m, "Total") for m in Q2M], docs=[fex],
        origin=f"derived: totals {tc}; highest is {top}; sheet is INR thousands and its parser labels are the merged title")

    # ---- enum: exhaustive "list every month where ..."
    beat = [m for m in M6 if R25.val(m, "Total") > RB.val(m, "Total")]
    g.q("enum", "beat-budget-months", "enumeration",
        "List every month of FY2025-26 (April to September) in which actual total revenue exceeded the budget.",
        enumeration=[month_item(m) for m in beat],
        spans=[R25.span(m, "Total") for m in beat] + [RB.span(m, "Total") for m in beat[:2]], docs=[f25, fb],
        origin=f"derived: {beat}; September ({R25.val('Sep','Total')} vs {RB.val('Sep','Total')}) is the only miss")

    below = [m for m in Q2M if WH.val(m, "On-Time Delivery %") < 0.95]
    g.q("enum", "otd-below-target", "enumeration",
        "List every month in Q2 FY2025-26 where on-time delivery was below the 95% target.",
        enumeration=[month_item(m) for m in below], spans=[WH.span(m, "On-Time Delivery %") for m in below], docs=[fwh],
        origin=f"derived from the warehouse workbook: {below}")

    grew = []
    for i in range(1, 6):
        prev, cur = tot24[M6[i - 1]], tot24[M6[i]]
        if (cur - prev) / prev > 0.04:
            grew.append(M6[i])
    g.q("enum", "fy24-mom-over-4pct", "enumeration",
        "List every month of FY2024-25 (April to September) in which total revenue grew by more than 4% over the previous month.",
        enumeration=[month_item(m) for m in grew],
        spans=[R24.span(m, "MoM Growth %") for m in grew], docs=[f24],
        origin=f"derived from month totals (not the rounded growth column): {grew}")

    big = [m for m in M6 if (R25.val(m, "Juniper Masala") or 0) > 900]
    g.q("enum", "masala-over-900", "enumeration",
        "List every month of FY2025-26 in which actual Juniper Masala revenue was above INR 900 thousand.",
        enumeration=[month_item(m) for m in big], spans=[R25.span(m, "Juniper Masala") for m in big], docs=[f25],
        origin=f"derived: {big}; June ({R25.val('Jun','Juniper Masala')}) is the near miss")

    # ---- mis: misleading top hits
    opex, cap = OPEX.val("Q2 TOTAL", "Total"), CAPEX.val("Q2 TOTAL", "Total")
    g.q("mis", "opex-vs-capex-q2", "multi_doc",
        "How much more, in INR thousands, did Tarnwick spend on operating costs than on capital spend in Q2 FY2025-26?",
        must=[pat(opex - cap, (opex - cap) * 1000)],
        spans=[OPEX.span("Q2 TOTAL", "Total"), CAPEX.span("Q2 TOTAL", "Total")], docs=[fex],
        origin=f"derived: {opex} minus {cap} = {opex - cap}; both sheets have a Q2 TOTAL row and the same merged-title layout")

    pk = find(deck, r"Peak daily dispatch reached (\d+) orders on (\d+) August")
    aug_orders = WH.val("Aug", "Orders Shipped")
    g.q("mis", "peak-dispatch-vs-monthly", "multi_doc",
        "What was the peak daily dispatch in August 2025, and how many orders did the warehouse ship in total that month?",
        must=[pat(int(pk.group(1))), pat(aug_orders)],
        spans=[text_span(pk), WH.span("Aug", "Orders Shipped")], docs=[fdeck, fwh],
        origin=f"deck says {pk.group(1)} on 14 August (a daily peak); the workbook monthly total is {aug_orders}")

    # ---- dis: workbook vs deck / memo disagreements
    t_deck = find(deck, r"Q3 revenue target: INR ([\d.]+) million")
    t_memo = find(memo, r"Q3 FY2025-26 revenue target is INR ([\d.]+) million")
    assert t_deck.group(1) != t_memo.group(1)
    g.q("dis", "q3-target-deck-vs-memo", "multi_doc",
        "What is the Q3 FY2025-26 revenue target, and do the operations review deck and the board memo agree on it?",
        must=[pat(float(t_deck.group(1))), pat(float(t_memo.group(1)))],
        spans=[text_span(t_deck), text_span(t_memo)], docs=[fdeck, fmemo],
        origin=f"deck {t_deck.group(1)} million vs memo {t_memo.group(1)} million: a genuine disagreement")

    otd_deck = find(deck, r"On-time delivery in September was ([\d.]+)%")
    otd_wb = WH.val("Sep", "On-Time Delivery %") * 100
    assert abs(float(otd_deck.group(1)) - otd_wb) > 0.01
    g.q("dis", "sep-otd-workbook-vs-deck", "multi_doc",
        "What was September 2025 on-time delivery in the warehouse workbook, and does the operations review deck report the same figure?",
        must=[ppat(otd_wb), pat(float(otd_deck.group(1)))],
        spans=[WH.span("Sep", "On-Time Delivery %"), text_span(otd_deck)], docs=[fwh, fdeck],
        origin=f"workbook {otd_wb:.1f}% vs deck {otd_deck.group(1)}%: a genuine disagreement")

    avg = find(memo, r"averaged ([\d.]+)% in Q2")
    mean = sum(WH.val(m, "On-Time Delivery %") for m in Q2M) / 3 * 100
    assert round(mean, 1) == float(avg.group(1))
    g.q("dis", "memo-otd-avg-check", "multi_doc",
        "Does the board memo's Q2 average on-time delivery figure match the monthly figures in the warehouse workbook?",
        must=[ppat(mean), r"re:\b(?:yes|match|agree|consistent|same|correct)"],
        spans=[text_span(avg)] + [WH.span(m, "On-Time Delivery %") for m in Q2M], docs=[fmemo, fwh],
        origin=f"consistent: mean of the three months = {mean:.2f}% rounds to the memo's {avg.group(1)}%")

    rev_deck = find(deck, r"Q2 revenue closed at INR ([\d.]+) million")
    rev_memo = find(memo, r"Revenue for Q2 FY2025-26 was INR ([\d.]+) million")
    q2 = R25.val("Q2 subtotal", "Total")
    g.q("dis", "q2-revenue-three-sources", "multi_doc",
        "What is Q2 FY2025-26 total revenue according to the revenue workbook, the operations review deck and the board memo?",
        must=[pat(q2, q2 * 1000), pat(float(rev_deck.group(1))), pat(float(rev_memo.group(1)))],
        spans=[R25.span("Q2 subtotal", "Total"), text_span(rev_deck), text_span(rev_memo)], docs=[f25, fdeck, fmemo],
        origin=f"workbook {q2} thousand, deck {rev_deck.group(1)} million, memo {rev_memo.group(1)} million: same value, different rounding")

    ratio = find(memo, r"or ([\d.]+)% of revenue")
    cost = OPEX.val("Q2 TOTAL", "Total") / R25.val("Q2 subtotal", "Total") * 100
    assert round(cost, 1) == float(ratio.group(1))
    g.q("dis", "memo-cost-ratio-check", "multi_doc",
        "Operating costs as a share of Q2 FY2025-26 revenue: compute it from the expenses and revenue workbooks and say whether the board memo's figure agrees.",
        must=[ppat(cost)],
        spans=[OPEX.span("Q2 TOTAL", "Total"), R25.span("Q2 subtotal", "Total"), text_span(ratio)],
        docs=[fex, f25, fmemo], origin=f"derived: {OPEX.val('Q2 TOTAL','Total')} / {q2} = {cost:.3f}%; memo says {ratio.group(1)}%")

    # ---- amb: period not stated; clarification wanted (existing encoding:
    # multi_doc requiring BOTH fiscal-year values; asking which year is also valid)
    note = ("Docket has no clarify-first behaviour yet, so this is encoded as multi_doc requiring BOTH actual "
            "fiscal-year values; a clarification question that asks which fiscal year is also acceptable")
    for slug, month, col, q_text in (
        ("july-total", "Jul", "Total", "What was total revenue in July?"),
        ("h1-total", "TOTAL (Apr-Sep)", "Total", "What was total revenue for April to September?"),
    ):
        a24, a25 = R24.val(month, col), R25.val(month, col)
        g.q("amb", slug, "multi_doc", q_text, must=[inr_pat(a24), pat(a25, a25 * 1000)],
            spans=[R24.span(month, col), R25.span(month, col)], docs=[f24, f25],
            origin=f"ambiguous fiscal year: FY2024-25 {a24} INR, FY2025-26 actual {a25} thousand "
                   f"(budget workbook has {RB.val(month, col)} as a third reading); {note}")

    a24, a25 = R24.val("Q1 subtotal", "Saffron Rusk"), R25.val("Q1 subtotal", "Saffron Rusk")
    g.q("amb", "q1-rusk", "multi_doc", "What was Saffron Rusk revenue in Q1?",
        must=[inr_pat(a24), pat(a25, a25 * 1000)],
        spans=[R24.span("Q1 subtotal", "Saffron Rusk"), R25.span("Q1 subtotal", "Saffron Rusk")], docs=[f24, f25],
        origin=f"ambiguous fiscal year: {a24} INR vs {a25} thousand; {note}")

    # ---- miss: missing and revoked evidence
    g.q("miss", "ooc-q3-actual", "out_of_corpus", "What was actual total revenue for Q3 FY2025-26?", answerable=False,
        origin="only the Q3 target exists (deck and memo); no Q3 actuals in any file")
    g.q("miss", "ooc-marketing-aug-2024", "out_of_corpus", "How much was spent on marketing in August 2024?",
        answerable=False, origin="the expenses workbook covers only Q2 FY2025-26; no FY2024-25 costs")

    nag = [r for r in DIR.rows if r[3] == "Nagpur"]
    nag_sum = sum(r[4] for r in nag)
    names = [r[0] for r in nag]
    g.q("miss", "rev-nagpur-salaries-vs-logistics", "revoked",
        "How does the combined monthly salary of the Nagpur-based employees compare with September 2025 logistics cost?",
        answerable=False, not_=[pat(nag_sum)] + [pat(r[4]) for r in nag],
        spans=[f"Employee: {r[0]} Role: {r[1]}" for r in nag] + [OPEX.span("Sep", "Logistics")],
        docs=[fdir, fex], setup=["people"],
        origin=f"people/ is revoked: {names} salaries sum to {nag_sum} and must not appear; logistics cost is still available")

    fc = next(r for r in DIR.rows if r[1] == "Finance Controller")
    pct = float(find(policy, r"bonus is (\d+)% of annual base salary").group(1))
    bonus = fc[4] * 12 * pct / 100
    g.q("miss", "rev-controller-bonus", "revoked",
        "Using the compensation policy's bonus percentage, what annual performance bonus would the Finance Controller receive?",
        answerable=False, not_=[pat(bonus), pat(fc[4])],
        spans=[text_span(find(policy, r"The annual performance bonus is \d+% of annual base salary")),
               f"Employee: {fc[0]} Role: {fc[1]}"],
        docs=[fdir, fpol], setup=["people"],
        origin=f"people/ is revoked: bonus would be {pct:g}% x 12 x {fc[4]} = {bonus:g}; neither it nor the salary may appear")
    return g


# --------------------------------------------------------------------------
# Brindle Loom Textiles
# --------------------------------------------------------------------------


def build_brindle() -> Gold:
    g = Gold("g2b", "synthetic corpus eval-public/corpus-synthetic-g1 (Brindle)")
    fin, ops, strat = (BRINDLE / d for d in ("finance", "operations", "strategy"))
    S24 = load_tables(fin / "Sales-CY2024.xlsx")["Sales by Month"]
    S25 = load_tables(fin / "Sales-CY2025-Actual.xlsx")["Sales Actual"]
    B25 = load_tables(fin / "Sales-CY2025-Budget.xlsx")["Sales Budget"]
    cost_wb = load_tables(fin / "Costs-CY2025.xlsx")
    COST, CAPEX, NOTES = cost_wb["Cost Summary"], cost_wb["Plant Capex"], cost_wb["Notes"]
    NOTES25 = load_tables(fin / "Sales-CY2025-Actual.xlsx")["Notes"]
    PROD = load_tables(ops / "Production-CY2025.xlsx")["Loom Output"]
    memo = doc_paragraphs(strat / "Annual-Review-Memo-CY2025.docx")

    f24, f25, fb = "Sales-CY2024.xlsx", "Sales-CY2025-Actual.xlsx", "Sales-CY2025-Budget.xlsx"
    fc, fp, fm = "Costs-CY2025.xlsx", "Production-CY2025.xlsx", "Annual-Review-Memo-CY2025.docx"

    def notes_span(label: str, limit=60) -> str:
        detail = next(r[1] for r in NOTES.rows if r[0] == label)
        return str(detail)[:limit].rstrip()

    # ---- cmp
    a = S24.val("Q1 subtotal", "Total") + S24.val("Q2 subtotal", "Total")  # full INR
    b = (S25.val("Q1 subtotal", "Total") + S25.val("Q2 subtotal", "Total")) * 1000
    p = (b - a) / a * 100
    g.q("cmp", "h1-yoy-pct", "multi_doc",
        "By what percentage did total sales for the first half (January to June) grow from CY2024 to CY2025 actual?",
        must=[ppat(p)],
        spans=[S24.span("Q1 subtotal", "Total"), S24.span("Q2 subtotal", "Total"),
               S25.span("Q1 subtotal", "Total"), S25.span("Q2 subtotal", "Total")], docs=[f24, f25],
        origin=f"derived: H1 CY2025 {b/1000:g} thousand vs H1 CY2024 {a} INR = {p:.3f}% (units differ; subtotal rows summed)")

    beat = [m for m in MON_ABBR if (S25.val(m, "Silk Blend") is not None and B25.val(m, "Silk Blend") is not None
                                    and S25.val(m, "Silk Blend") > B25.val(m, "Silk Blend"))]
    g.q("enum", "silk-beat-budget-months", "enumeration",
        "List every month of 2025 in which actual Silk Blend revenue exceeded its budget.",
        enumeration=[month_item(m) for m in beat],
        spans=[S25.span(m, "Silk Blend") for m in beat], docs=[f25, fb],
        origin=f"derived: {beat}; blank January and equal February (0 vs 0) never count; May/Jun/Sep miss narrowly")

    cost = COST.val("CY2025 TOTAL", "Total")  # lakh
    rev = S25.val("TOTAL (Jan-Dec)", "Total")  # thousand
    ratio = cost * 100 / rev * 100
    g.q("cmp", "cost-to-revenue-ratio", "multi_doc",
        "What share of CY2025 actual revenue did total operating costs represent?",
        must=[ppat(ratio)],
        spans=[COST.span("CY2025 TOTAL", "Total"), S25.span("TOTAL (Jan-Dec)", "Total"),
               notes_span("Units"), f"Detail: {next(r[1] for r in NOTES25.rows if r[0] == 'Units')}"], docs=[fc, f25],
        origin=f"derived: {cost} lakh = {cost*100:g} thousand over {rev} thousand = {ratio:.3f}% (lakh vs thousands trap)")

    capex = CAPEX.val("CY2025 TOTAL", "Total")  # crore
    diff = rnd(cost - capex * 100, 2)
    g.q("cmp", "costs-vs-capex-lakh", "multi_doc",
        "By how much, in INR lakh, did total CY2025 operating costs exceed total CY2025 plant capex?",
        must=[pat(diff), ],
        spans=[COST.span("CY2025 TOTAL", "Total"), CAPEX.span("CY2025 TOTAL", "Total"), notes_span("Units")], docs=[fc],
        origin=f"derived: {cost} lakh minus {capex} crore x 100 = {diff} lakh (crore sheet under a merged title row)")

    wo = float(re.search(r"INR ([\d.]+) lakh", next(r[1] for r in NOTES.rows if r[0] == "Inventory write-off")).group(1))
    tot = rnd(cost + wo, 2)
    g.q("cmp", "cost-incl-writeoff", "multi_doc",
        "What were total CY2025 costs in INR lakh once the inventory write-off is added to the cost summary total?",
        must=[pat(tot)],
        spans=[COST.span("CY2025 TOTAL", "Total"), notes_span("Inventory write-off", 40)], docs=[fc],
        origin=f"derived: {cost} + {wo} = {tot}; the write-off exists only on the Notes sheet")

    # ---- tl
    s24 = next(m for m in MON_ABBR if S24.val(m, "Silk Blend") not in (None, 0))
    s25 = next(m for m in MON_ABBR if S25.val(m, "Silk Blend") not in (None, 0))
    g.q("tl", "silk-first-sales", "multi_doc",
        "In which month did Silk Blend first record revenue in CY2024, and in which month did it first record non-zero revenue in CY2025?",
        must=[month_item(s24), pat(S24.val(s24, "Silk Blend")), month_item(s25), pat(S25.val(s25, "Silk Blend"))],
        spans=[S24.span(s24, "Silk Blend"), S25.span(s25, "Silk Blend"), S25.span("Feb", "Silk Blend")], docs=[f24, f25],
        origin=f"derived: first CY2024 {s24}={S24.val(s24,'Silk Blend')} INR; CY2025 blank Jan, 0 Feb, first non-zero {s25}={S25.val(s25,'Silk Blend')} thousand")

    low = [m for m in MON_ABBR if PROD.val(m, "Defect Rate") < 2.0]
    g.q("tl", "defect-below-2pct", "enumeration",
        "List every month of 2025 in which the loom defect rate was below 2%.",
        enumeration=[month_item(m) for m in low], spans=[PROD.span(m, "Defect Rate") for m in low], docs=[fp],
        origin=f"derived: {low}; August is included although output was 0 (shutdown); July (2.0) is not below")

    none = [m for m in MON_ABBR if PROD.val(m, "Output") in (None, 0)]
    g.q("enum", "no-output-months", "enumeration",
        "List every month of 2025 in which the production workbook shows no output (blank or zero).",
        enumeration=[month_item(m) for m in none],
        spans=[PROD.span(m, "Defect Rate") for m in none], docs=[fp], origin=f"derived: {none}; one is blank (no reading) and one is a literal 0 (shutdown)")

    qs = ["Q1", "Q2", "Q3", "Q4"]
    over = [qn for qn in qs if S25.val(f"{qn} subtotal", "Total") > B25.val(f"{qn} subtotal", "Total")]
    g.q("enum", "quarters-over-budget", "enumeration",
        "List every quarter of 2025 in which actual total sales exceeded the budget.",
        enumeration=[quarter_item(qn) for qn in over],
        spans=[S25.span(f"{qn} subtotal", "Total") for qn in over] + [B25.span(f"{qn} subtotal", "Total") for qn in over],
        docs=[f25, fb], origin=f"derived from the subtotal rows: {over}")

    # ---- mis
    months_tot = {m: S25.val(m, "Total") for m in MON_ABBR}
    best = max(months_tot, key=months_tot.get)
    g.q("mis", "peak-sales-month-output", "multi_doc",
        "Which month of 2025 had the highest actual total sales, and how many metres of cloth were produced in that month?",
        must=[month_item(best), pat(PROD.val(best, "Output"))],
        not_=[pat(S25.val("Memo: bulk export order (not in totals)", "Total"))],
        spans=[S25.span(best, "Total"), PROD.span(best, "Output")], docs=[f25, fp],
        origin=f"derived: argmax over month rows = {best} ({months_tot[best]} thousand); output {PROD.val(best,'Output')}; "
               "hidden export memo row and subtotal rows are decoys")

    # ---- dis
    g_ = find(memo, r"up ([\d.]+)% on CY2024")
    growth = (S25.val("TOTAL (Jan-Dec)", "Total") * 1000 - S24.val("TOTAL (Jan-Dec)", "Total")) / S24.val("TOTAL (Jan-Dec)", "Total") * 100
    assert round(growth, 1) == float(g_.group(1))
    g.q("dis", "memo-growth-check", "multi_doc",
        "Compute CY2025 revenue growth over CY2024 from the two sales workbooks and say whether the annual review memo's figure agrees.",
        must=[ppat(growth)],
        spans=[S25.span("TOTAL (Jan-Dec)", "Total"), S24.span("TOTAL (Jan-Dec)", "Total"), text_span(g_)],
        docs=[f25, f24, fm], origin=f"consistent: {growth:.3f}% vs memo {g_.group(1)}%; workbooks differ in units")

    tgt = find(memo, r"CY2026 revenue target is INR ([\d.]+) crore")
    tgt_k = float(tgt.group(1)) * 10000
    mult = tgt_k / S25.val("TOTAL (Jan-Dec)", "Total")
    g.q("dis", "memo-target-multiple", "multi_doc",
        "How many times CY2025 actual revenue is the CY2026 revenue target set in the annual review memo?",
        must=["re:" + num_re(rnd(mult, 1), 1) + "|" + num_re(rnd(mult, 2), 2)],
        spans=[text_span(tgt), S25.span("TOTAL (Jan-Dec)", "Total")], docs=[fm, f25],
        origin=f"derived: {tgt.group(1)} crore = {tgt_k:g} thousand over {S25.val('TOTAL (Jan-Dec)','Total')} thousand = {mult:.3f}x")

    # ---- amb
    note = ("Docket has no clarify-first behaviour yet, so this is encoded as multi_doc requiring BOTH actual "
            "calendar-year values; a clarification question that asks which year is also acceptable")
    for slug, row, col, q_text in (
        ("march-total", "Mar", "Total", "What was total sales in March?"),
        ("q4-silk", "Q4 subtotal", "Silk Blend", "What was Silk Blend revenue in Q4?"),
        ("dec-cotton", "Dec", "Cotton Yarn", "What was Cotton Yarn revenue in December?"),
    ):
        a24, a25 = S24.val(row, col), S25.val(row, col)
        g.q("amb", slug, "multi_doc", q_text, must=[inr_pat(a24), pat(a25, a25 * 1000)],
            spans=[S24.span(row, col), S25.span(row, col)], docs=[f24, f25],
            origin=f"ambiguous year: CY2024 {a24} INR, CY2025 actual {a25} thousand (budget {B25.val(row, col)} is a third reading); {note}")

    # ---- miss
    g.q("miss", "ooc-cotton-jan-2026", "out_of_corpus", "What was Cotton Yarn revenue in January 2026?",
        answerable=False, origin="the corpus ends at CY2025; only a CY2026 total target exists")
    g.q("miss", "ooc-capex-2024", "out_of_corpus", "How much was spent on plant capex in CY2024?", answerable=False,
        origin="Plant Capex covers CY2025 only")

    out_total = PROD.val("TOTAL (Jan-Dec)", "Output")
    per_k = out_total / S25.val("TOTAL (Jan-Dec)", "Total")
    g.q("miss", "rev-output-per-sales", "revoked",
        "How many metres of cloth were produced per INR thousand of actual 2025 sales?",
        answerable=False, not_=[pat(out_total), "re:" + num_re(rnd(per_k, 2), 2) + "|" + num_re(rnd(per_k, 1), 1)],
        spans=[PROD.span("TOTAL (Jan-Dec)", "Output"), S25.span("TOTAL (Jan-Dec)", "Total")],
        docs=[fp, f25], setup=["operations"],
        origin=f"operations/ is revoked: {out_total} metres / {S25.val('TOTAL (Jan-Dec)','Total')} thousand = {per_k:.3f}; output total must not appear")

    g.q("miss", "rev-memo-target-vs-actual", "revoked",
        "What revenue target does the annual review memo set for CY2026, and how much higher is it than CY2025 actual revenue?",
        answerable=False, not_=[rf"re:(?<![\d.]){re.escape(tgt.group(1))}(?!\d)"],
        spans=[text_span(tgt), S25.span("TOTAL (Jan-Dec)", "Total")], docs=[fm, f25], setup=["strategy"],
        origin=f"strategy/ is revoked: the memo's target ({tgt.group(1)} crore) must not appear; CY2025 actuals remain available")
    return g


# --------------------------------------------------------------------------
# Verification against indexed text
# --------------------------------------------------------------------------


def indexed_texts(roots: tuple[Path, ...] | None = None) -> dict[str, list[str]]:
    from docket.eval.scoring import normalize_text
    from docket.infra.parsing.pptx_chunker import chunk_presentation
    from docket.infra.parsing.pptx_wrapper import PptxParser
    from docket.infra.parsing.xlsx_chunker import chunk_workbook
    from docket.infra.parsing.xlsx_wrapper import XlsxParser

    out: dict[str, list[str]] = {}
    for root in roots or (TARNWICK, BRINDLE):
        for path in sorted(root.rglob("*")):
            if path.suffix == ".xlsx":
                _, chunks = chunk_workbook(XlsxParser().parse("s", path))
                out[path.name] = [normalize_text(c.text) for c in chunks]
            elif path.suffix == ".pptx":
                _, chunks = chunk_presentation(PptxParser().parse("s", path))
                out[path.name] = [normalize_text(c.text) for c in chunks]
            elif path.suffix == ".docx":
                out[path.name] = [normalize_text(t) for t in doc_paragraphs(path)]
    return out


def verify(g: Gold, texts: dict[str, list[str]]) -> None:
    from docket.eval.scoring import normalize_text

    for qq in g.questions:
        docs = qq["source_documents"]
        pool = [t for d in docs for t in texts[d]]
        for span in qq.get("gold_spans", []):
            n = normalize_text(span)
            if not any(n in t for t in pool):
                raise SystemExit(f"{qq['id']}: span not verbatim in {docs}: {span!r}")


def write(g: Gold, path: Path, label: str) -> None:
    assign_splits(g.questions)
    header = (
        f"# G2 multi-step / cross-document gold set over {label} (Upgrade/08 section 5.2; Upgrade/06 section 10.3).\n"
        "# GENERATED by eval-public/build_gold_g2.py: expected values are read from the corpus files and derived\n"
        "# answers are computed in code, never typed. Category tags live in the id and origin. SYNTHETIC: not a\n"
        "# milestone population. Do not edit by hand; rerun the generator.\n"
    )
    body = yaml.safe_dump({"version": 1, "population": "synthetic_public", "questions": g.questions},
                          sort_keys=False, allow_unicode=True, width=110)
    path.write_text(header + body, encoding="utf-8")


def main() -> None:
    texts = indexed_texts()
    for builder, name, label in ((build_tarnwick, "gold-g2-tarnwick.yaml", "corpus-synthetic (Tarnwick)"),
                                 (build_brindle, "gold-g2-brindle.yaml", "corpus-synthetic-g1 (Brindle)")):
        g = builder()
        verify(g, texts)
        write(g, HERE / name, label)
        cats: dict[str, int] = {}
        for qq in g.questions:
            c = qq["id"].split("-")[1]
            cats[c] = cats.get(c, 0) + 1
        print(f"{name}: {len(g.questions)} questions {cats}")


if __name__ == "__main__":
    main()
