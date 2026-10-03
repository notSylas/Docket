"""Splits a parsed `.xlsx` workbook (`xlsx_wrapper.ParsedWorkbook`) into
`EvidenceUnitDraft`/`ChunkDraft` lists -- the same intermediate shape
`docket.infra.parsing.chunker.chunk_document` produces for Docling markdown, so
`docket.services.ingestion.chunk_writer.ChunkWriter.persist_units_and_chunks` and
`docket.services.ingestion.pipeline.IngestionPipeline` can consume either without
caring which parser produced it.

Granularity decision (EvidenceUnit = one populated row, not one cell):

Doc 02 section 4 requires reading "values together with row/column headers"
and section 6 warns against flattening "a table into unrelated numbers". A
single numeric cell, cited alone, loses exactly the context that makes it
evidence rather than a bare number -- which column it's under, which sheet,
whether it's hidden/filtered, whether it's a formula result or a literal.
One `EvidenceUnit`/`Chunk` per populated spreadsheet row keeps every cell's
header label attached to it in the same citable unit, mirrors how a person
reads a spreadsheet row as one statement ("for January, revenue was X and
cost was Y"), and keeps units small enough that retrieval doesn't have to
pull in unrelated rows to get header context. `unit_kind="range"` reflects
that a row is a (possibly single-cell) horizontal range of cells, not one
isolated cell; `locator_json["range"]` is always a `"<start>:<end>"` cell
range string, even for a row with only one populated column (e.g.
`"B5:B5"`), so the locator shape never depends on how many cells happened to
be populated. A future adapter that genuinely needs single-cell citations
(e.g. a lookup pinpointing one value with no useful row context) can
introduce `unit_kind="cell"` without this module needing to change.

Every chunk gets `provenance="extracted"` here -- this checkpoint never
produces a Docket-computed derivation, and a formula's cached result is
still "what the workbook states", not something Docket computed (doc 02
section 4, case 2 vs. case 3). A real `"derived"` producer is future work;
nothing here blocks that column from being used later.
"""

from __future__ import annotations

import hashlib
import json
from datetime import date, datetime

from openpyxl.utils import get_column_letter

from docket.core.config import settings
from docket.infra.parsing.chunker import ChunkDraft, EvidenceUnitDraft
from docket.infra.parsing.xlsx_context import build_sheet_contexts, row_context
from docket.infra.parsing.tokens import TokenCounter, get_token_counter, part_locator, split_to_fit
from docket.infra.parsing.xlsx_wrapper import (
    MISSING_FORMULA_CACHE_ERROR,
    CellData,
    ParsedWorkbook,
    RowData,
    SheetData,
)


def _sha256_hex(text: str) -> str:
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


def _format_raw_value(cell: CellData) -> str:
    value = cell.value
    if cell.is_date and isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value)


def _error_reason(cell: CellData) -> str:
    if cell.error_text == MISSING_FORMULA_CACHE_ERROR:
        return "formula has no cached result stored in this workbook"
    if cell.is_formula:
        return "cached formula result is an error"
    return "stored value is an error literal"


def _cell_fragment(sheet: SheetData, cell: CellData) -> str | None:
    """Render one cell as a `"Header: value (...)"` text fragment, or `None`
    for a merge's non-anchor cell (no independent value -- see
    `xlsx_wrapper._build_sheet`'s docstring)."""
    if cell.merged_non_anchor:
        return None

    label = sheet.headers.get(cell.column, f"Column {get_column_letter(cell.column)}")

    if cell.is_error:
        # Doc 02 section 4: reject/flag formula-result errors rather than
        # silently reporting them as a valid value. The error literal is
        # kept front-and-center in the text (never dropped or replaced by a
        # guess) precisely so it can never be mistaken for a usable number.
        detail = f"{label}: [UNRESOLVED -- {cell.error_text}] ({_error_reason(cell)}"
        if cell.formula_text:
            detail += f"; formula: {cell.formula_text}"
        detail += ")"
        return detail

    fragment = f"{label}: {_format_raw_value(cell)}"

    annotations: list[str] = []
    if cell.is_formula:
        annotations.append(
            f"formula: {cell.formula_text}, cached result as of this workbook version"
        )
    if cell.is_percentage and isinstance(cell.value, (int, float)):
        annotations.append(f"~{cell.value * 100:g}% per cell formatting")
    if cell.scaling_hint:
        annotations.append(f"display scale: {cell.scaling_hint}")
    if cell.merged_range:
        annotations.append(f"merged {cell.merged_range}")

    if annotations:
        fragment += " (" + "; ".join(annotations) + ")"
    return fragment


def _row_text(sheet: SheetData, row: RowData) -> str:
    header_line = f"Sheet: {sheet.name} | Row: {row.row}"
    notes: list[str] = []
    if row.hidden:
        notes.append("hidden row")
    if sheet.filtered:
        notes.append("sheet has an active filter")
    if notes:
        header_line += " [" + "; ".join(notes) + "]"

    fragments = [f for f in (_cell_fragment(sheet, c) for c in row.cells) if f]
    if not fragments:
        return ""
    return header_line + "\n" + "\n".join(fragments)


def _row_locator(sheet: SheetData, row: RowData) -> dict:
    columns = [c.column for c in row.cells]
    min_col, max_col = min(columns), max(columns)
    start = f"{get_column_letter(min_col)}{row.row}"
    end = f"{get_column_letter(max_col)}{row.row}"

    locator: dict = {"sheet": sheet.name, "range": f"{start}:{end}"}
    if row.hidden:
        locator["hidden"] = True
    if sheet.filtered:
        locator["filtered"] = True

    error_cells = [c.coordinate for c in row.cells if c.is_error]
    if error_cells:
        locator["error_cells"] = error_cells

    merged_ranges = sorted({c.merged_range for c in row.cells if c.merged_range})
    if merged_ranges:
        locator["merged_ranges"] = merged_ranges

    return locator


def _split_row_text(text: str, counter: TokenCounter, cap: int) -> list[str]:
    """`text` unchanged when within `cap` tokens; otherwise split on
    `Header: value` line boundaries (then sentence/word, only for a single
    over-long value), repeating the leading `Sheet: ... | Row: ...` line in
    every part."""
    if counter.count(text) <= cap:
        return [text]
    header_line, _, body = text.partition("\n")
    return [f"{header_line}\n{part}" for part in split_to_fit(body, counter, cap, header_line)]


def chunk_workbook(
    workbook: ParsedWorkbook,
    counter: TokenCounter | None = None,
    max_tokens: int | None = None,
) -> tuple[list[EvidenceUnitDraft], list[ChunkDraft]]:
    """Build one `EvidenceUnitDraft` + one `ChunkDraft` per populated row
    across every sheet in `workbook`, in sheet order then row order.

    Page provenance doesn't apply to spreadsheets (no concept of a "page" in
    a workbook) -- `page_start`/`page_end` are left `None` throughout,
    exactly like a Docling document with no page markers.

    A row over `max_tokens` (default `settings.chunk_max_tokens`) becomes
    several unit/chunk pairs sharing the row's locator plus
    `{"part": i, "of": n}`; a row within the cap is emitted unchanged.
    """
    counter = counter or get_token_counter()
    cap = max_tokens or settings.chunk_max_tokens
    units: list[EvidenceUnitDraft] = []
    chunks: list[ChunkDraft] = []
    ordinal = 0
    # Sheet/period context (doc 05 step 3): indexed text only, never `Chunk.text`.
    sheet_contexts = (
        build_sheet_contexts(workbook) if settings.xlsx_period_context_enabled else {}
    )

    for sheet in workbook.sheets:
        for row in sheet.rows:
            text = _row_text(sheet, row)
            if not text:
                continue

            locator = _row_locator(sheet, row)
            if sheet_contexts:
                context = row_context(sheet_contexts.get(sheet.name, []), sheet, row)
                if context:
                    locator["context"] = context
            parts = _split_row_text(text, counter, cap)

            for part_index, part_text in enumerate(parts):
                unit_index = len(units)
                content_hash = _sha256_hex(part_text)
                part_loc = part_locator(locator, part_index, len(parts))

                units.append(
                    EvidenceUnitDraft(
                        unit_index=unit_index,
                        heading=sheet.name,
                        text=part_text,
                        content_hash=content_hash,
                        unit_kind="range",
                        locator_json=json.dumps(part_loc, sort_keys=True),
                    )
                )
                chunks.append(
                    ChunkDraft(
                        evidence_unit_index=unit_index,
                        ordinal=ordinal,
                        heading=sheet.name,
                        text=part_text,
                        content_hash=content_hash,
                        provenance="extracted",
                    )
                )
                ordinal += 1

    return units, chunks
