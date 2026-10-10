"""Thin wrapper around python-pptx used by ingestion for native `.pptx`
parsing. Mirrors `xlsx_wrapper.py`'s shape: callers get a plain
`ParsedPresentation` back, or a `ParseError` they can catch to skip a single
bad file without crashing a batch ingestion run. See Upgrade doc 02 section 3
(the PowerPoint extraction route) and section 7 (required PowerPoint
validation cases) for the requirements this module implements.

Scope for this checkpoint (native extraction only, no rendering/vision):

- Visible slide text, from any shape with a text frame (text boxes,
  placeholders, autoshapes with text) -- doc 02 section 3: "Native
  text/table/chart extraction".
- Tables on slides, preserved as a row/column grid rather than flattened
  (doc 02 section 5/6: "Do not flatten a table into unrelated numbers" --
  the same principle `xlsx_chunker.py` applies to spreadsheet rows).
- Embedded chart *data* via python-pptx's chart API (categories/series/
  values) -- doc 02's explicit reference link
  (https://python-pptx.readthedocs.io/en/stable/api/chart.html). Only data
  python-pptx exposes natively; a chart rendered as a picture (no
  `GraphicFrame.chart`) is NOT a chart this module can read data from -- see
  `UnextractedShapeData`.
- Speaker notes, kept on `SlideData.notes_text`, structurally separate from
  `SlideData.text_shapes` -- doc 02 section 3: "distinguish speaker notes
  from visible slide text". Never merged into the same field or unit.

Explicitly NOT implemented here (see the checkpoint brief): rendering slides
to images for OCR/vision on image-only content, legacy `.ppt`, animations/
transitions, embedded video/audio. A shape that carries none of
text/table/chart content (most commonly a `PICTURE` -- which is also what an
"image-only chart" looks like to python-pptx, since it has no `.chart`
object to read data from) is recorded on `SlideData.unextracted_shapes`
rather than silently dropped, so a slide whose only content is such a shape
produces an honest empty result instead of masquerading as fully extracted
(doc 02 section 3: "must produce explicit coverage/errors instead of empty
successful results").
"""

from __future__ import annotations

from dataclasses import dataclass
from importlib.metadata import version
from pathlib import Path

from docket.infra.parsing.filecheck import check_file_signature


@dataclass
class ChartSeriesData:
    name: str | None
    # One value per category, in category order. `None` for a data point
    # python-pptx itself reports as missing -- never coerced to 0, mirroring
    # xlsx_wrapper's "a missing cell is not zero" treatment of blank cells.
    values: list[float | None]


@dataclass
class ChartData:
    shape_id: int
    shape_name: str
    # `XL_CHART_TYPE` member name (e.g. "COLUMN_CLUSTERED"), or `None` if
    # python-pptx couldn't resolve the chart's type from the XML.
    chart_type: str | None
    title: str | None
    categories: list[str]
    series: list[ChartSeriesData]


@dataclass
class TableData:
    shape_id: int
    shape_name: str
    # Full grid of cell text, row-major, INCLUDING the header row (if any) --
    # `header_row` below is a convenience view onto `rows[0]`, not a separate
    # copy of different data.
    rows: list[list[str]]
    header_row: list[str] | None  # None when the table has 0 or 1 rows
    n_rows: int
    n_cols: int


@dataclass
class TextShapeData:
    shape_id: int
    shape_name: str
    text: str  # full text-frame text (paragraphs newline-joined); never blank


@dataclass
class UnextractedShapeData:
    """A shape structurally present on the slide but carrying none of
    text/table/chart content this checkpoint extracts -- most commonly a
    `PICTURE` (doc 02's "image-only charts" validation case: a screenshot of
    a chart has no `.chart` object for python-pptx to read data from).
    Recorded as a coverage signal, never turned into evidence text -- doing
    so would risk inventing content that was never actually extracted.
    """

    shape_id: int
    shape_name: str
    shape_type: str


@dataclass
class SlideData:
    slide_number: int  # 1-indexed, matches how a user would refer to a slide
    text_shapes: list[TextShapeData]
    tables: list[TableData]
    charts: list[ChartData]
    # None when the slide has no notes slide at all, or its notes text frame
    # is empty/whitespace-only -- distinguished from "" so an empty-notes
    # slide never produces a spurious empty notes unit downstream.
    notes_text: str | None
    unextracted_shapes: list[UnextractedShapeData]


@dataclass
class ParsedPresentation:
    slides: list[SlideData]
    source_path: Path
    parser_name: str
    parser_version: str


class ParseError(Exception):
    """Raised for any failure parsing ``path`` for ``source_id``.

    Wraps the underlying python-pptx (or filesystem) exception as ``cause``
    so callers never need to catch python-pptx's own exception types
    directly -- this also covers corrupt/non-OPC files, which python-pptx
    cannot open at all. Mirrors `xlsx_wrapper.ParseError`.
    """

    def __init__(self, source_id: str, path: Path, cause: Exception):
        self.source_id = source_id
        self.path = path
        self.cause = cause
        super().__init__(f"failed to parse {path} for source {source_id}: {cause}")


class UnsupportedPresentationFormatError(Exception):
    """Raised for `.ppt` (legacy binary PowerPoint format).

    python-pptx only reads the OOXML `.pptx`/`.potx` zip container; it
    cannot open the legacy binary format at all. Gets an explicit, named
    rejection rather than an untested silent attempt -- doc 02 section 3:
    "Unsupported ... inputs must produce explicit coverage/errors instead of
    empty successful results." Mirrors
    `xlsx_wrapper.UnsupportedSpreadsheetFormatError`.
    """

    def __init__(self, path: Path):
        self.path = path
        super().__init__(
            f"unsupported presentation format for {path.suffix}: only .pptx "
            "is implemented in this checkpoint (legacy .ppt is not supported "
            "-- python-pptx cannot read the binary PowerPoint format at all)"
        )


def _shape_text(shape) -> str | None:
    if not shape.has_text_frame:
        return None
    text = shape.text_frame.text
    if not text or not text.strip():
        return None
    return text


def _table_data(shape) -> TableData:
    table = shape.table
    rows: list[list[str]] = [[cell.text for cell in row.cells] for row in table.rows]
    n_rows = len(rows)
    n_cols = len(rows[0]) if rows else 0
    # Mirrors xlsx_wrapper's header-row convention: the first row is treated
    # as the header when there's at least one more row to be data. A
    # single-row table has nothing to be "header for", so it's left
    # unset -- its one row is then rendered as data in pptx_chunker with
    # generic "Column N" labels rather than consuming its only row as a
    # header with no data left to attach it to.
    header_row = rows[0] if n_rows > 1 else None
    return TableData(
        shape_id=shape.shape_id,
        shape_name=shape.name,
        rows=rows,
        header_row=header_row,
        n_rows=n_rows,
        n_cols=n_cols,
    )


def _chart_data(shape) -> ChartData:
    chart = shape.chart

    try:
        chart_type = chart.chart_type.name if chart.chart_type is not None else None
    except Exception:  # noqa: BLE001 -- an unrecognized chart-type XML value must not crash parsing
        chart_type = None

    title: str | None = None
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            title_text = chart.chart_title.text_frame.text
            title = title_text.strip() if title_text and title_text.strip() else None
    except Exception:  # noqa: BLE001
        title = None

    categories: list[str] = []
    try:
        plots = chart.plots
        if plots:
            categories = [str(c) if c is not None else "" for c in plots[0].categories]
    except Exception:  # noqa: BLE001
        categories = []

    series_list: list[ChartSeriesData] = []
    for series in chart.series:
        values = [v if v is None else float(v) for v in series.values]
        series_list.append(ChartSeriesData(name=series.name, values=values))

    return ChartData(
        shape_id=shape.shape_id,
        shape_name=shape.name,
        chart_type=chart_type,
        title=title,
        categories=categories,
        series=series_list,
    )


def _iter_shapes(shapes):
    """Recursively walk a shape collection, descending into `GROUP` shapes.

    A table/chart/text box nested inside a group (common when an author
    groups slide elements for layout) must still be discovered -- grouping
    is a layout convenience, not a reason for its content to become
    invisible to extraction.
    """
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _build_slide(slide_number: int, slide) -> SlideData:
    text_shapes: list[TextShapeData] = []
    tables: list[TableData] = []
    charts: list[ChartData] = []
    unextracted: list[UnextractedShapeData] = []

    for shape in _iter_shapes(slide.shapes):
        if shape.has_chart:
            charts.append(_chart_data(shape))
            continue
        if shape.has_table:
            tables.append(_table_data(shape))
            continue
        if shape.has_text_frame:
            text = _shape_text(shape)
            if text is not None:
                text_shapes.append(
                    TextShapeData(shape_id=shape.shape_id, shape_name=shape.name, text=text)
                )
            continue
        # Not natively extractable -- most commonly a picture. See
        # `UnextractedShapeData`'s docstring for why this is recorded rather
        # than dropped.
        shape_type = shape.shape_type
        unextracted.append(
            UnextractedShapeData(
                shape_id=shape.shape_id,
                shape_name=shape.name,
                shape_type=shape_type.name if shape_type is not None else "UNKNOWN",
            )
        )

    notes_text: str | None = None
    if slide.has_notes_slide:
        raw_notes = slide.notes_slide.notes_text_frame.text
        if raw_notes and raw_notes.strip():
            notes_text = raw_notes

    return SlideData(
        slide_number=slide_number,
        text_shapes=text_shapes,
        tables=tables,
        charts=charts,
        notes_text=notes_text,
        unextracted_shapes=unextracted,
    )


class PptxParser:
    """Parses `.pptx` presentations via python-pptx into a plain
    `ParsedPresentation`.

    Cheap to construct (no model loading, unlike `DoclingParser`) -- a fresh
    instance per call is fine, but the ingestion pipeline still holds one
    instance for parity with how `XlsxParser` is wired in.
    """

    def __init__(self) -> None:
        self._parser_name = "python-pptx"
        self._parser_version = self._detect_version()

    @staticmethod
    def _detect_version() -> str:
        return version("python-pptx")

    @property
    def parser_name(self) -> str:
        return self._parser_name

    @property
    def parser_version(self) -> str:
        return self._parser_version

    def parse(self, source_id: str, path: Path) -> ParsedPresentation:
        """Load `path` and build a `ParsedPresentation`. Any failure
        (corrupt file, not actually a zip/pptx at all) is caught and
        re-raised as `ParseError` -- never propagated raw, matching
        `XlsxParser`."""
        from pptx import Presentation

        if path.suffix.lower() != ".pptx":
            # Defense in depth: `docket.services.ingestion.pipeline` already
            # routes `.ppt` to an explicit per-file failure before ever
            # calling this parser, but a direct caller (script, future CLI
            # command, test) bypassing that dispatch still gets a named
            # rejection instead of python-pptx failing with an unrelated
            # low-level error (or, for a `.ppt` that happens to not even be
            # a zip, an opaque one).
            raise UnsupportedPresentationFormatError(path)

        try:
            check_file_signature(path)  # cheap "not really a pptx" rejection
            presentation = Presentation(str(path))
        except Exception as exc:  # noqa: BLE001 -- intentionally broad, see ParseError docstring
            raise ParseError(source_id, path, exc) from exc

        try:
            slides = [
                _build_slide(i, slide) for i, slide in enumerate(presentation.slides, start=1)
            ]
        except Exception as exc:  # noqa: BLE001
            raise ParseError(source_id, path, exc) from exc

        return ParsedPresentation(
            slides=slides,
            source_path=path,
            parser_name=self._parser_name,
            parser_version=self._parser_version,
        )
