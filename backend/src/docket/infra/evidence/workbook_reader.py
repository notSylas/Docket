"""Deterministic spreadsheet range reader (Upgrade doc 05 section 8).

Backs the agent's `read_range` and `calculate` tools. It answers "what exactly
is in these cells of the workbook this chunk came from", from the ORIGINAL
stored bytes, never from a search snippet and never from a model's memory:

* The chunk id only PINS which evidence version (workbook) is read. Only a READY
  version of an ACTIVE source is readable (the same eligibility the resolver
  and `hybrid_search` apply); a revoked, superseded, pending or unknown chunk
  is refused with a clear error.
* Bytes come from the content-addressed store (`ContentAddressedStore.get`),
  so the live file may have changed or been deleted.
* openpyxl loads the workbook twice (cached values, formula text), exactly as
  ingestion does (`xlsx_wrapper`). Formulas are NEVER evaluated and macros are
  never run (`keep_vba` stays False, the bytes are only parsed).
* Header labels come from the chunker's own header detection
  (`xlsx_wrapper._build_sheet`), so a merged title row that the chunker took as
  the header is reported as such instead of being given invented labels.
* A blank cell is `null` with `blank: true`; a formula whose cached result is
  missing is `cached_value_missing: true` (value null) -- neither is ever 0.
* Output size is bounded: at most `MAX_CELLS` cells and `MAX_ROWS` rows per call
  (tool use); the coverage object says exactly what was and was not returned.

Provenance: everything returned here is `extracted` (what the workbook states,
including a formula's cached result). Derived numbers come from
`docket.services.agent.calculator`, which resolves its inputs through this
same reader.
"""

from __future__ import annotations

import io
import json
import re
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl.utils import column_index_from_string, get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import sessionmaker

from docket.core.db.models import (
    Chunk,
    EvidenceUnit,
    EvidenceVersion,
    Source,
    SourceStatus,
    VersionStatus,
)
from docket.infra.evidence.store import ContentAddressedStore, ObjectNotFoundError
from docket.infra.parsing.xlsx_context import sheet_scope_lines
from docket.infra.parsing.xlsx_wrapper import (
    MISSING_FORMULA_CACHE_ERROR,
    ParsedWorkbook,
    SheetData,
    _build_sheet,
    _cell_value_and_format,
    _merge_lookup,
    _scaling_hint,
)
from docket.infra.retrieval.resolver import _citation_label

MAX_CELLS = 200
MAX_ROWS = 60
_MAX_COLUMN = 16384  # XFD
_MAX_ROW = 1048576
_BUNDLE_CACHE_SIZE = 4

NOT_STATED = "not stated"
PERCENT = "percent"
AMBIGUOUS_PREFIX = "ambiguous"


class WorkbookReadError(Exception):
    """A clear, model-facing reason a read could not be done."""


# -- A1 range parsing -----------------------------------------------------

_CELL = r"\$?([A-Za-z]{1,3})\$?([0-9]+)"
_RANGE_RE = re.compile(rf"^\s*{_CELL}(?:\s*:\s*{_CELL})?\s*$")


@dataclass(frozen=True)
class A1Range:
    min_col: int
    min_row: int
    max_col: int
    max_row: int

    @property
    def text(self) -> str:
        start = f"{get_column_letter(self.min_col)}{self.min_row}"
        end = f"{get_column_letter(self.max_col)}{self.max_row}"
        return start if start == end else f"{start}:{end}"

    @property
    def n_columns(self) -> int:
        return self.max_col - self.min_col + 1

    @property
    def n_rows(self) -> int:
        return self.max_row - self.min_row + 1


def parse_a1_range(text: Any) -> A1Range:
    """`B5`, `A1:F12` (also `$B$5:$F$12`); a reversed range is normalised.
    Whole-column (`A:A`), whole-row (`3:3`), open-ended and multi-area forms
    are rejected with a clear error."""
    if not isinstance(text, str) or not text.strip():
        raise WorkbookReadError("range is required, e.g. 'B5:F5' or 'C7'")
    match = _RANGE_RE.match(text)
    if match is None:
        raise WorkbookReadError(
            f"invalid range '{text}': use A1 notation for a single cell ('C7') or a "
            "rectangle ('A1:F12'); whole-column ('A:A'), whole-row ('3:3') and "
            "multi-area ranges are not supported"
        )
    c1, r1, c2, r2 = match.groups()
    if c2 is None:
        c2, r2 = c1, r1
    try:
        cols = [column_index_from_string(c.upper()) for c in (c1, c2)]
    except ValueError as exc:
        raise WorkbookReadError(f"invalid column in range '{text}'") from exc
    rows = [int(r1), int(r2)]
    if min(rows) < 1 or max(rows) > _MAX_ROW or max(cols) > _MAX_COLUMN:
        raise WorkbookReadError(f"range '{text}' is outside the valid cell grid")
    return A1Range(min(cols), min(rows), max(cols), max(rows))


# -- units ------------------------------------------------------------------

_CURRENCY_WORDS = {
    "inr": "INR", "rs": "INR", "rupee": "INR", "rupees": "INR", "₹": "INR",
    "usd": "USD", "dollar": "USD", "dollars": "USD", "$": "USD",
    "eur": "EUR", "euro": "EUR", "euros": "EUR", "€": "EUR",
    "gbp": "GBP", "pound": "GBP", "pounds": "GBP", "£": "GBP",
}
_SCALE_WORDS = {
    "thousand": "thousands", "thousands": "thousands",
    "million": "millions", "millions": "millions",
    "billion": "billions", "billions": "billions",
    "lakh": "lakhs", "lakhs": "lakhs",
    "crore": "crores", "crores": "crores",
}
_WORD_RE = re.compile(r"[A-Za-z]+|[₹$€£]")
_PAREN_RE = re.compile(r"\(([^)]*)\)")


def parse_unit_text(text: str) -> tuple[set[str], set[str]]:
    """(currencies, scales) a piece of text states, e.g. `INR thousands`."""
    currencies: set[str] = set()
    scales: set[str] = set()
    for word in _WORD_RE.findall(text):
        key = word.lower()
        if key in _CURRENCY_WORDS:
            currencies.add(_CURRENCY_WORDS[key])
        elif key in _SCALE_WORDS:
            scales.add(_SCALE_WORDS[key])
    return currencies, scales


def _format_unit(currencies: set[str], scales: set[str]) -> str:
    if len(currencies) > 1 or len(scales) > 1:
        parts = sorted(currencies) + sorted(scales)
        return f"{AMBIGUOUS_PREFIX}: {', '.join(parts)}"
    parts = sorted(currencies) + sorted(scales)
    return " ".join(parts) if parts else NOT_STATED


def context_units(lines: list[str]) -> str:
    """The unit the sheet/workbook context lines state ("INR thousands"), or
    `not stated`, or `ambiguous: ...` when they name more than one currency or
    scale. Never guessed from the numbers."""
    currencies: set[str] = set()
    scales: set[str] = set()
    for line in lines:
        c, s = parse_unit_text(line)
        currencies |= c
        scales |= s
    return _format_unit(currencies, scales)


def cell_units(header: str | None, is_percent: bool, sheet_units: str) -> str:
    """Unit of one cell: percent-formatted cells are `percent`; a currency/scale
    in parentheses in the column header ("Total (INR thousands)") wins over the
    sheet-level statement, which fills in whatever the header leaves out."""
    if is_percent or (header and "%" in header):
        return PERCENT
    if header:
        h_cur: set[str] = set()
        h_scale: set[str] = set()
        for inner in _PAREN_RE.findall(header):
            c, s = parse_unit_text(inner)
            h_cur |= c
            h_scale |= s
        if h_cur or h_scale:
            if sheet_units.startswith(AMBIGUOUS_PREFIX) or sheet_units == NOT_STATED:
                return _format_unit(h_cur, h_scale)
            s_cur, s_scale = parse_unit_text(sheet_units)
            return _format_unit(h_cur or s_cur, h_scale or s_scale)
    return sheet_units


# -- data model ---------------------------------------------------------------


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


@dataclass
class CellRead:
    address: str
    row: int
    column: int
    value: Any
    header: str | None
    number_format: str | None
    blank: bool
    is_formula: bool
    formula_text: str | None
    cached_value_missing: bool
    error: str | None
    hidden_row: bool
    hidden_column: bool
    is_percent: bool
    units: str
    merged_range: str | None = None
    merged_into: str | None = None
    display_scale: str | None = None

    @property
    def is_number(self) -> bool:
        return (
            isinstance(self.value, (int, float))
            and not isinstance(self.value, bool)
            and not self.cached_value_missing
            and self.error is None
        )

    def to_json(self, sheet_units: str) -> dict[str, Any]:
        if self.blank and not self.is_formula:
            # Compact: a blank cell carries no header/format/units noise.
            out: dict[str, Any] = {"address": self.address, "value": None, "blank": True}
            if self.merged_into:
                out["merged_into"] = self.merged_into
            if self.hidden_column:
                out["hidden_column"] = True
            return out
        out = {
            "address": self.address,
            "value": _json_value(self.value),
            "header": self.header,
            "formula": self.is_formula,
        }
        if self.number_format and self.number_format != "General":
            out["number_format"] = self.number_format
        if self.is_formula:
            out["formula_text"] = self.formula_text
        if self.cached_value_missing:
            out["cached_value_missing"] = True
            out["note"] = "no cached value"
        if self.error and not self.cached_value_missing:
            out["error"] = self.error
        if self.is_percent and isinstance(self.value, (int, float)):
            out["percent"] = True
        if self.hidden_column:
            out["hidden_column"] = True
        if self.merged_range:
            out["merged_range"] = self.merged_range
        if self.merged_into:
            out["merged_into"] = self.merged_into
        if self.display_scale:
            out["display_scale"] = self.display_scale
        if self.units != sheet_units:
            out["units"] = self.units
        return out


@dataclass
class RowRead:
    row: int
    hidden: bool
    cells: list[CellRead]


@dataclass
class RangeRead:
    chunk_id: str
    source: str
    file_path: str | None
    evidence_version_id: str
    sheet: str
    requested_range: str
    returned_range: str
    headers: list[dict[str, Any]]
    header_note: str | None
    rows: list[RowRead]
    context: dict[str, list[str]]
    units: str
    coverage: dict[str, Any]
    chunk_ids: list[str] = field(default_factory=list)
    citation_labels: list[str] = field(default_factory=list)

    @property
    def cells(self) -> list[CellRead]:
        return [c for r in self.rows for c in r.cells]

    def to_json(self) -> dict[str, Any]:
        out: dict[str, Any] = {
            "source": self.source,
            "evidence_version_id": self.evidence_version_id,
            "sheet": self.sheet,
            "range": self.requested_range,
            "returned_range": self.returned_range,
            "provenance": "extracted",
            "headers": self.headers,
        }
        if self.header_note:
            out["header_note"] = self.header_note
        out["context"] = self.context
        out["units"] = self.units
        out["rows"] = [
            {"row": r.row, "hidden": r.hidden, "cells": [c.to_json(self.units) for c in r.cells]}
            for r in self.rows
        ]
        out["coverage"] = self.coverage
        out["chunk_ids"] = self.chunk_ids
        out["citation_labels"] = self.citation_labels
        return out


# -- workbook bundle (parsed once per stored object) ---------------------------


@dataclass
class _Bundle:
    parsed: ParsedWorkbook
    value_sheets: dict[str, Any]
    formula_sheets: dict[str, Any]
    sheet_data: dict[str, SheetData]
    merge_lookups: dict[str, dict]
    hidden_columns: dict[str, set[int]]


def _load_bundle(data: bytes, file_name: str) -> _Bundle:
    import openpyxl

    try:
        # Two loads (cached values / formula text); formulas are never
        # evaluated, macros never executed (keep_vba defaults to False).
        value_wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
        formula_wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, read_only=False)
    except Exception as exc:  # noqa: BLE001 -- not a readable workbook
        raise WorkbookReadError(
            f"the stored file '{file_name}' could not be opened as an .xlsx workbook ({exc})"
        ) from exc
    names = list(value_wb.sheetnames)
    sheet_data = {n: _build_sheet(n, value_wb[n], formula_wb[n]) for n in names}
    hidden_cols: dict[str, set[int]] = {}
    for n in names:
        cols: set[int] = set()
        for dim in value_wb[n].column_dimensions.values():
            if dim.hidden and dim.min and dim.max:
                cols.update(range(dim.min, min(dim.max, dim.min + 2000) + 1))
        hidden_cols[n] = cols
    return _Bundle(
        parsed=ParsedWorkbook(
            sheets=[sheet_data[n] for n in names],
            source_path=Path(file_name),
            parser_name="openpyxl",
            parser_version=openpyxl.__version__,
        ),
        value_sheets={n: value_wb[n] for n in names},
        formula_sheets={n: formula_wb[n] for n in names},
        sheet_data=sheet_data,
        merge_lookups={n: _merge_lookup(value_wb[n]) for n in names},
        hidden_columns=hidden_cols,
    )


@dataclass
class _Pinned:
    chunk_id: str
    evidence_version_id: str
    content_hash: str
    file_path: str | None
    source_name: str
    unit_sheet: str | None


class WorkbookReader:
    """See the module docstring. `session_factory` and `store` are the app's
    existing database sessions and content-addressed blob store."""

    def __init__(self, *, session_factory: sessionmaker, store: ContentAddressedStore):
        self._session_factory = session_factory
        self._store = store
        self._bundles: OrderedDict[tuple[str, str], _Bundle] = OrderedDict()

    # -- lookup ---------------------------------------------------------------

    def _pin(self, chunk_id: str) -> _Pinned:
        if not isinstance(chunk_id, str) or not chunk_id:
            raise WorkbookReadError("chunk_id is required (use one returned by search_knowledge)")
        with self._session_factory() as session:
            row = session.execute(
                select(Chunk, Source, EvidenceVersion, EvidenceUnit)
                .join(Source, Chunk.source_id == Source.id)
                .join(EvidenceVersion, Chunk.evidence_version_id == EvidenceVersion.id)
                .join(EvidenceUnit, Chunk.evidence_unit_id == EvidenceUnit.id)
                .where(Chunk.id == chunk_id)
                .where(Source.status == SourceStatus.ACTIVE)
                .where(EvidenceVersion.status == VersionStatus.READY)
            ).first()
        if row is None:
            raise WorkbookReadError(
                f"chunk not found or no longer available: {chunk_id} (it may be unknown, "
                "revoked, or from a superseded version); search again"
            )
        chunk, source, version, unit = row
        try:
            locator = json.loads(unit.locator_json) if unit.locator_json else {}
        except ValueError:
            locator = {}
        sheet = locator.get("sheet") if isinstance(locator, dict) else None
        display = version.file_path or source.path
        return _Pinned(
            chunk_id=chunk.id,
            evidence_version_id=version.id,
            content_hash=version.content_hash,
            file_path=version.file_path,
            source_name=Path(display).name,
            unit_sheet=sheet if isinstance(sheet, str) else None,
        )

    def _bundle(self, pinned: _Pinned) -> _Bundle:
        key = (pinned.content_hash, pinned.source_name)
        cached = self._bundles.get(key)
        if cached is not None:
            self._bundles.move_to_end(key)
            return cached
        if pinned.unit_sheet is None and not pinned.source_name.lower().endswith(".xlsx"):
            raise WorkbookReadError(
                f"chunk {pinned.chunk_id} is not from a spreadsheet ({pinned.source_name}); "
                "read_range only works on .xlsx workbooks"
            )
        try:
            data = self._store.get(pinned.content_hash)
        except ObjectNotFoundError as exc:
            raise WorkbookReadError(
                f"the stored copy of {pinned.source_name} is no longer available"
            ) from exc
        bundle = _load_bundle(data, pinned.source_name)
        self._bundles[key] = bundle
        while len(self._bundles) > _BUNDLE_CACHE_SIZE:
            self._bundles.popitem(last=False)
        return bundle

    def _overlapping_chunks(
        self, pinned: _Pinned, sheet: str, area: A1Range, source_name: str
    ) -> tuple[list[str], list[str]]:
        with self._session_factory() as session:
            rows = session.execute(
                select(Chunk.id, EvidenceUnit.locator_json)
                .join(EvidenceUnit, Chunk.evidence_unit_id == EvidenceUnit.id)
                .where(Chunk.evidence_version_id == pinned.evidence_version_id)
                .where(EvidenceUnit.unit_kind == "range")
                .order_by(Chunk.ordinal)
            ).all()
        found: list[tuple[int, int, str]] = []
        for chunk_id, locator_json in rows:
            try:
                locator = json.loads(locator_json) if locator_json else {}
                if locator.get("sheet") != sheet or not locator.get("range"):
                    continue
                span = parse_a1_range(locator["range"])
            except (ValueError, AttributeError, WorkbookReadError):
                continue
            if (
                span.max_row < area.min_row or span.min_row > area.max_row
                or span.max_col < area.min_col or span.min_col > area.max_col
            ):
                continue
            found.append((span.min_row, len(found), chunk_id))
        found.sort()
        ids = [c for _, _, c in found]
        return ids, [_citation_label(source_name, c) for c in ids]

    # -- reading --------------------------------------------------------------

    @staticmethod
    def _resolve_sheet(bundle: _Bundle, wanted: str) -> str:
        names = list(bundle.sheet_data)
        if wanted in bundle.sheet_data:
            return wanted
        folded = [n for n in names if n.strip().lower() == wanted.strip().lower()]
        if len(folded) == 1:
            return folded[0]
        raise WorkbookReadError(f"sheet '{wanted}' not found; this workbook has: {names}")

    def read(
        self,
        chunk_id: str,
        range_text: str,
        sheet: str | None = None,
        *,
        max_cells: int = MAX_CELLS,
        max_rows: int = MAX_ROWS,
        truncate: bool = True,
    ) -> RangeRead:
        """Read `range_text` from the workbook `chunk_id` belongs to. With
        `truncate=False` (arithmetic) a range over the caps is an error instead
        of being cut short."""
        area = parse_a1_range(range_text)
        if area.n_columns > max_cells:
            raise WorkbookReadError(
                f"range '{range_text}' is {area.n_columns} columns wide; at most {max_cells} cells per call"
            )
        pinned = self._pin(chunk_id)
        bundle = self._bundle(pinned)
        wanted = sheet or pinned.unit_sheet
        if not wanted:
            raise WorkbookReadError("sheet is required: this chunk does not name a sheet")
        sheet_name = self._resolve_sheet(bundle, wanted)

        sd = bundle.sheet_data[sheet_name]
        value_ws = bundle.value_sheets[sheet_name]

        # Rows below the sheet's last used row are not data: a generous request
        # such as A3:F100 on a 6-row sheet returns the real rows, not ~90 blank
        # ones (those swamped the model in the first real-stack run). A range
        # wholly outside the used area returns just its first row (blank cells)
        # so a single-cell question about an empty cell still gets an answer.
        used_rows = value_ws.max_row or 0
        used_range = (
            A1Range(1, 1, max(value_ws.max_column or 1, 1), used_rows).text if used_rows else None
        )
        cells_requested = area.n_rows * area.n_columns
        eff_max_row = area.max_row
        clipped = False
        if used_rows and area.min_row <= used_rows < area.max_row:
            eff_max_row, clipped = used_rows, True
        elif area.min_row > used_rows and area.n_rows > 1:
            eff_max_row, clipped = area.min_row, True
        eff_rows = eff_max_row - area.min_row + 1
        allowed_rows = min(max_rows, max_cells // area.n_columns)
        if eff_rows > allowed_rows and not truncate:
            raise WorkbookReadError(
                f"range '{range_text}' has {eff_rows * area.n_columns} cells in {eff_rows} rows; "
                f"at most {max_cells} cells / {max_rows} rows are allowed here (narrow the range)"
            )
        rows_returned = min(eff_rows, allowed_rows)
        last_row = area.min_row + rows_returned - 1
        returned = A1Range(area.min_col, area.min_row, area.max_col, last_row)

        formula_ws = bundle.formula_sheets[sheet_name]
        merges = bundle.merge_lookups[sheet_name]
        hidden_cols = bundle.hidden_columns[sheet_name]

        merged_title = len(sd.headers) > 1 and len(set(sd.headers.values())) == 1
        header_note: str | None = None
        if merged_title:
            header_note = (
                f"Row {sd.header_row} is the detected header row but it is a single merged "
                f"title ('{next(iter(sd.headers.values()))}'), not column labels, so no column "
                "labels are given; the real column labels, if any, are in a row below it."
            )

        def header_for(col: int) -> str | None:
            return None if merged_title else sd.headers.get(col)

        context = sheet_scope_lines(bundle.parsed, sheet_name)
        sheet_units = context_units(context["units_and_scope"])

        rows: list[RowRead] = []
        for r in range(returned.min_row, returned.max_row + 1):
            row_dim = value_ws.row_dimensions.get(r)
            hidden_row = bool(row_dim.hidden) if row_dim is not None else False
            cells: list[CellRead] = []
            for c in range(returned.min_col, returned.max_col + 1):
                cells.append(
                    self._read_cell(
                        sheet_name, r, c, value_ws, formula_ws, merges, hidden_row,
                        c in hidden_cols, header_for(c), sheet_units,
                    )
                )
            rows.append(RowRead(row=r, hidden=hidden_row, cells=cells))

        truncated = rows_returned < eff_rows
        coverage: dict[str, Any] = {
            "cells_requested": cells_requested,
            "cells_returned": rows_returned * area.n_columns,
            "rows_in_range": area.n_rows,
            "rows_returned": rows_returned,
            "truncated": truncated,
            "sheet_used_range": used_range,
        }
        if clipped:
            coverage["clipped_to_used_range"] = True
            coverage["clip_message"] = (
                f"Rows after {last_row if not truncated else eff_max_row} are beyond the sheet's "
                f"last used row ({used_rows}) and were not returned."
                if area.min_row <= used_rows
                else f"The range is entirely below the sheet's last used row ({used_rows}); "
                "only its first row is shown."
            )
        if truncated:
            coverage["message"] = (
                f"Range truncated to rows {returned.min_row}-{returned.max_row} "
                f"(limit {max_cells} cells / {max_rows} rows per call); call read_range "
                f"again for rows {last_row + 1}-{eff_max_row}."
            )
        ids, labels = self._overlapping_chunks(pinned, sheet_name, returned, pinned.source_name)
        headers = [
            {"column": get_column_letter(c), "header": header_for(c)}
            for c in range(returned.min_col, returned.max_col + 1)
        ]
        return RangeRead(
            chunk_id=pinned.chunk_id,
            source=pinned.source_name,
            file_path=pinned.file_path,
            evidence_version_id=pinned.evidence_version_id,
            sheet=sheet_name,
            requested_range=area.text,
            returned_range=returned.text,
            headers=headers,
            header_note=header_note,
            rows=rows,
            context=context,
            units=sheet_units,
            coverage=coverage,
            chunk_ids=ids,
            citation_labels=labels,
        )

    @staticmethod
    def _read_cell(
        sheet: str, r: int, c: int, value_ws: Any, formula_ws: Any, merges: dict,
        hidden_row: bool, hidden_col: bool, header: str | None, sheet_units: str,
    ) -> CellRead:
        address = f"{get_column_letter(c)}{r}"
        merge = merges.get((r, c))
        if merge is not None and not merge[1]:
            return CellRead(
                address=address, row=r, column=c, value=None, header=header,
                number_format=None, blank=True, is_formula=False, formula_text=None,
                cached_value_missing=False, error=None, hidden_row=hidden_row,
                hidden_column=hidden_col, is_percent=False,
                units=cell_units(header, False, sheet_units), merged_into=merge[0],
            )
        if r > (value_ws.max_row or 0) or c > (value_ws.max_column or 0):
            return CellRead(
                address=address, row=r, column=c, value=None, header=header,
                number_format=None, blank=True, is_formula=False, formula_text=None,
                cached_value_missing=False, error=None, hidden_row=hidden_row,
                hidden_column=hidden_col, is_percent=False,
                units=cell_units(header, False, sheet_units),
            )
        value_cell = value_ws.cell(row=r, column=c)
        formula_cell = formula_ws.cell(row=r, column=c)
        value, is_formula, formula_text, is_error, error_text = _cell_value_and_format(
            value_cell, formula_cell
        )
        number_format = value_cell.number_format
        is_percent = bool(number_format and "%" in number_format)
        missing = is_error and error_text == MISSING_FORMULA_CACHE_ERROR
        return CellRead(
            address=address, row=r, column=c, value=value, header=header,
            number_format=number_format, blank=value is None and not is_formula,
            is_formula=is_formula, formula_text=formula_text,
            cached_value_missing=missing,
            error=error_text if is_error else None,
            hidden_row=hidden_row, hidden_column=hidden_col, is_percent=is_percent,
            units=cell_units(header, is_percent, sheet_units),
            merged_range=merge[0] if merge is not None else None,
            display_scale=_scaling_hint(number_format),
        )
