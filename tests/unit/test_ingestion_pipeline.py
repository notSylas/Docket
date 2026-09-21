"""Tests for `attest.ingestion.pipeline.IngestionPipeline` -- the end-to-end
incremental ingestion wiring: real Docling parsing (via small generated
`.docx` fixture files, same technique as `test_docling_wrapper.py`'s
real-file smoke test), `FakeInferenceGateway` for embeddings (no Ollama
dependency, so this stays a fast-ish unit test despite exercising real
Docling), and a real migrated SQLite DB + real LanceDB dir (both under
`tmp_path`).

This is the regression suite proving the whole pipeline -- not just
`IndexManager` in isolation (see `test_index_manager.py`) -- correctly skips
re-embedding unchanged files end to end, correctly reconciles stale chunks
when a file's content changes, and correctly isolates one bad file's failure
from the rest of the batch.
"""

from __future__ import annotations

from pathlib import Path
from types import SimpleNamespace

import docx
import pytest
from sqlalchemy import Engine, select

from attest.db.engine import get_session_factory
from attest.db.models import Chunk, EvidenceVersion, IngestionJob, IngestionJobStatus
from attest.evidence.manager import EvidenceManager
from attest.evidence.store import ContentAddressedStore
from attest.index.fts_index import FtsIndexWriter
from attest.index.manager import IndexManager
from attest.index.vector_index import LanceIndexWriter
from attest.inference.gateway import FakeInferenceGateway
from attest.ingestion.pipeline import IngestionPipeline
from attest.parsing.docling_wrapper import DoclingParser
from attest.parsing.recipes import DEFAULT_SPLITTER, ChunkRecipe
from attest.sources.manager import SourceManager


def _write_docx(path: Path, heading: str, body: str) -> None:
    document = docx.Document()
    document.add_heading(heading, level=1)
    document.add_paragraph(body)
    document.save(str(path))


@pytest.fixture(scope="module")
def parser() -> DoclingParser:
    # Constructed once per test module: Docling loads its models on
    # DocumentConverter() construction (expensive the first time).
    return DoclingParser()


@pytest.fixture()
def env(
    migrated_sqlite_engine: Engine, tmp_path: Path, parser: DoclingParser
) -> SimpleNamespace:
    session_factory = get_session_factory(migrated_sqlite_engine)

    store_root = tmp_path / "evidence_store"
    for sub in ("objects", "manifests", "quarantine", "trash"):
        (store_root / sub).mkdir(parents=True, exist_ok=True)
    store = ContentAddressedStore(store_root)
    evidence_manager = EvidenceManager(store, session_factory)

    gateway = FakeInferenceGateway()
    fts = FtsIndexWriter(migrated_sqlite_engine)
    vector = LanceIndexWriter(tmp_path / "lancedb")
    index_manager = IndexManager(fts, vector, gateway)

    chunk_recipe = ChunkRecipe(
        chunk_size=200,
        overlap=40,
        splitter=DEFAULT_SPLITTER,
        parser_name=parser.parser_name,
        parser_version=parser.parser_version,
    )

    pipeline = IngestionPipeline(
        session_factory=session_factory,
        evidence_manager=evidence_manager,
        parser=parser,
        index_manager=index_manager,
        chunk_recipe=chunk_recipe,
    )

    source_manager = SourceManager(session_factory)
    folder = tmp_path / "docs"
    folder.mkdir()
    source = source_manager.register_source(folder)

    return SimpleNamespace(
        session_factory=session_factory,
        gateway=gateway,
        fts=fts,
        vector=vector,
        pipeline=pipeline,
        source_manager=source_manager,
        source=source,
        folder=folder,
    )


def _chunk_ids_for_file(env: SimpleNamespace, file_path: Path) -> set[str]:
    """Chunk ids belonging to `file_path`'s *current* EvidenceVersion only
    -- a file's earlier (superseded) version keeps its own old Chunk rows
    around in SQLite for audit purposes even after reconciliation removes
    them from the indexes, so this must not match those too."""
    with env.session_factory() as session:
        rows = session.execute(
            select(Chunk.id)
            .join(EvidenceVersion, Chunk.evidence_version_id == EvidenceVersion.id)
            .where(
                EvidenceVersion.file_path == str(file_path),
                EvidenceVersion.is_current.is_(True),
            )
        ).scalars().all()
        return set(rows)


def _job_status(env: SimpleNamespace, job_id: str) -> IngestionJobStatus:
    with env.session_factory() as session:
        job = session.get(IngestionJob, job_id)
        return job.status


# ---------------------------------------------------------------------------
# First run: everything is new.
# ---------------------------------------------------------------------------


def test_first_run_ingests_all_files_and_indexes(env: SimpleNamespace) -> None:
    doc1 = env.folder / "doc1.docx"
    doc2 = env.folder / "doc2.docx"
    _write_docx(doc1, "Section One", "Alpha content about penguins and glaciers.")
    _write_docx(doc2, "Section Two", "Beta content about volcanoes and coffee.")

    result = env.pipeline.run_ingestion_for_source(env.source.id)

    assert result.status == "succeeded"
    assert result.files_processed == 2
    assert result.files_failed == 0
    assert {r.status for r in result.file_results} == {"ingested"}
    assert all(r.chunks_written > 0 for r in result.file_results)

    assert _job_status(env, result.job_id) == IngestionJobStatus.SUCCEEDED

    with env.session_factory() as session:
        chunk_count = len(
            session.execute(
                select(Chunk.id).where(Chunk.source_id == env.source.id)
            ).scalars().all()
        )
    assert chunk_count == sum(r.chunks_written for r in result.file_results)
    assert chunk_count > 0

    # Chunks actually landed in both indexes (not just SQLite).
    assert len(env.gateway.embed_calls) > 0
    for file_path in (doc1, doc2):
        chunk_ids = _chunk_ids_for_file(env, file_path)
        assert chunk_ids
        assert env.vector.existing_chunk_ids(chunk_ids) == chunk_ids
        assert env.fts.existing_chunk_ids(chunk_ids) == chunk_ids


# ---------------------------------------------------------------------------
# Second run, unchanged: the critical no-re-embed regression assertion.
# ---------------------------------------------------------------------------


def test_second_run_unchanged_files_are_skipped_and_not_reembedded(
    env: SimpleNamespace,
) -> None:
    doc1 = env.folder / "doc1.docx"
    doc2 = env.folder / "doc2.docx"
    _write_docx(doc1, "Section One", "Alpha content about penguins and glaciers.")
    _write_docx(doc2, "Section Two", "Beta content about volcanoes and coffee.")

    first = env.pipeline.run_ingestion_for_source(env.source.id)
    assert {r.status for r in first.file_results} == {"ingested"}
    embed_calls_after_first_run = len(env.gateway.embed_calls)
    assert embed_calls_after_first_run > 0

    second = env.pipeline.run_ingestion_for_source(env.source.id)

    assert second.status == "succeeded"
    assert second.files_processed == 2
    assert second.files_failed == 0
    assert {r.status for r in second.file_results} == {"unchanged"}
    assert all(r.chunks_written == 0 for r in second.file_results)

    # The whole-pipeline regression assertion: re-running over unchanged
    # files must not trigger a single additional embed call anywhere in the
    # stack (evidence layer -> parser -> chunker -> index manager).
    assert len(env.gateway.embed_calls) == embed_calls_after_first_run


# ---------------------------------------------------------------------------
# Content change: re-ingest + reconcile stale chunks.
# ---------------------------------------------------------------------------


def test_modified_file_is_reingested_and_stale_chunks_are_reconciled(
    env: SimpleNamespace,
) -> None:
    doc1 = env.folder / "doc1.docx"
    doc2 = env.folder / "doc2.docx"
    _write_docx(doc1, "Section One", "Alpha content about penguins and glaciers.")
    _write_docx(doc2, "Section Two", "Beta content about volcanoes and coffee.")

    env.pipeline.run_ingestion_for_source(env.source.id)
    old_doc1_chunk_ids = _chunk_ids_for_file(env, doc1)
    old_doc2_chunk_ids = _chunk_ids_for_file(env, doc2)
    assert old_doc1_chunk_ids and old_doc2_chunk_ids

    _write_docx(doc1, "Section One Revised", "Completely different words: orchestras and tides.")
    result = env.pipeline.run_ingestion_for_source(env.source.id)

    results_by_path = {r.path: r for r in result.file_results}
    assert results_by_path[doc1].status == "ingested"
    assert results_by_path[doc2].status == "unchanged"

    new_doc1_chunk_ids = _chunk_ids_for_file(env, doc1)
    assert new_doc1_chunk_ids
    assert new_doc1_chunk_ids.isdisjoint(old_doc1_chunk_ids)

    # Old doc1 chunks are gone from both indexes...
    assert env.vector.existing_chunk_ids(old_doc1_chunk_ids) == set()
    assert env.fts.existing_chunk_ids(old_doc1_chunk_ids) == set()
    # ...new doc1 chunks are indexed...
    assert env.vector.existing_chunk_ids(new_doc1_chunk_ids) == new_doc1_chunk_ids
    assert env.fts.existing_chunk_ids(new_doc1_chunk_ids) == new_doc1_chunk_ids
    # ...and doc2's untouched chunks were NOT swept up in the reconcile.
    assert env.vector.existing_chunk_ids(old_doc2_chunk_ids) == old_doc2_chunk_ids
    assert env.fts.existing_chunk_ids(old_doc2_chunk_ids) == old_doc2_chunk_ids


# ---------------------------------------------------------------------------
# One bad file must not abort the batch.
# ---------------------------------------------------------------------------


def test_one_corrupt_file_fails_without_aborting_the_batch(env: SimpleNamespace) -> None:
    doc1 = env.folder / "doc1.docx"
    doc2 = env.folder / "doc2.docx"
    corrupt = env.folder / "corrupt.docx"
    _write_docx(doc1, "Section One", "Alpha content about penguins and glaciers.")
    _write_docx(doc2, "Section Two", "Beta content about volcanoes and coffee.")
    corrupt.write_text("this is plain text, not a real .docx file")

    result = env.pipeline.run_ingestion_for_source(env.source.id)

    assert result.files_processed == 3
    assert result.files_failed == 1
    assert result.status == "partial"
    assert _job_status(env, result.job_id) == IngestionJobStatus.PARTIAL

    results_by_path = {r.path: r for r in result.file_results}
    assert results_by_path[doc1].status == "ingested"
    assert results_by_path[doc2].status == "ingested"
    assert results_by_path[corrupt].status == "failed"
    assert results_by_path[corrupt].error is not None
